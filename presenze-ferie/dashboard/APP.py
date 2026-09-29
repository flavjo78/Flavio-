import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
import numpy as np
import calendar
import datetime
import io
import os
import json
import functools
import glob
import hashlib
import re
import secrets
import smtplib
import ssl
import urllib.parse
import email.message
import email.utils

# La dashboard lavora sempre nella cartella dove si trova questo file (anche se
# viene avviata da un'altra cartella): così il file di avvio non ha più bisogno
# del comando "pushd", che ad ogni avvio lasciava collegata una nuova lettera di
# unità di rete "web (\\ServerNas)" (N:, O:, P: ...).
try:
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
except Exception:
    pass

st.set_page_config(page_title="Dashboard Presenze", layout="wide", page_icon="📊")

# Etichetta di versione mostrata in cima alla pagina: serve a verificare a
# colpo d'occhio se il PC sta effettivamente eseguendo QUESTO file (utile
# quando si copia/sostituisce APP.py e si vuole essere sicuri che non stia
# girando una copia precedente rimasta in memoria o in un'altra cartella).
VERSIONE_APP = "1.04"  # aumentare ad ogni modifica (1.04, 1.05, ...)

MESI_IT = {
    1: "Gennaio", 2: "Febbraio", 3: "Marzo", 4: "Aprile",
    5: "Maggio", 6: "Giugno", 7: "Luglio", 8: "Agosto",
    9: "Settembre", 10: "Ottobre", 11: "Novembre", 12: "Dicembre"
}
GIORNI_IT = ["Lun", "Mar", "Mer", "Gio", "Ven", "Sab", "Dom"]
GIORNI_IT_COMPLETI = ["Lunedì", "Martedì", "Mercoledì", "Giovedì", "Venerdì", "Sabato", "Domenica"]
COLONNE_GIORNO_SETTIMANA = {"Lun": 0, "Mar": 1, "Mer": 2, "Gio": 3, "Ven": 4, "Sab": 5, "Dom": 6}
VALORI_OFF = {"OFF", "-", "RIPOSO", "X", "R"}


def _mese_precedente(anno, mese):
    return (anno - 1, 12) if mese == 1 else (anno, mese - 1)


def _mese_successivo(anno, mese):
    return (anno + 1, 1) if mese == 12 else (anno, mese + 1)


def _colonne_fisse(n, **kw):
    """Colonne che NON vanno a capo su schermi stretti (con fallback per
    versioni di Streamlit più vecchie)."""
    try:
        return st.columns(n, wrap=False, **kw)
    except TypeError:
        return st.columns(n, **kw)


def _scegli_giorno_calendario(data, chiave_popover, key_prefix):
    st.session_state["giorno_rif"] = data
    st.session_state[f"{key_prefix}_mese_vista"] = (data.year, data.month)
    # chiude il riquadro del calendario dopo la scelta
    try:
        st.session_state[chiave_popover] = False
    except Exception:
        pass


def _cambia_mese_calendario(key_prefix, anno, mese):
    st.session_state[f"{key_prefix}_mese_vista"] = (anno, mese)


def disegna_calendario(data_corrente, oggi, date_con_dati, key_prefix, chiave_popover):
    """Calendario mensile compatto in stile 'agenda': mese con frecce, giorni
    della settimana, giorni come cerchietti (blu = giorno selezionato, bordo
    blu = oggi, grigio chiaro = sabato/domenica o giorni senza timbrature,
    giorni futuri non cliccabili). L'aspetto è dato dal CSS in inietta_css()."""
    stato_key = f"{key_prefix}_mese_vista"
    rif_key = f"{key_prefix}_mese_vista_rif"
    if st.session_state.get(rif_key) != data_corrente or stato_key not in st.session_state:
        st.session_state[stato_key] = (data_corrente.year, data_corrente.month)
        st.session_state[rif_key] = data_corrente
    anno_v, mese_v = st.session_state[stato_key]

    with st.container(key=f"calendario_{key_prefix}"):
        col_p, col_lbl, col_s = _colonne_fisse([1, 4, 1], vertical_alignment="center")
        with col_p:
            st.button("‹", key=f"calnav_{key_prefix}_prec", type="tertiary", on_click=_cambia_mese_calendario,
                      args=(key_prefix, *_mese_precedente(anno_v, mese_v)))
        with col_lbl:
            st.markdown(f"<div class='calendario-titolo'>{MESI_IT[mese_v]} {anno_v}</div>", unsafe_allow_html=True)
        with col_s:
            st.button("›", key=f"calnav_{key_prefix}_succ", type="tertiary", on_click=_cambia_mese_calendario,
                      args=(key_prefix, *_mese_successivo(anno_v, mese_v)),
                      disabled=(anno_v, mese_v) >= (oggi.year, oggi.month))

        st.markdown(
            "<div class='calendario-intestazione'>" + "".join(f"<span>{g}</span>" for g in GIORNI_IT) + "</div>",
            unsafe_allow_html=True,
        )

        settimane = calendar.Calendar(firstweekday=0).monthdatescalendar(anno_v, mese_v)
        for settimana in settimane:
            colonne = _colonne_fisse(7, gap=None)
            for c, giorno in zip(colonne, settimana):
                with c:
                    if giorno.month != mese_v:
                        st.markdown("<div class='calendario-vuoto'></div>", unsafe_allow_html=True)
                        continue
                    if giorno == data_corrente:
                        tipo = "sel"
                    elif giorno > oggi:
                        tipo = "fut"
                    elif giorno == oggi:
                        tipo = "oggi"
                    elif nome_festivo(giorno):
                        tipo = "fe"
                    elif giorno.weekday() >= 5:
                        tipo = "we"
                    elif giorno not in date_con_dati:
                        tipo = "nd"
                    else:
                        tipo = "ok"
                    st.button(
                        str(giorno.day), key=f"calg_{tipo}_{key_prefix}_{giorno:%Y%m%d}",
                        disabled=(tipo == "fut"),
                        on_click=_scegli_giorno_calendario, args=(giorno, chiave_popover, key_prefix),
                    )

        festa = nome_festivo(data_corrente)
        if festa:
            st.markdown(f"<div class='calendario-festa'>🎉 {data_corrente:%d/%m} è festivo: {festa}</div>",
                        unsafe_allow_html=True)
        if st.session_state.get("modalita_correzione"):
            st.button(
                (f"➖ Togli festivo del {data_corrente:%d/%m}" if festa else f"🎉 Segna {data_corrente:%d/%m} come festivo"),
                key=f"calfesta_{key_prefix}", use_container_width=True,
                on_click=_cb_festivo, args=(data_corrente.isoformat(), not festa, "Festivo"),
            )
        col_o, col_i = st.columns(2)
        col_o.button("Oggi", key=f"calrapido_{key_prefix}_oggi", use_container_width=True,
                     on_click=_scegli_giorno_calendario, args=(oggi, chiave_popover, key_prefix))
        col_i.button("Ieri", key=f"calrapido_{key_prefix}_ieri", use_container_width=True,
                     on_click=_scegli_giorno_calendario,
                     args=(oggi - datetime.timedelta(days=1), chiave_popover, key_prefix))


def crea_popover(etichetta, key):
    try:
        return st.popover(etichetta, key=key, on_change="rerun", use_container_width=True)
    except TypeError:
        return st.popover(etichetta, use_container_width=True)


SOGLIA_RITARDO_DEFAULT = datetime.time(9, 15)

# ----------------------------------------------------------------------------
# Impostazioni modificabili dalla pagina "Impostazioni" (salvate in
# impostazioni.json nella cartella dell'app, quindi uguali per tutti i PC).
# ----------------------------------------------------------------------------
PERCORSO_IMPOSTAZIONI = "impostazioni.json"
OPZIONI_GIORNO_APERTURA = {
    "ieri": "Ieri",
    "ultimo_lavorativo": "Ultimo giorno lavorativo (se oggi è lunedì mostra venerdì)",
    "oggi": "Oggi",
}
IMPOSTAZIONI_PREDEFINITE = {
    "scostamento_massimo_min": 240,
    "tolleranza_ritardo_ingresso_min": 0,
    "tolleranza_ritardo_fine_pausa_min": 10,
    "tolleranza_uscita_anticipata_min": 5,
    "finestra_doppia_timbratura_sec": 60,
    "soglia_ritardo_default": "09:15",
    "giorno_apertura": "ieri",
    "soglia_mezza_giornata_pct": 60,
    "togli_pausa_se_non_timbrata": True,
    "festivi_nazionali": True,          # feste nazionali italiane automatiche
    "festivi_extra": {},                # {"2026-12-24": "Chiusura aziendale", ...}
    "festivi_esclusi": [],              # feste nazionali da NON considerare (date iso)
}


def carica_impostazioni():
    dati = dict(IMPOSTAZIONI_PREDEFINITE)
    try:
        if os.path.exists(PERCORSO_IMPOSTAZIONI):
            with open(PERCORSO_IMPOSTAZIONI, "r", encoding="utf-8") as f:
                letti = json.load(f)
            if isinstance(letti, dict):
                dati.update({k: v for k, v in letti.items() if k in IMPOSTAZIONI_PREDEFINITE})
    except Exception:
        pass
    return dati


def salva_impostazioni(dati):
    with open(PERCORSO_IMPOSTAZIONI, "w", encoding="utf-8") as f:
        json.dump(dati, f, ensure_ascii=False, indent=2)


IMPOSTAZIONI = carica_impostazioni()


def imp(chiave):
    return IMPOSTAZIONI.get(chiave, IMPOSTAZIONI_PREDEFINITE[chiave])


# ----------------------------------------------------------------------------
# Correzioni manuali delle timbrature e giustificativi (Malattia / Ferie).
# Le timbrature originali sul NAS NON vengono mai toccate: le correzioni sono
# salvate a parte (correzioni_timbrature.json nella cartella dell'app, quindi
# uguali per tutti i PC) e applicate "sopra" ai dati prima dei calcoli.
# Ogni correzione si può annullare.
# ----------------------------------------------------------------------------
PERCORSO_CORREZIONI = "correzioni_timbrature.json"
GIUSTIFICATIVI = {"M": "Malattia", "F": "Ferie", "F½": "Ferie ½ giornata"}


def carica_correzioni():
    vuoto = {"aggiunte": [], "eliminate": [], "giustificativi": {}}
    try:
        if os.path.exists(PERCORSO_CORREZIONI):
            with open(PERCORSO_CORREZIONI, "r", encoding="utf-8") as f:
                dati = json.load(f)
            if isinstance(dati, dict):
                for k in vuoto:
                    if isinstance(dati.get(k), type(vuoto[k])):
                        vuoto[k] = dati[k]
    except Exception:
        pass
    return vuoto


def salva_correzioni(correzioni):
    temporaneo = PERCORSO_CORREZIONI + ".tmp"
    with open(temporaneo, "w", encoding="utf-8") as f:
        json.dump(correzioni, f, ensure_ascii=False, indent=2)
    os.replace(temporaneo, PERCORSO_CORREZIONI)


def chiave_giorno(data, nome):
    return f"{data.isoformat()}|{nome}"


def applica_correzioni_eventi(df_eventi, correzioni):
    """Toglie le timbrature eliminate a mano e aggiunge quelle inserite a mano
    (tabella grezza Data/Orario/Nome/Evento, nomi già corretti)."""
    if not correzioni or (not correzioni.get("aggiunte") and not correzioni.get("eliminate")):
        return df_eventi
    df_eventi = df_eventi.copy()
    if not df_eventi.empty and correzioni.get("eliminate"):
        date_iso = pd.to_datetime(df_eventi["Data"], dayfirst=True, errors="coerce").dt.strftime("%Y-%m-%d")
        da_togliere = {(e["data"], e["nome"], str(e["orario"]).strip()) for e in correzioni["eliminate"]}
        chiavi = list(zip(date_iso, df_eventi["Nome"].astype(str).str.strip(), df_eventi["Orario"].astype(str).str.strip()))
        df_eventi = df_eventi[[k not in da_togliere for k in chiavi]]
    if correzioni.get("aggiunte"):
        nuove = pd.DataFrame([{
            "Data": datetime.date.fromisoformat(a["data"]).strftime("%d-%m-%Y"),
            "Orario": a["orario"], "Nome": a["nome"], "Evento": f'{a["evento"]} MANUALE',
        } for a in correzioni["aggiunte"]])
        df_eventi = pd.concat([df_eventi, nuove], ignore_index=True) if not df_eventi.empty else nuove
    return df_eventi


def codice_presenza(nome, data, riga, orari_lavoro, giustificativi, oggi):
    """Simbolo del giorno per un dipendente:
    M = malattia, F = ferie, F½ = mezza giornata di ferie (inseriti nella scheda
    Correzioni o approvati dalla scheda Ferie),
    P = presente, ½ = mezza giornata (ore lavorate sotto la soglia % delle ore
    previste, solo per giorni già conclusi), A = assente,
    - = giorno futuro o di riposo."""
    g = (giustificativi or {}).get(chiave_giorno(data, nome))
    if g in GIUSTIFICATIVI:
        return g
    if riga is not None:
        durata = riga.get("DurataTotale_sec") or 0
        if durata > 0 and data < oggi:
            previste = ore_previste_secondi(orari_lavoro, nome)
            if previste and durata < previste * float(imp("soglia_mezza_giornata_pct")) / 100.0:
                return "½"
        return "P"
    if data > oggi or not lavorativo_il(orari_lavoro, nome, data):
        return "-"
    return "A"


def controlla_orari(orari):
    """Segnala orari impostati in modo incoerente (es. uscita alle 06:30 con
    ingresso alle 07:30: probabilmente si voleva scrivere 16:30)."""
    avvisi = []
    for nome, info in sorted((orari or {}).items()):
        ing, usc = info.get("default"), info.get("uscita")
        p1, p2 = info.get("pausa_inizio"), info.get("pausa_fine")
        if ing and usc and usc <= ing:
            avvisi.append(f"**{nome}**: l'uscita ({usc:%H:%M}) è prima dell'ingresso ({ing:%H:%M}), "
                          "probabilmente è un errore di battitura (es. 06:30 invece di 16:30).")
        if p1 and p2 and p2 <= p1:
            avvisi.append(f"**{nome}**: la fine pausa ({p2:%H:%M}) è prima dell'inizio pausa ({p1:%H:%M}).")
        if ing and usc and usc > ing and ((p1 and not (ing < p1 < usc)) or (p2 and not (ing < p2 < usc))):
            avvisi.append(f"**{nome}**: la pausa non è compresa tra ingresso ({ing:%H:%M}) e uscita ({usc:%H:%M}).")
    return avvisi


# ----------------------------------------------------------------------------
# Giorni festivi: feste nazionali italiane calcolate in automatico (compresa
# la Pasquetta, che cambia ogni anno, e il 4 ottobre San Francesco, festa
# nazionale dal 2026) + giorni aggiunti a mano (santo patrono, ponti,
# chiusure aziendali). Un festivo non conta come assenza.
# ----------------------------------------------------------------------------

def _pasqua(anno):
    a, b, c = anno % 19, anno // 100, anno % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mese = (h + l - 7 * m + 114) // 31
    giorno = ((h + l - 7 * m + 114) % 31) + 1
    return datetime.date(anno, mese, giorno)


def festivi_nazionali(anno):
    fissi = {
        (1, 1): "Capodanno", (1, 6): "Epifania", (4, 25): "Festa della Liberazione",
        (5, 1): "Festa dei Lavoratori", (6, 2): "Festa della Repubblica", (8, 15): "Ferragosto",
        (11, 1): "Ognissanti", (12, 8): "Immacolata", (12, 25): "Natale", (12, 26): "Santo Stefano",
    }
    if anno >= 2026:
        fissi[(10, 4)] = "San Francesco (festa nazionale)"
    risultato = {datetime.date(anno, m, g): nome for (m, g), nome in fissi.items()}
    pasqua = _pasqua(anno)
    risultato[pasqua] = "Pasqua"
    risultato[pasqua + datetime.timedelta(days=1)] = "Lunedì dell'Angelo (Pasquetta)"
    return risultato


_CACHE_FESTIVI = {}


def festivi_anno(anno):
    """Tutti i festivi validi di un anno: {data: nome}."""
    if anno not in _CACHE_FESTIVI:
        risultato = {}
        if imp("festivi_nazionali"):
            esclusi = set(imp("festivi_esclusi") or [])
            risultato = {d: n for d, n in festivi_nazionali(anno).items() if d.isoformat() not in esclusi}
        for iso, nome in (imp("festivi_extra") or {}).items():
            try:
                d = datetime.date.fromisoformat(iso)
            except ValueError:
                continue
            if d.year == anno:
                risultato[d] = nome or "Festivo"
        _CACHE_FESTIVI[anno] = risultato
    return _CACHE_FESTIVI[anno]


def nome_festivo(data):
    return festivi_anno(data.year).get(data)


def lavorativo_il(orari_lavoro, nome, data):
    """Giorno lavorativo per quella persona in quella DATA (orario + festivi)."""
    return giorno_lavorativo_per(orari_lavoro, nome, data.weekday()) and not nome_festivo(data)


def _aggiorna_festivo(data_iso, rendi_festivo, nome=""):
    """Segna/toglie un festivo (salvato in impostazioni.json, uguale per tutti i PC)."""
    dati = carica_impostazioni()
    extra = dict(dati.get("festivi_extra") or {})
    esclusi = set(dati.get("festivi_esclusi") or [])
    d = datetime.date.fromisoformat(data_iso)
    e_nazionale = d in festivi_nazionali(d.year)
    if rendi_festivo:
        esclusi.discard(data_iso)
        if not e_nazionale:
            extra[data_iso] = nome or "Festivo"
    else:
        extra.pop(data_iso, None)
        if e_nazionale:
            esclusi.add(data_iso)
    dati["festivi_extra"] = extra
    dati["festivi_esclusi"] = sorted(esclusi)
    salva_impostazioni(dati)


def _sposta_orario(t, minuti):
    """Aggiunge (o toglie, se negativi) dei minuti a un orario, restando nella giornata."""
    if t is None:
        return None
    tot = t.hour * 60 + t.minute + float(minuti or 0)
    tot = max(0, min(tot, 23 * 60 + 59))
    return datetime.time(int(tot // 60), int(tot % 60), t.second if not minuti else 0)
PERCORSO_ORARI_SALVATI = "orari_lavoro.json"
PERCORSO_PASSWORD = "password_orari.txt"
PASSWORD_DEFAULT = "0"
PERCORSO_CONFIG_CARTELLA_LOG = "config_cartella_log.txt"


def carica_percorso_cartella_log():
    """Percorso (di rete o locale) della cartella dove 'salva_presenze.php'
    scrive un CSV per ogni timbratura. Se impostata e raggiungibile, la
    dashboard legge i dati da qui invece che dal file Excel: niente più
    bisogno di aprire Excel e premere "Aggiorna dati" per vedere le
    timbrature più recenti. Persistita su disco per sopravvivere ai riavvii."""
    if os.path.exists(PERCORSO_CONFIG_CARTELLA_LOG):
        try:
            with open(PERCORSO_CONFIG_CARTELLA_LOG, "r", encoding="utf-8") as f:
                return f.read().strip()
        except Exception:
            return ""
    return ""


def salva_percorso_cartella_log(percorso):
    with open(PERCORSO_CONFIG_CARTELLA_LOG, "w", encoding="utf-8") as f:
        f.write((percorso or "").strip())


PERCORSO_ALIAS_NOMI = "alias_nomi.json"
ETICHETTA_ALIAS_IGNORA = "🚫 Ignora (scarta queste timbrature)"


def carica_alias_nomi():
    """Mappa 'nome scritto male' -> 'nome corretto', usata per correggere
    automaticamente i refusi (es. un dispositivo che manda 'Paolo Angona'
    invece di 'Paolo Ancona') PRIMA di qualunque calcolo, così le timbrature
    non vengono perse ma vengono attribuite alla persona giusta."""
    if os.path.exists(PERCORSO_ALIAS_NOMI):
        try:
            with open(PERCORSO_ALIAS_NOMI, "r", encoding="utf-8") as f:
                dati = json.load(f)
            return dati if isinstance(dati, dict) else {}
        except Exception:
            return {}
    return {}


def salva_alias_nomi(mappa):
    with open(PERCORSO_ALIAS_NOMI, "w", encoding="utf-8") as f:
        json.dump(mappa, f, ensure_ascii=False, indent=2)


def applica_alias_nomi(df, alias_nomi):
    """Sostituisce nella colonna Nome ogni 'nome scritto male' con il
    corrispondente 'nome corretto' (confronto senza distinguere maiuscole/
    minuscole e ignorando spazi iniziali/finali). Se per un nome sbagliato è
    stato scelto ETICHETTA_ALIAS_IGNORA, le timbrature con quel nome vengono
    scartate del tutto invece di essere corrette."""
    if not alias_nomi or df.empty:
        return df
    mappa_lower = {
        str(sbagliato).strip().lower(): str(corretto).strip()
        for sbagliato, corretto in alias_nomi.items()
        if str(sbagliato).strip()
    }
    if not mappa_lower:
        return df

    chiave_lower = df["Nome"].astype(str).str.strip().str.lower()

    nomi_da_ignorare = {k for k, v in mappa_lower.items() if v == ETICHETTA_ALIAS_IGNORA}
    if nomi_da_ignorare:
        df = df[~chiave_lower.isin(nomi_da_ignorare)].copy()
        chiave_lower = df["Nome"].astype(str).str.strip().str.lower()

    df["Nome"] = [mappa_lower.get(k, n) for k, n in zip(chiave_lower, df["Nome"])]
    return df


def carica_password():
    """Legge la password che protegge la modifica degli Orari Dipendenti.
    Se non è mai stata cambiata, la password iniziale è '0'."""
    if os.path.exists(PERCORSO_PASSWORD):
        try:
            with open(PERCORSO_PASSWORD, "r", encoding="utf-8") as f:
                testo = f.read().strip()
            return testo if testo else PASSWORD_DEFAULT
        except Exception:
            return PASSWORD_DEFAULT
    return PASSWORD_DEFAULT


def salva_password(nuova_password):
    with open(PERCORSO_PASSWORD, "w", encoding="utf-8") as f:
        f.write(nuova_password.strip())


# ----------------------------------------------------------------------------
# FERIE (v1.04): richieste dei dipendenti dall'app "Le mie ferie" (cartella
# ferie/ sul NAS, vedi ferie/ e DOCUMENTAZIONE_LOGICA.md sezione 4).
# Contratto: l'APP scrive SOLO richieste_ferie/<id>.json. La DASHBOARD gestisce
# ferie_saldi.json, ferie_config.json, l'esito delle richieste e, quando approva,
# scrive "F" (giorno intero) o "F½" (mezza giornata) nei giustificativi di
# correzioni_timbrature.json per ogni giorno lavorativo del periodo.
# Le regole sui giorni sono quelle di sempre (lavorativo_il): app e dashboard
# contano i giorni allo stesso modo.
# ----------------------------------------------------------------------------
# FERIE-INIZIO
PERCORSO_FERIE_SALDI = "ferie_saldi.json"
CARTELLA_RICHIESTE_FERIE = "richieste_ferie"
PERCORSO_FERIE_CONFIG = "ferie_config.json"
FERIE_MEZZA = "F½"
PIN_ITERAZIONI = 100000
STATI_RICHIESTA = {"in_attesa": "⏳ In attesa", "approvata": "✅ Approvata",
                   "rifiutata": "❌ Rifiutata", "annullata": "↩️ Annullata"}
CONFIG_FERIE_PREDEFINITA = {
    "attiva": False, "smtp_host": "", "smtp_porta": 465, "smtp_sicurezza": "ssl",
    "smtp_utente": "", "smtp_password": "", "mittente": "", "mittente_nome": "Ferie",
    "destinatari_admin": [], "verifica_certificato": True, "url_app": "",
}


def _scrivi_json_atomico(percorso, dati):
    cartella = os.path.dirname(percorso)
    if cartella:
        os.makedirs(cartella, exist_ok=True)
    temporaneo = f"{percorso}.{secrets.token_hex(4)}.tmp"
    with open(temporaneo, "w", encoding="utf-8") as f:
        json.dump(dati, f, ensure_ascii=False, indent=2)
    os.replace(temporaneo, percorso)


def _leggi_json(percorso, predefinito):
    try:
        if os.path.exists(percorso):
            with open(percorso, "r", encoding="utf-8-sig") as f:
                dati = json.load(f)
            if predefinito is None or isinstance(dati, type(predefinito)):
                return dati
    except Exception:
        pass
    return predefinito


def formatta_giorni(valore):
    """5 -> '5', 0.5 -> '0,5', 12.5 -> '12,5' (i giorni di ferie possono essere a mezzi)."""
    return f"{float(valore):g}".replace(".", ",")


# ---- saldi, PIN, impostazioni -------------------------------------------------

def carica_saldi_ferie():
    return _leggi_json(PERCORSO_FERIE_SALDI, {})


def salva_saldi_ferie(saldi):
    _scrivi_json_atomico(PERCORSO_FERIE_SALDI, saldi)


def crea_pin_hash(pin):
    """PIN salvato come 'pbkdf2$iterazioni$sale$hash' (mai in chiaro). Lo stesso
    formato lo verifica l'app PHP (hash_pbkdf2 sha256)."""
    sale = secrets.token_bytes(16)
    h = hashlib.pbkdf2_hmac("sha256", str(pin).encode("utf-8"), sale, PIN_ITERAZIONI)
    return f"pbkdf2${PIN_ITERAZIONI}${sale.hex()}${h.hex()}"


def verifica_pin_hash(pin, pin_hash):
    try:
        nome, iterazioni, sale, atteso = str(pin_hash).split("$")
        if nome != "pbkdf2":
            return False
        h = hashlib.pbkdf2_hmac("sha256", str(pin).encode("utf-8"), bytes.fromhex(sale), int(iterazioni))
        return secrets.compare_digest(h.hex(), atteso.lower())
    except Exception:
        return False


def pin_valido(pin):
    return bool(re.fullmatch(r"[0-9]{4,8}", str(pin)))


def nuovo_pin_casuale(cifre=4):
    return "".join(secrets.choice("0123456789") for _ in range(cifre))


def carica_config_ferie():
    cfg = dict(CONFIG_FERIE_PREDEFINITA)
    cfg.update({k: v for k, v in _leggi_json(PERCORSO_FERIE_CONFIG, {}).items() if k in CONFIG_FERIE_PREDEFINITA})
    return cfg


def salva_config_ferie(cfg):
    _scrivi_json_atomico(PERCORSO_FERIE_CONFIG, cfg)


# ---- richieste ----------------------------------------------------------------

def _percorso_richiesta(id_richiesta):
    if not re.fullmatch(r"[A-Za-z0-9._-]+", str(id_richiesta)):
        raise ValueError("Identificativo di richiesta non valido.")
    return os.path.join(CARTELLA_RICHIESTE_FERIE, f"{id_richiesta}.json")


def leggi_richiesta_ferie(id_richiesta):
    r = _leggi_json(_percorso_richiesta(id_richiesta), None)
    return r if isinstance(r, dict) else None


def carica_richieste_ferie():
    """Tutte le richieste, dalla più recente. I file rovinati vengono ignorati."""
    risultato = []
    for percorso in glob.glob(os.path.join(CARTELLA_RICHIESTE_FERIE, "*.json")):
        r = _leggi_json(percorso, None)
        if isinstance(r, dict) and all(k in r for k in ("id", "nome", "dal", "al", "stato")):
            risultato.append(r)
    return sorted(risultato, key=lambda r: r.get("creata_il") or "", reverse=True)


def giorni_lavorativi_periodo(orari_lavoro, nome, dal, al):
    """Giorni lavorativi (orario + festivi) di una persona tra due date comprese."""
    return [dal + datetime.timedelta(days=k) for k in range((al - dal).days + 1)
            if lavorativo_il(orari_lavoro, nome, dal + datetime.timedelta(days=k))]


def giorni_richiesta_ferie(orari_lavoro, r):
    """(giorni lavorativi, peso di ogni giorno: 1 o 0,5) di una richiesta."""
    giorni = giorni_lavorativi_periodo(orari_lavoro, r["nome"], datetime.date.fromisoformat(r["dal"]),
                                       datetime.date.fromisoformat(r["al"]))
    return giorni, (0.5 if r.get("mezza_giornata") else 1.0)


def _valore_ferie(codice):
    return 1.0 if codice == "F" else 0.5 if codice == FERIE_MEZZA else 0.0


def saldo_ferie(nome, anno, orari_lavoro, giustificativi, richieste, saldi, oggi):
    """Saldo ferie dell'anno (stesse regole dell'app):
    godute = F con data <= oggi, programmate = F con data > oggi (F½ vale 0,5),
    in attesa = giorni delle richieste in attesa,
    disponibili = spettanti + residuo - godute - programmate."""
    voce = saldi.get(nome) or {}
    configurato = int(voce.get("anno") or 0) == anno and "giorni_spettanti" in voce
    spettanti = float(voce["giorni_spettanti"]) if configurato else 0.0
    residuo = float(voce.get("residuo_anno_precedente") or 0) if configurato else 0.0
    godute = programmate = in_attesa = 0.0
    for chiave, codice in (giustificativi or {}).items():
        data_iso, _, n = chiave.partition("|")
        if n != nome or not data_iso.startswith(f"{anno:04d}-"):
            continue
        if data_iso <= oggi.isoformat():
            godute += _valore_ferie(codice)
        else:
            programmate += _valore_ferie(codice)
    for r in richieste:
        if r["nome"] == nome and r["stato"] == "in_attesa":
            giorni, peso = giorni_richiesta_ferie(orari_lavoro, r)
            in_attesa += peso * sum(1 for g in giorni if g.year == anno)
    disponibili = spettanti + residuo - godute - programmate
    return {"configurato": configurato, "spettanti": spettanti, "residuo": residuo, "godute": godute,
            "programmate": programmate, "in_attesa": in_attesa, "disponibili": disponibili,
            "disponibili_se_approvate": disponibili - in_attesa}


def conflitti_richiesta_ferie(orari_lavoro, r, giustificativi):
    """Giorni della richiesta già segnati come malattia o ferie (da non sovrascrivere)."""
    giorni, _ = giorni_richiesta_ferie(orari_lavoro, r)
    trovati = []
    for g in giorni:
        codice = (giustificativi or {}).get(chiave_giorno(g, r["nome"]))
        if codice:
            trovati.append((g, codice))
    return trovati


def _ora_iso():
    return datetime.datetime.now().replace(microsecond=0).isoformat()


def approva_richiesta_ferie(id_richiesta, da, orari_lavoro):
    """Approva: scrive F (o F½) su ogni giorno lavorativo del periodo in
    correzioni_timbrature.json e segna la richiesta 'approvata'. Rilegge il
    file appena prima, così non approva una richiesta annullata nel frattempo.
    I giorni già segnati come malattia NON vengono sovrascritti.
    Ritorna (esito, messaggio)."""
    r = leggi_richiesta_ferie(id_richiesta)
    if r is None:
        return False, "Richiesta non trovata."
    if r["stato"] != "in_attesa":
        return False, f"La richiesta non è più in attesa (ora è: {STATI_RICHIESTA.get(r['stato'], r['stato'])})."
    giorni, _ = giorni_richiesta_ferie(orari_lavoro, r)
    if not giorni:
        return False, "Nel periodo non ci sono giorni lavorativi: niente da approvare."
    codice = FERIE_MEZZA if r.get("mezza_giornata") else "F"
    corr = carica_correzioni()
    saltati = []
    for g in giorni:
        chiave = chiave_giorno(g, r["nome"])
        if corr["giustificativi"].get(chiave) == "M":
            saltati.append(g)
            continue
        corr["giustificativi"][chiave] = codice
    salva_correzioni(corr)
    r.update({"stato": "approvata", "deciso_il": _ora_iso(), "deciso_da": da, "motivo_rifiuto": None})
    _scrivi_json_atomico(_percorso_richiesta(id_richiesta), r)
    scritti = len(giorni) - len(saltati)
    msg = f"✅ Approvata: {r['nome']}, {formatta_giorni(scritti * (0.5 if codice == FERIE_MEZZA else 1))} giorni di ferie segnati."
    if saltati:
        msg += " Non toccati perché già in malattia: " + ", ".join(f"{g:%d/%m}" for g in saltati) + "."
    return True, msg


def rifiuta_richiesta_ferie(id_richiesta, motivo, da):
    motivo = (motivo or "").strip()
    if not motivo:
        return False, "Scrivi il motivo del rifiuto: il dipendente lo vedrà nell'app."
    r = leggi_richiesta_ferie(id_richiesta)
    if r is None:
        return False, "Richiesta non trovata."
    if r["stato"] != "in_attesa":
        return False, f"La richiesta non è più in attesa (ora è: {STATI_RICHIESTA.get(r['stato'], r['stato'])})."
    r.update({"stato": "rifiutata", "deciso_il": _ora_iso(), "deciso_da": da, "motivo_rifiuto": motivo})
    _scrivi_json_atomico(_percorso_richiesta(id_richiesta), r)
    return True, f"❌ Rifiutata la richiesta di {r['nome']}."


# ---- email ---------------------------------------------------------------------

def invia_mail_ferie(destinatari, oggetto, testo, cfg=None):
    """Invia una mail di testo via SMTP con le impostazioni di ferie_config.json.
    Ritorna (esito, errore). Non lancia mai eccezioni: un problema di posta non
    deve bloccare l'approvazione."""
    cfg = cfg or carica_config_ferie()
    destinatari = [d.strip() for d in destinatari if d and re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", d.strip())]
    if not cfg.get("attiva"):
        return False, "notifiche email non attive"
    if not destinatari:
        return False, "nessun indirizzo valido"
    if not cfg.get("smtp_host") or not cfg.get("mittente"):
        return False, "server o mittente non impostati"
    try:
        msg = email.message.EmailMessage()
        msg["From"] = email.utils.formataddr((cfg.get("mittente_nome") or "", cfg["mittente"]))
        msg["To"] = ", ".join(destinatari)
        msg["Subject"] = oggetto
        msg.set_content(testo, charset="utf-8", cte="base64")  # accenti sicuri anche con server senza 8BITMIME
        contesto = ssl.create_default_context() if cfg.get("verifica_certificato", True) else ssl._create_unverified_context()
        sic = cfg.get("smtp_sicurezza", "ssl")
        porta = int(cfg.get("smtp_porta") or (465 if sic == "ssl" else 587))
        if sic == "ssl":
            server = smtplib.SMTP_SSL(cfg["smtp_host"], porta, timeout=10, context=contesto)
        else:
            server = smtplib.SMTP(cfg["smtp_host"], porta, timeout=10)
        with server:
            if sic == "tls":
                server.starttls(context=contesto)
            if cfg.get("smtp_utente"):
                server.login(cfg["smtp_utente"], cfg.get("smtp_password", ""))
            server.send_message(msg)
        return True, ""
    except Exception as e:
        return False, str(e)


def testo_periodo_ferie(r):
    p = datetime.date.fromisoformat(r["dal"]).strftime("%d/%m/%Y")
    if r["al"] != r["dal"]:
        p += " - " + datetime.date.fromisoformat(r["al"]).strftime("%d/%m/%Y")
    if r.get("mezza_giornata"):
        p += f" (mezza giornata, {r['mezza_giornata']})"
    return p


def notifica_esito_ferie(id_richiesta):
    """Avvisa il dipendente (se ha un'email in ferie_saldi.json) dell'esito. Ritorna una frase per il pannello."""
    r = leggi_richiesta_ferie(id_richiesta)
    if not r:
        return ""
    indirizzo = ((carica_saldi_ferie().get(r["nome"]) or {}).get("email") or "").strip()
    if not indirizzo:
        return ""
    if r["stato"] == "approvata":
        oggetto, testo = "Ferie approvate", f"Ciao {r['nome'].split()[0]},\nla tua richiesta di ferie ({testo_periodo_ferie(r)}) è stata APPROVATA.\n"
    else:
        oggetto = "Ferie non approvate"
        testo = (f"Ciao {r['nome'].split()[0]},\nla tua richiesta di ferie ({testo_periodo_ferie(r)}) non è stata approvata.\n"
                 f"Motivo: {r.get('motivo_rifiuto') or '-'}\n")
    ok, errore = invia_mail_ferie([indirizzo], oggetto, testo)
    return f" ✉️ Avviso inviato a {indirizzo}." if ok else f" (Email al dipendente non inviata: {errore}.)"
# FERIE-FINE


# ----------------------------------------------------------------------------
# Aspetto grafico: palette azzurra coerente in tutte le pagine
# ----------------------------------------------------------------------------

def inietta_css():
    st.markdown(
        """
        <style>
        .stApp { background-color: #f4f8fc; }
        section[data-testid="stSidebar"] { background-color: #eaf2fb; }
        h1, h2, h3, h4 { color: #14324d; }
        div[data-testid="stMetric"] {
            background-color: #ffffff; border: 1px solid #dbe7f3;
            border-radius: 10px; padding: 10px 14px;
        }
        .card {
            background-color: #ffffff; border: 1px solid #dbe7f3;
            border-radius: 12px; padding: 16px 20px; margin-bottom: 14px;
        }
        .dipendente-riga {
            display: flex; justify-content: space-between; align-items: center;
            padding: 10px 4px; border-bottom: 1px solid #eef2f7;
        }
        .dot { height: 11px; width: 11px; border-radius: 50%; display: inline-block; margin-right: 6px; }
        .dot-verde { background-color: #2ecc71; }
        .dot-rosso { background-color: #e74c3c; }
        .dot-arancio { background-color: #e69100; }
        .dot-azzurro { background-color: #3498db; }
        .dot-grigio { background-color: #b0bec5; }
        .dot-viola { background-color: #8e44ad; }
        .badge-giust {
            background-color: #efe3f7; color: #6c2d8c; border-radius: 6px;
            padding: 1px 7px; font-size: 0.78em; margin-left: 6px;
        }
        .dot-mini { height: 8px; width: 8px; margin-right: 3px; }
        .sotto-nome { color: #6b7d8f; font-size: 0.85em; }
        .timeline-riga {
            display: flex; align-items: center; gap: 12px; padding: 9px 0;
            border-bottom: 1px solid #eef2f7;
        }
        .timeline-orario { min-width: 55px; font-weight: 600; color: #14324d; }
        .timeline-punto { height: 9px; width: 9px; border-radius: 50%; background-color: #9db3c8; }
        .timeline-punto-verde { background-color: #2ecc71; }
        .timeline-punto-arancio { background-color: #e69100; }
        .timeline-punto-rosso { background-color: #e74c3c; }
        .timeline-riga-scartata { opacity: 0.55; text-decoration: line-through; }
        .badge-ritardo {
            background-color: #fdecdc; color: #a15c00; border-radius: 6px;
            padding: 1px 7px; font-size: 0.78em; margin-left: 6px;
        }
        .badge-uscita {
            background-color: #fde8e8; color: #a11616; border-radius: 6px;
            padding: 1px 7px; font-size: 0.78em; margin-left: 6px;
        }

        /* Menu principale (Oggi / Ricerca Storica / Orari / Impostazioni): SEMPRE
           visibile in alto anche scorrendo la pagina. Si "aggancia" sotto la barra
           di Streamlit (alta 3.75rem). Solo CSS: niente script che ricontrollano
           la pagina di continuo (quello di prima rallentava il browser e non
           funzionava con le versioni recenti di Streamlit). */
        .st-key-menu_principale > div > [role="tablist"] {
            position: sticky; top: 3.75rem; z-index: 999;
            background-color: #e3edf9; padding: 10px; gap: 10px; border-radius: 12px;
            box-shadow: 0 3px 8px rgba(20,50,77,0.15);
        }
        .st-key-menu_principale > div > [role="tablist"] [role="tab"] {
            height: 56px; padding: 0 28px; font-size: 1.15em;
        }
        .stTabs [role="tablist"] { gap: 6px; background-color: #e3edf9; padding: 6px; border-radius: 12px; }
        .stTabs [role="tab"] { border-radius: 9px; padding: 0 18px; font-weight: 600; color: #45607a; }
        .stTabs [role="tab"][aria-selected="true"] { background-color: #2f6fd1 !important; color: #ffffff !important; }
        .stTabs [role="tab"][aria-selected="true"] p { color: #ffffff !important; }

        /* Versione in alto a destra */
        .versione-app {
            position: fixed; top: 0.95rem; right: 9rem; z-index: 1000001;
            font-size: 0.75rem; font-weight: 700; color: #45607a; background: #e3edf9;
            padding: 2px 10px; border-radius: 10px; pointer-events: none;
        }
        .banner-correzione {
            background: #fff3cd; border: 2px solid #f0b400; color: #6b4e00; border-radius: 10px;
            padding: 8px 14px; margin: 6px 0 10px 0; font-size: 0.95em;
        }
        iframe[height="0"] { display: none; }

        /* ---- Calendario a comparsa (pulsante "Calendario" della pagina Oggi) ---- */
        .calendario-titolo { text-align: center; font-weight: 800; font-size: 1.05em; color: #14324d; }
        .calendario-intestazione {
            display: grid; grid-template-columns: repeat(7, 1fr); text-align: center;
            font-size: 0.75em; font-weight: 700; color: #7d8fa1; text-transform: uppercase;
            margin: 2px 0 4px 0;
        }
        .calendario-vuoto { height: 36px; }
        div[class*="st-key-calendario_"] [data-testid="stHorizontalBlock"] { gap: 2px !important; }
        div[class*="st-key-calendario_"] [data-testid="stColumn"] { min-width: 0 !important; }
        div[class*="st-key-calg_"] { display: flex; justify-content: center; }
        div[class*="st-key-calg_"] button {
            width: 36px; height: 36px; min-height: 36px; padding: 0; border-radius: 50%;
            border: none; background: transparent; color: #14324d; font-weight: 600;
            box-shadow: none; transition: background-color .12s;
        }
        div[class*="st-key-calg_"] button p { font-size: 0.9em; }
        div[class*="st-key-calg_"] button:hover { background: #e3edf9; color: #14324d; }
        div[class*="st-key-calg_we_"] button, div[class*="st-key-calg_nd_"] button { color: #a9b6c3; font-weight: 500; }
        div[class*="st-key-calg_oggi_"] button { box-shadow: inset 0 0 0 2px #2f6fd1; }
        div[class*="st-key-calg_sel_"] button, div[class*="st-key-calg_sel_"] button:hover {
            background: #2f6fd1; color: #ffffff;
        }
        div[class*="st-key-calg_fut_"] button { color: #d5dde5; background: transparent; }
        div[class*="st-key-calg_fe_"] button { color: #c0392b; font-weight: 800; }
        .calendario-festa { text-align: center; color: #c0392b; font-weight: 700; font-size: 0.9em; margin: 4px 0; }
        div[class*="st-key-calnav_"] button p { font-size: 1.5em; font-weight: 800; color: #2f6fd1; }
        div[data-testid="stPopoverBody"]:has(div[class*="st-key-calendario_"]) { min-width: 330px; }

        /* Tabs più "carine": segmented control azzurro invece delle linguette di default */
        .stTabs [data-baseweb="tab-list"] {
            gap: 6px; background-color: #e3edf9; padding: 6px; border-radius: 12px;
        }
        .stTabs [data-baseweb="tab"] {
            height: 44px; border-radius: 9px; padding: 0 18px; background-color: transparent;
            font-weight: 600; color: #45607a;
        }
        .stTabs [aria-selected="true"] {
            background-color: #2f6fd1 !important; color: #ffffff !important;
        }

        /* Bottoni-nome nella pagina Oggi: sembrano testo cliccabile, non pulsanti */
        .lista-dipendenti .stButton > button {
            background: transparent; border: none; padding: 4px 2px; text-align: left;
            font-weight: 700; color: #14324d; width: 100%;
        }
        .lista-dipendenti .stButton > button:hover { color: #2f6fd1; text-decoration: underline; }

        /* Calendario manuale (griglia di pulsanti) della pagina Oggi: celle più leggibili */
        div[data-testid="stVerticalBlockBorderWrapper"] .stButton > button {
            padding: 6px 4px;
        }
        .dot-giallo { background-color: #c7d92e; }
        .timeline-punto-giallo { background-color: #c7d92e; }

        /* Tabella "Riepilogo Generale": cella con entrata/uscita/ore su tre righe */
        .tabella-riepilogo-wrap {
            overflow-x: auto; max-width: 100%; border: 1px solid #dbe7f3;
            border-radius: 8px; margin-bottom: 10px;
        }
        .tabella-riepilogo { border-collapse: collapse; font-size: 0.8em; white-space: nowrap; }
        .tabella-riepilogo th, .tabella-riepilogo td {
            border: 1px solid #eef2f7; padding: 4px 8px; text-align: center;
        }
        .tabella-riepilogo thead th {
            background-color: #e3edf9; position: sticky; top: 0; z-index: 1; color: #14324d;
        }
        .tabella-riepilogo .cella-nome, .tabella-riepilogo .cella-nome-header {
            text-align: left; font-weight: 700; position: sticky; left: 0;
            background-color: #ffffff; z-index: 2; min-width: 140px; color: #14324d;
        }
        .tabella-riepilogo thead .cella-nome-header { z-index: 3; background-color: #e3edf9; }
        .tabella-riepilogo .cella-weekend { background-color: rgba(128,128,128,0.12); }
        .tabella-riepilogo .riga-e, .tabella-riepilogo .riga-u, .tabella-riepilogo .riga-o {
            line-height: 1.4;
        }
        .tabella-riepilogo .icona-mov { color: #8a97a3; font-weight: 700; margin-right: 3px; }
        /* Separazione più marcata tra la riga di un dipendente e quella successiva */
        .tabella-riepilogo tbody tr td { border-bottom: 3px solid #aebdcc; }

        </style>
        """,
        unsafe_allow_html=True,
    )


# ----------------------------------------------------------------------------
# Utility di formattazione
# ----------------------------------------------------------------------------

def formatta_orario(val):
    """Converte un valore eterogeneo (Timedelta, time, stringa, frazione di
    giorno excel...) in una stringa HH:MM leggibile, oppure '-' se assente."""
    if val is None:
        return "-"
    # pd.isna() copre in un colpo solo NaN, NaT (Timedelta/Timestamp mancante)
    # e None: prima veniva controllato solo il caso "float NaN", per cui un
    # valore NaT (tipico quando entrata o uscita mancano) sfuggiva al
    # controllo e finiva per essere mostrato a video come testo "NaT".
    try:
        if pd.isna(val):
            return "-"
    except (TypeError, ValueError):
        pass
    if isinstance(val, str) and val.strip() in ["", "-", "NaT", "None", "nan"]:
        return "-"
    try:
        if isinstance(val, pd.Timedelta):
            tot_sec = int(val.total_seconds())
        elif isinstance(val, (datetime.time,)):
            return val.strftime("%H:%M")
        elif isinstance(val, pd.Timestamp):
            return val.strftime("%H:%M")
        elif isinstance(val, (int, float)):
            tot_sec = int(val * 86400)
        else:
            td = pd.to_timedelta(str(val))
            tot_sec = int(td.total_seconds())
        h, m = tot_sec // 3600, (tot_sec % 3600) // 60
        return f"{h:02d}:{m:02d}"
    except Exception:
        txt = str(val).strip()
        return txt[:5] if len(txt) >= 5 else (txt or "-")


def a_time(val):
    """Prova a convertire un valore (Timedelta, time, stringa 'HH:MM' o
    'HH:MM:SS', frazione di giorno excel...) in datetime.time, altrimenti None."""
    if val is None:
        return None
    try:
        if pd.isna(val):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(val, datetime.time):
        return val
    if isinstance(val, pd.Timestamp):
        return val.time()
    if isinstance(val, str):
        testo = val.strip()
        for formato in ("%H:%M:%S", "%H:%M"):
            try:
                return datetime.datetime.strptime(testo, formato).time()
            except ValueError:
                continue
    try:
        if isinstance(val, (int, float)):
            tot_sec = int(val * 86400)
        else:
            td = pd.to_timedelta(str(val))
            tot_sec = int(td.total_seconds())
        if tot_sec < 0 or tot_sec >= 86400:
            return None
        h, m = tot_sec // 3600, (tot_sec % 3600) // 60
        return datetime.time(h, m)
    except Exception:
        return None


def a_secondi(val):
    """Converte una durata (Timedelta/stringa/numero) in secondi (float)."""
    if val is None:
        return 0.0
    try:
        if pd.isna(val):
            return 0.0
    except (TypeError, ValueError):
        pass
    try:
        if isinstance(val, pd.Timedelta):
            return val.total_seconds()
        if isinstance(val, (int, float)):
            return float(val) * 86400
        return pd.to_timedelta(str(val)).total_seconds()
    except Exception:
        return 0.0


def _lista_timbrature(r):
    """Restituisce l'elenco ordinato (per orario) di tutte le timbrature di
    entrata/uscita registrate in una riga (Entrata1, Uscita1, Entrata2, Uscita2),
    ciascuna etichettata con il segmento a cui appartiene (per la colorazione)."""
    eventi = []
    for campo, tipo, segmento in (
        ("Entrata1", "Entrata", "entrata"),
        ("Uscita1", "Uscita", "inizio_pausa"),
        ("Entrata2", "Entrata", "fine_pausa"),
        ("Uscita2", "Uscita", "uscita_finale"),
    ):
        t = a_time(r[campo])
        if t is not None:
            eventi.append((t, tipo, segmento))
    eventi.sort(key=lambda x: x[0])
    return eventi


_SEGMENTO_DA_SLOT = {
    "Ingresso": "entrata", "Inizio pausa": "inizio_pausa",
    "Fine pausa": "fine_pausa", "Uscita": "uscita_finale",
}


def _lista_timbrature_grezze(r):
    """Tutte le timbrature originali di un giorno, comprese quelle scartate
    (doppioni) o non riconosciute, per la vista di controllo nel log
    dettagliato: nulla sparisce dalla vista, ma ogni voce dice il suo esito.
    Ogni elemento: (orario, etichetta, stato, segmento) dove stato è
    'ok', 'corretto', 'doppia' o 'extra' (segmento è None per doppia/extra)."""
    grezze = r.get("TimbratureGrezze")
    if not isinstance(grezze, list) or not grezze:
        return [(t, tipo, "ok", segmento) for t, tipo, segmento in _lista_timbrature(r)]

    eventi = []
    for ann in grezze:
        t = a_time(ann.get("orario"))
        if t is None:
            continue
        manuale = " ✏️ (aggiunta a mano)" if "MANUALE" in str(ann.get("evento_originale", "")).upper() else ""
        if ann.get("scartata"):
            eventi.append((t, "Doppia timbratura", "doppia", None))
        elif ann.get("extra"):
            eventi.append((t, "Timbratura non riconosciuta" + manuale, "extra", None))
        else:
            slot = ann.get("slot")
            if not slot:
                continue
            stato = "corretto" if ann.get("corretto") else "ok"
            eventi.append((t, slot + manuale, stato, _SEGMENTO_DA_SLOT.get(slot)))
    eventi.sort(key=lambda x: x[0])
    return eventi


def _colore_evento(r, segmento):
    """Verde se l'orario è conforme, giallo-verde per un ritardo sul rientro
    dalla pausa (anomalia lieve), arancio per ritardo in entrata o uscita
    anticipata, rosso se il giorno presenta un'anomalia strutturale grave.
    L'inizio pausa non viene mai segnalato come ritardo."""
    if bool(r.get("AnomaliaGrave")):
        return "rosso"
    if segmento == "fine_pausa" and bool(r.get("RitardoFinePausa")):
        return "giallo"
    mappa_campo = {
        "entrata": "InRitardo",
        "uscita_finale": "UscitaAnticipata",
    }
    campo = mappa_campo.get(segmento)
    if campo and bool(r.get(campo)):
        return "arancio"
    return "verde"


def _colore_giorno(r):
    """Colore riassuntivo dell'intera giornata: il peggiore tra entrata,
    fine pausa e uscita finale."""
    ordine = {"rosso": 3, "arancio": 2, "giallo": 1, "verde": 0}
    colori = [
        _colore_evento(r, "entrata"),
        _colore_evento(r, "fine_pausa"),
        _colore_evento(r, "uscita_finale"),
    ]
    return max(colori, key=lambda c: ordine[c])


# ----------------------------------------------------------------------------
# Parsing dei log grezzi (foglio con colonne Data / Nome / Orario / Evento)
# ----------------------------------------------------------------------------



def _secondi_di_orario(val):
    t = a_time(val)
    if t is None:
        return None
    return t.hour * 3600 + t.minute * 60 + t.second


def _deduplica_eventi(eventi):
    """eventi: lista di (orario_grezzo, evento_str) già ordinata per orario.

    Se la stessa persona timbra due volte a meno di un minuto di distanza è,
    quasi certamente, un doppio click accidentale: la timbratura successiva
    viene scartata dal calcolo. Restituisce (eventi_puliti, eventi_annotati):
    - eventi_puliti: le timbrature da usare per i calcoli (senza i doppioni);
    - eventi_annotati: TUTTE le timbrature originali, ciascuna con il flag
      'scartata' (True se identificata come doppia), da mostrare per intero
      nel log dettagliato in modo che nulla sparisca dalla vista."""
    eventi_puliti = []
    eventi_annotati = []
    ultimo_secondi = None
    for orario, evento in eventi:
        secondi = _secondi_di_orario(orario)
        scartata = (
            ultimo_secondi is not None and secondi is not None
            and abs(secondi - ultimo_secondi) <= float(imp("finestra_doppia_timbratura_sec"))
        )
        eventi_annotati.append({"orario": orario, "evento_originale": str(evento), "scartata": scartata})
        if not scartata:
            eventi_puliti.append((orario, evento))
            ultimo_secondi = secondi
    return eventi_puliti, eventi_annotati

def elabora_log_grezzi(df_raw, alias_nomi=None):
    """Trasforma un log grezzo di timbrature (una riga per ogni ENTRA/ESCE)
    in una riga per dipendente/giorno, gestendo più turni (es. pausa pranzo)."""
    df_raw = df_raw.copy()
    if "Data" not in df_raw.columns or "Nome" not in df_raw.columns or "Orario" not in df_raw.columns:
        return pd.DataFrame()

    df_raw["Data"] = pd.to_datetime(df_raw["Data"], dayfirst=True, errors="coerce")
    df_raw = df_raw.dropna(subset=["Data", "Nome"])
    df_raw["Nome"] = df_raw["Nome"].astype(str).str.strip()
    # Corregge i nomi scritti male GIA' QUI, prima di raggruppare le
    # timbrature per Nome+Data: così, se lo stesso giorno c'è un turno
    # timbrato con il nome giusto e uno timbrato con un refuso, vengono
    # raggruppati insieme fin da subito come un'unica persona con due turni,
    # invece di restare due righe separate che nessun raggruppamento
    # successivo saprebbe più ricongiungere.
    df_raw = applica_alias_nomi(df_raw, alias_nomi)
    df_raw["Evento"] = (
        df_raw["Evento"].astype(str).str.upper().str.strip() if "Evento" in df_raw.columns else ""
    )

    records = []
    for (dt, nome), group in df_raw.groupby(["Data", "Nome"]):
        group = group.sort_values("Orario")
        eventi_grezzi = list(zip(group["Orario"], group.get("Evento", [""] * len(group))))
        eventi, eventi_annotati = _deduplica_eventi(eventi_grezzi)

        entrate, uscite = [], []
        for orario, evento in eventi:
            if "ENTRA" in str(evento):
                entrate.append(orario)
            elif "ESCE" in str(evento):
                uscite.append(orario)
        if not entrate and not uscite:
            orari = [o for o, _ in eventi]
            entrate, uscite = orari[0::2], orari[1::2]

        n_turni = max(len(entrate), len(uscite), 1)
        rec = {"Data": dt, "Nome": nome}
        durata_tot_sec = 0.0
        anomalia = []

        for i in range(n_turni):
            e = entrate[i] if i < len(entrate) else None
            u = uscite[i] if i < len(uscite) else None
            durata_sec = 0.0
            if e is not None and u is not None:
                try:
                    t_e, t_u = pd.to_timedelta(str(e)), pd.to_timedelta(str(u))
                    if t_u > t_e:
                        durata_sec = (t_u - t_e).total_seconds()
                    else:
                        anomalia.append(f"Orari incoerenti (turno {i + 1})")
                except Exception:
                    pass
            elif e is not None and u is None:
                anomalia.append(f"Uscita mancante (turno {i + 1})")

            durata_tot_sec += durata_sec
            rec[f"Entrata{i + 1}"] = e
            rec[f"Uscita{i + 1}"] = u
            rec[f"Durata{i + 1}"] = pd.Timedelta(seconds=durata_sec) if durata_sec > 0 else None

        rec["DurataTotale"] = pd.Timedelta(seconds=durata_tot_sec) if durata_tot_sec > 0 else None
        rec["Anomalia"] = "; ".join(anomalia) if anomalia else "OK"
        rec["TimbratureGrezze"] = eventi_annotati
        rec["EventiPuliti"] = eventi
        records.append(rec)

    return pd.DataFrame(records)


# ----------------------------------------------------------------------------
# Lettura diretta della cartella di rete dove "salva_presenze.php" scrive un
# file CSV per ogni singola timbratura inviata dall'app (uno o più righe a
# file). È la stessa cartella che Power Query legge dentro presenze.xlsx per
# la query "presenze" - qui la leggiamo direttamente in Python, così i dati
# sono sempre quelli reali, senza il passaggio manuale "Aggiorna dati" in
# Excel. Ogni file è fatto così:
#   sep=;                              <- riga guida per Excel, va ignorata
#   Data;Orario;Nome;Evento            <- intestazione
#   09-06-2026;08:06:01;Mario Rossi ;ENTRA   <- una o più righe di dati
# ----------------------------------------------------------------------------

def _leggi_righe_csv_timbrature(percorso_file):
    """Legge un singolo file CSV di timbrature grezze. Restituisce un
    DataFrame con le colonne del file (tipicamente Data/Orario/Nome/Evento),
    oppure None se il file è vuoto/illeggibile/malformato: viene semplicemente
    saltato, non deve bloccare la lettura di tutti gli altri file."""
    try:
        with open(percorso_file, "r", encoding="utf-8-sig", errors="replace") as f:
            testo = f.read()
    except Exception:
        return None

    righe = [r for r in testo.splitlines() if r.strip() != ""]
    if not righe:
        return None
    if righe[0].strip().lower().startswith("sep="):
        righe = righe[1:]
    if len(righe) < 2:
        return None

    intestazione = [c.strip() for c in righe[0].split(";")]
    righe_dati = []
    for riga in righe[1:]:
        valori = riga.split(";")
        if len(valori) < len(intestazione):
            continue
        righe_dati.append(valori[: len(intestazione)])
    if not righe_dati:
        return None

    return tuple(intestazione), righe_dati



def _parse_singolo_csv_timbrature(percorso_file):
    letto = _leggi_righe_csv_timbrature(percorso_file)
    if letto is None:
        return None
    intestazione, righe = letto
    return pd.DataFrame(righe, columns=list(intestazione))


SECONDI_AGGIORNAMENTO_CARTELLA = 30


@st.cache_data(ttl=SECONDI_AGGIORNAMENTO_CARTELLA, show_spinner=False)
def elenco_file_cartella(percorso_cartella):
    """Elenco dei file .csv della cartella con data di modifica e dimensione,
    ottenuto con UNA sola lettura della cartella (os.scandir: su Windows date e
    dimensioni arrivano insieme all'elenco, senza chiedere al NAS file per
    file). Prima invece, ad ogni clic, si chiedeva al NAS la data di ogni
    singolo file, due volte: con centinaia/migliaia di file in rete erano
    migliaia di richieste e diversi secondi di attesa. Il risultato resta
    valido per 30 secondi: le nuove timbrature compaiono entro mezzo minuto
    (o subito con il pulsante 'Ricarica dati')."""
    esito = {"errore": None, "voci": []}
    try:
        with os.scandir(percorso_cartella) as elenco:
            for voce in elenco:
                if not voce.name.lower().endswith(".csv"):
                    continue
                try:
                    info = voce.stat()
                    esito["voci"].append((voce.name, info.st_mtime, info.st_size))
                except OSError:
                    esito["voci"].append((voce.name, 0.0, -1))
    except Exception as e:
        esito["errore"] = str(e)
    esito["voci"].sort()
    return esito


PERCORSO_ARCHIVIO_LOCALE = os.path.join(
    __import__("tempfile").gettempdir(), "dashboard_presenze_archivio_timbrature.pkl"
)


@st.cache_resource(show_spinner=False)
def _archivio_file_letti():
    """Memoria dei file CSV già letti, identificati da nome + data + dimensione:
    un file già letto non viene mai più riletto dal NAS finché non cambia.
    Ad ogni aggiornamento si leggono quindi solo i file NUOVI. La memoria
    viene anche salvata sul disco di QUESTO computer (cartella temporanea di
    Windows), così anche alla prima apertura del mattino non si rileggono
    dal NAS tutti i file dei mesi passati, ma solo quelli arrivati dopo."""
    try:
        import pickle
        with open(PERCORSO_ARCHIVIO_LOCALE, "rb") as f:
            dati = pickle.load(f)
        if isinstance(dati, dict):
            return dati
    except Exception:
        pass
    return {}


def _salva_archivio_su_disco(archivio):
    try:
        import pickle
        temporaneo = PERCORSO_ARCHIVIO_LOCALE + ".tmp"
        with open(temporaneo, "wb") as f:
            pickle.dump(archivio, f, protocol=pickle.HIGHEST_PROTOCOL)
        os.replace(temporaneo, PERCORSO_ARCHIVIO_LOCALE)
    except Exception:
        pass


def _righe_file_con_archivio(percorso_file, mtime, dimensione):
    archivio = _archivio_file_letti()
    chiave = (percorso_file, mtime, dimensione)
    if chiave not in archivio:
        archivio[chiave] = _leggi_righe_csv_timbrature(percorso_file)
    return archivio[chiave]


def carica_eventi_grezzi_cartella(percorso_cartella, voci):
    """Unisce le timbrature di tutti i file della cartella in un'unica tabella
    (Data/Orario/Nome/Evento), leggendo dal NAS solo i file non ancora letti."""
    archivio = _archivio_file_letti()
    chiavi = [(os.path.join(percorso_cartella, n), m, d) for n, m, d in (voci or [])]
    da_leggere = [k for k in chiavi if k not in archivio]
    if da_leggere:
        # I file nuovi si leggono in parallelo (16 alla volta): in rete il
        # tempo è quasi tutto attesa del NAS, quindi farlo insieme è molto
        # più veloce che uno dopo l'altro.
        from concurrent.futures import ThreadPoolExecutor
        with ThreadPoolExecutor(max_workers=16) as esecutore:
            letti = list(esecutore.map(lambda k: _leggi_righe_csv_timbrature(k[0]), da_leggere))
        for k, letto in zip(da_leggere, letti):
            archivio[k] = letto
        # teniamo in memoria solo i file che esistono ancora (niente accumulo)
        attuali = set(chiavi)
        for k in [k for k in archivio if k[0].startswith(percorso_cartella) and k not in attuali]:
            del archivio[k]
        _salva_archivio_su_disco(archivio)

    righe_per_intestazione = {}
    for chiave in chiavi:
        letto = archivio.get(chiave)
        if letto is None:
            continue
        intestazione, righe = letto
        righe_per_intestazione.setdefault(intestazione, []).extend(righe)
    if not righe_per_intestazione:
        return pd.DataFrame()
    tabelle = [pd.DataFrame(righe, columns=list(intest)) for intest, righe in righe_per_intestazione.items()]
    df_tutto = pd.concat(tabelle, ignore_index=True) if len(tabelle) > 1 else tabelle[0]
    if not {"Data", "Orario", "Nome", "Evento"}.issubset(df_tutto.columns):
        return pd.DataFrame()
    df_tutto["Nome"] = df_tutto["Nome"].astype(str).str.strip()
    df_tutto["Evento"] = df_tutto["Evento"].astype(str).str.strip()
    df_tutto = df_tutto[df_tutto["Nome"] != ""]
    return df_tutto


def diagnostica_cartella_log(percorso_cartella, elenco=None):
    """Controlla cosa vede davvero il processo Python sulla cartella di rete:
    quanti file .csv ci sono, qual è il più recente e se quel file si legge
    correttamente (così si distingue subito tra cartella non raggiungibile,
    cartella ferma a una certa data, o formato di un file non riconosciuto).
    Usa l'elenco già letto (una sola richiesta al NAS) invece di chiedere la
    data di ogni file uno per uno. Se la cartella non è raggiungibile,
    'errore' contiene il vero messaggio di Windows (es. WinError 3/5/53)."""
    info = {
        "raggiungibile": False, "n_file": 0, "file_piu_recente": None,
        "data_piu_recente": None, "file_piu_recente_leggibile": None,
        "errore": None,
    }
    if not percorso_cartella:
        return info
    if elenco is None:
        elenco = elenco_file_cartella(percorso_cartella)
    if elenco["errore"]:
        info["errore"] = elenco["errore"]
        return info
    voci = elenco["voci"]
    info["raggiungibile"] = True
    info["n_file"] = len(voci)
    if not voci:
        return info
    nome, mtime, dimensione = max(voci, key=lambda v: v[1])
    info["file_piu_recente"] = nome
    if mtime:
        info["data_piu_recente"] = datetime.datetime.fromtimestamp(mtime)
    try:
        letto = _righe_file_con_archivio(os.path.join(percorso_cartella, nome), mtime, dimensione)
        info["file_piu_recente_leggibile"] = bool(letto and letto[1])
    except Exception as e:
        info["errore"] = str(e)
    return info


# ----------------------------------------------------------------------------
# Caricamento e normalizzazione del file Excel (fogli generati dall'app Android)
# ----------------------------------------------------------------------------

COLONNE_STD = [
    "Data", "Nome", "DurataTotale", "Anomalia",
    "Entrata1", "Uscita1", "Durata1", "Entrata2", "Uscita2", "Durata2",
]


def _trova_intestazione(df_raw, max_righe=20):
    """Cerca fra le prime righe del foglio quella che sembra un'intestazione,
    restituendo (indice_riga, lista_valori) oppure (None, None).

    Nei fogli pivot di Excel/Power Query la riga sopra l'intestazione vera
    contiene spesso etichette come "Somma di Durata1" o "Etichette di
    colonna": per questo confrontiamo le celle per uguaglianza esatta
    (non per sottostringa) così da non scambiarle per l'intestazione vera."""
    for idx in range(min(max_righe, len(df_raw))):
        valori = [str(v).strip() for v in df_raw.iloc[idx].values if pd.notna(v) and str(v).strip() != ""]
        if len(valori) >= 2:
            valori_lower = [v.lower() for v in valori]
            ha_data_e_nome = "data" in valori_lower and "nome" in valori_lower
            e_pivot_righe = valori_lower[0] == "etichette di riga"
            if ha_data_e_nome or e_pivot_righe:
                return idx, valori
    return None, None


def _estrai_tabella(df_raw, idx):
    df = df_raw.iloc[idx + 1:].copy()
    df.columns = [str(c).strip() for c in df_raw.iloc[idx].values]
    return df.dropna(how="all")


def _parse_foglio(df_raw):
    """Analizza un singolo foglio e restituisce (dataframe_normalizzato, qualita)
    dove qualita è un intero: 3 = formato ricco (Entrata1/Uscita1...),
    2 = log grezzo elaborato, 1 = solo pivot con il totale ore, 0 = niente."""
    idx, intestazione = _trova_intestazione(df_raw)
    if idx is None:
        return pd.DataFrame(), 0

    colonne_lower = [c.lower() for c in intestazione]

    if any("data" in c for c in colonne_lower) and any("nome" in c for c in colonne_lower):
        df = _estrai_tabella(df_raw, idx)
        if "Entrata1" in df.columns or "entrata1" in [c.lower() for c in df.columns]:
            df["Data"] = pd.to_datetime(df["Data"], dayfirst=True, errors="coerce")
            df = df.dropna(subset=["Data", "Nome"])
            df["Nome"] = df["Nome"].astype(str).str.strip()
            if "Anomalia" not in df.columns:
                df["Anomalia"] = "OK"
            for col in COLONNE_STD:
                if col not in df.columns:
                    df[col] = None
            return df[COLONNE_STD], 3

        if "Orario" in df.columns or "Evento" in df.columns:
            df_el = elabora_log_grezzi(df)
            if df_el.empty:
                return pd.DataFrame(), 0
            for col in COLONNE_STD:
                if col not in df_el.columns:
                    df_el[col] = None
            return df_el[COLONNE_STD], 2

        df["Data"] = pd.to_datetime(df["Data"], dayfirst=True, errors="coerce")
        df = df.dropna(subset=["Data", "Nome"])
        df["Nome"] = df["Nome"].astype(str).str.strip()
        for col in COLONNE_STD:
            if col not in df.columns:
                df[col] = None
        return df[COLONNE_STD], 1

    if any("etichette" in c for c in colonne_lower) or any("data" in c for c in colonne_lower):
        df = _estrai_tabella(df_raw, idx)
        cols = list(df.columns)
        first_col = cols[0]
        df_melted = df.melt(id_vars=[first_col], var_name="Nome", value_name="DurataTotale")
        df_melted = df_melted.rename(columns={first_col: "Data"})
        df_melted["Data"] = pd.to_datetime(df_melted["Data"], dayfirst=True, errors="coerce")
        df_melted = df_melted.dropna(subset=["Data", "Nome", "DurataTotale"])
        df_melted["Nome"] = df_melted["Nome"].astype(str).str.strip()
        df_melted = df_melted[~df_melted["Nome"].str.lower().str.contains("totale", na=False)]
        for col in COLONNE_STD:
            if col not in df_melted.columns:
                df_melted[col] = None
        return df_melted[COLONNE_STD], 1

    return pd.DataFrame(), 0


@st.cache_data(show_spinner="Caricamento e analisi del file presenze...")
def load_all_data(file_bytes, file_name):
    """Legge tutti i fogli del file Excel e restituisce un unico dataframe
    normalizzato, dando priorità ai fogli con il dettaglio più ricco
    (Entrata1/Uscita1/Durata2...) rispetto ai riepiloghi grezzi o ai pivot."""
    excel_file = pd.ExcelFile(io.BytesIO(file_bytes))
    candidati = []

    for sheet_name in excel_file.sheet_names:
        if sheet_name.strip().lower() == "orarilavoro":
            continue
        try:
            df_raw = pd.read_excel(excel_file, sheet_name=sheet_name, header=None)
        except Exception:
            continue
        if df_raw.empty:
            continue
        try:
            df_norm, qualita = _parse_foglio(df_raw)
        except Exception:
            continue
        if qualita > 0 and not df_norm.empty:
            candidati.append((qualita, df_norm))

    if not candidati:
        return pd.DataFrame(columns=COLONNE_STD)

    candidati.sort(key=lambda x: x[0])
    df_completo = pd.concat([d for _, d in candidati], ignore_index=True)
    df_completo = df_completo.drop_duplicates(subset=["Data", "Nome"], keep="last")
    return df_completo.reset_index(drop=True)


# ----------------------------------------------------------------------------
# Orario di lavoro previsto per dipendente: lettura dal foglio Excel
# "OrariLavoro" (bootstrap iniziale) + gestione modificabile e persistente
# in un file separato "orari_lavoro.json", per non rischiare di scrivere
# sopra il file presenze.xlsx che viene rigenerato dall'app Android.
# ----------------------------------------------------------------------------

def _individua_colonne_orari(intestazioni):
    col_nome = col_ingresso = col_uscita = None
    col_pausa_inizio = col_pausa_fine = col_pausa_extra = None
    col_giorni = {}

    for i, grezzo in enumerate(intestazioni):
        testo = str(grezzo).strip().lower() if pd.notna(grezzo) else ""
        if not testo:
            continue
        if testo == "nome":
            col_nome = i
        elif "ingres" in testo or "ingrasso" in testo:
            col_ingresso = i
        elif "uscita" in testo:
            col_uscita = i
        elif "inizio" in testo and "pausa" in testo:
            col_pausa_inizio = i
        elif "fine" in testo and "pausa" in testo:
            col_pausa_fine = i
        elif "altr" in testo and "pausa" in testo:
            col_pausa_extra = i
        else:
            for abbr, weekday in COLONNE_GIORNO_SETTIMANA.items():
                if testo.startswith(abbr.lower()):
                    col_giorni[weekday] = i
                    break

    return {
        "nome": col_nome, "ingresso": col_ingresso, "uscita": col_uscita,
        "pausa_inizio": col_pausa_inizio, "pausa_fine": col_pausa_fine,
        "pausa_extra": col_pausa_extra, "giorni": col_giorni,
    }


def _durata_a_secondi(t):
    return None if t is None else t.hour * 3600 + t.minute * 60 + t.second


@st.cache_data(show_spinner=False)
def carica_orari_da_excel(file_bytes, file_name):
    """Legge il foglio 'OrariLavoro' (o 'orarilavoro') dal file Excel, se
    presente. Usato solo come base di partenza iniziale: dopo il primo
    salvataggio dalla pagina 'Orari Dipendenti', fa fede il file
    orari_lavoro.json (vedi carica_orari_lavoro_effettivi)."""
    try:
        excel_file = pd.ExcelFile(io.BytesIO(file_bytes))
    except Exception:
        return None
    nome_foglio = next(
        (s for s in excel_file.sheet_names if s.strip().lower() == "orarilavoro"), None
    )
    if nome_foglio is None:
        return None

    try:
        df_raw = pd.read_excel(excel_file, sheet_name=nome_foglio, header=None)
    except Exception:
        return None
    if df_raw.empty:
        return None

    idx_intestazione = None
    for i in range(min(10, len(df_raw))):
        valori = [str(v).strip().lower() for v in df_raw.iloc[i].values if pd.notna(v)]
        if any(("orario" in v) or ("ingres" in v) or ("ingrasso" in v) for v in valori):
            idx_intestazione = i
            break
    if idx_intestazione is None:
        return None

    intestazioni = list(df_raw.iloc[idx_intestazione].values)
    colonne = _individua_colonne_orari(intestazioni)
    dati = df_raw.iloc[idx_intestazione + 1:].reset_index(drop=True)

    if colonne["nome"] is None:
        colonne_usate = {c for c in (
            colonne["ingresso"], colonne["uscita"], colonne["pausa_inizio"],
            colonne["pausa_fine"], colonne["pausa_extra"],
        ) if c is not None}
        colonne_usate |= set(colonne["giorni"].values())
        migliore, punteggio_migliore = None, -1
        for c in range(dati.shape[1]):
            if c in colonne_usate:
                continue
            punteggio = sum(
                1 for v in dati[c].dropna()
                if isinstance(v, str) and len(v.strip()) > 2 and a_time(v) is None
            )
            if punteggio > punteggio_migliore:
                migliore, punteggio_migliore = c, punteggio
        colonne["nome"] = migliore

    if colonne["nome"] is None or colonne["ingresso"] is None:
        return None

    risultato = {}
    for _, riga in dati.iterrows():
        nome = riga[colonne["nome"]]
        if pd.isna(nome) or str(nome).strip() == "":
            continue
        nome = str(nome).strip()

        ingresso = a_time(riga[colonne["ingresso"]])
        uscita = a_time(riga[colonne["uscita"]]) if colonne["uscita"] is not None else None
        pausa_inizio = a_time(riga[colonne["pausa_inizio"]]) if colonne["pausa_inizio"] is not None else None
        pausa_fine = a_time(riga[colonne["pausa_fine"]]) if colonne["pausa_fine"] is not None else None
        pausa_extra = a_time(riga[colonne["pausa_extra"]]) if colonne["pausa_extra"] is not None else None

        per_day = {}
        for weekday, c in colonne["giorni"].items():
            val = riga[c]
            if pd.isna(val) or str(val).strip() == "":
                continue
            testo = str(val).strip().upper()
            if testo in VALORI_OFF:
                per_day[weekday] = "OFF"
            else:
                t = a_time(val)
                if t is not None:
                    per_day[weekday] = t

        risultato[nome] = {
            "default": ingresso,
            "uscita": uscita,
            "pausa_inizio": pausa_inizio,
            "pausa_fine": pausa_fine,
            "pausa_extra_sec": _durata_a_secondi(pausa_extra) or 0.0,
            "per_day": per_day,
        }

    return risultato if risultato else None


def _time_a_stringa(t):
    return t.strftime("%H:%M") if isinstance(t, datetime.time) else None


def _stringa_a_time(s):
    return a_time(s) if s else None


def salva_orari_lavoro(orari):
    """Scrive l'orario di lavoro corrente su orari_lavoro.json (formato
    leggibile, indipendente dal presenze.xlsx che viene rigenerato in
    automatico dall'app Android)."""
    serializzabile = {}
    for nome, info in orari.items():
        per_day_ser = {
            str(wd): (v if v == "OFF" else _time_a_stringa(v))
            for wd, v in info.get("per_day", {}).items()
        }
        serializzabile[nome] = {
            "ingresso": _time_a_stringa(info.get("default")),
            "uscita": _time_a_stringa(info.get("uscita")),
            "pausa_inizio": _time_a_stringa(info.get("pausa_inizio")),
            "pausa_fine": _time_a_stringa(info.get("pausa_fine")),
            "pausa_extra_sec": info.get("pausa_extra_sec", 0.0),
            "per_day": per_day_ser,
        }
    with open(PERCORSO_ORARI_SALVATI, "w", encoding="utf-8") as f:
        json.dump(serializzabile, f, ensure_ascii=False, indent=2)


def carica_orari_da_json():
    """Legge orari_lavoro.json se esiste (creato da un salvataggio precedente
    dalla pagina 'Orari Dipendenti'), altrimenti restituisce None."""
    if not os.path.exists(PERCORSO_ORARI_SALVATI):
        return None
    try:
        with open(PERCORSO_ORARI_SALVATI, "r", encoding="utf-8") as f:
            dati = json.load(f)
    except Exception:
        return None

    risultato = {}
    for nome, info in dati.items():
        per_day = {}
        for wd_str, v in info.get("per_day", {}).items():
            per_day[int(wd_str)] = "OFF" if v == "OFF" else _stringa_a_time(v)
        risultato[nome] = {
            "default": _stringa_a_time(info.get("ingresso")),
            "uscita": _stringa_a_time(info.get("uscita")),
            "pausa_inizio": _stringa_a_time(info.get("pausa_inizio")),
            "pausa_fine": _stringa_a_time(info.get("pausa_fine")),
            "pausa_extra_sec": info.get("pausa_extra_sec", 0.0),
            "per_day": per_day,
        }
    return risultato if risultato else None


def _secondi_a_durata_stringa(sec):
    sec = int(sec or 0)
    h, m = sec // 3600, (sec % 3600) // 60
    return f"{h:02d}:{m:02d}:00"


def scrivi_orari_in_excel(percorso_file, orari):
    """Scrive (o sostituisce) il foglio 'OrariLavoro' direttamente nel file
    Excel indicato, lasciando intatti tutti gli altri fogli. Usata quando
    l'app lavora sul file locale presenze.xlsx (non su un file caricato dal
    browser, per il quale non abbiamo un percorso su cui scrivere)."""
    from openpyxl import load_workbook

    wb = load_workbook(percorso_file)
    for nome_foglio in list(wb.sheetnames):
        if nome_foglio.strip().lower() == "orarilavoro":
            del wb[nome_foglio]
    ws = wb.create_sheet("OrariLavoro")
    ws.append(["Nome", "OrarioIngresso", "OrarioInizioPausa", "OrarioFinePausa",
               "OrarioUscita", "AltraPausa"] + GIORNI_IT)

    for nome, info in sorted(orari.items()):
        riga = [
            nome,
            _time_a_stringa(info.get("default")) or "",
            _time_a_stringa(info.get("pausa_inizio")) or "",
            _time_a_stringa(info.get("pausa_fine")) or "",
            _time_a_stringa(info.get("uscita")) or "",
            _secondi_a_durata_stringa(info.get("pausa_extra_sec", 0)),
        ]
        for wd in range(7):
            v = info.get("per_day", {}).get(wd)
            if v == "OFF":
                riga.append("OFF")
            elif isinstance(v, datetime.time):
                riga.append(_time_a_stringa(v))
            else:
                riga.append("")
        ws.append(riga)

    wb.save(percorso_file)


def ottieni_orari_lavoro(file_bytes, file_name):
    """Orario di lavoro effettivo da usare nell'app: se esiste un salvataggio
    fatto dalla pagina 'Orari Dipendenti' (orari_lavoro.json) ha sempre la
    precedenza; altrimenti si parte da quanto scritto nel foglio Excel
    'OrariLavoro', se presente."""
    orari_salvati = carica_orari_da_json()
    if orari_salvati:
        return orari_salvati
    return carica_orari_da_excel(file_bytes, file_name)


def soglia_ritardo_per(orari_lavoro, nome, weekday, soglia_default):
    if orari_lavoro and nome in orari_lavoro:
        info = orari_lavoro[nome]
        override = info["per_day"].get(weekday)
        if override == "OFF":
            return None
        if isinstance(override, datetime.time):
            return override
        return info["default"] if info["default"] is not None else soglia_default
    return soglia_default


def giorno_lavorativo_per(orari_lavoro, nome, weekday):
    if orari_lavoro and nome in orari_lavoro:
        override = orari_lavoro[nome]["per_day"].get(weekday)
        if override == "OFF":
            return False
        if isinstance(override, datetime.time):
            return True
    return weekday < 5


def ore_previste_secondi(orari_lavoro, nome):
    if not orari_lavoro or nome not in orari_lavoro:
        return None
    info = orari_lavoro[nome]
    ingresso, uscita = info.get("default"), info.get("uscita")
    if ingresso is None or uscita is None:
        return None
    totale = _durata_a_secondi(uscita) - _durata_a_secondi(ingresso)
    if info.get("pausa_inizio") is not None and info.get("pausa_fine") is not None:
        pausa = _durata_a_secondi(info["pausa_fine"]) - _durata_a_secondi(info["pausa_inizio"])
        totale -= max(pausa, 0)
    totale -= info.get("pausa_extra_sec", 0.0)
    return max(totale, 0)


# ----------------------------------------------------------------------------
# Riconoscimento intelligente delle timbrature in base all'orario personale.
#
# Ogni giorno le timbrature della persona (in ordine di orario) vengono
# abbinate ai 4 momenti previsti dal suo orario (ingresso, inizio pausa, fine
# pausa, uscita) RISPETTANDO L'ORDINE: la prima timbratura non può essere
# un'uscita se dopo ce n'è un'altra, ecc. Tra tutti gli abbinamenti possibili
# si sceglie quello "più ragionevole", cioè quello con la minor somma di
# minuti di scostamento dagli orari previsti, dove:
#   - lasciare un momento previsto senza timbratura "costa" 120 minuti;
#   - lasciare una timbratura senza abbinamento ("non riconosciuta") "costa"
#     altri 120 minuti;
#   - un tipo di evento diverso da quello atteso (es. ENTRA registrato al
#     posto di ESCE) aggiunge 15 minuti.
# In pratica: una timbratura viene riconosciuta anche se è in anticipo o in
# ritardo di parecchio (fino a circa 4 ore), purché sia nella posizione giusta
# della giornata. Prima invece si accettavano solo timbrature entro 20 minuti
# dall'orario previsto: chi entrava alle 06:53 con ingresso previsto alle
# 07:30 (37 minuti prima) veniva segnato "non riconosciuto" e quella giornata
# perdeva le ore lavorate.
# ----------------------------------------------------------------------------

PENALITA_TIPO_DIVERSO_MIN = 15


def _abbina_timbrature_a_momenti(candidati, momenti_attesi):
    """candidati: lista di dict con chiavi 'secondi' ed 'evento'.
    momenti_attesi: lista di (nome_slot, secondi_attesi) in ordine cronologico.
    Restituisce {nome_slot: candidato} scegliendo l'abbinamento a costo minimo
    che rispetta l'ordine cronologico (programmazione dinamica: pochi elementi,
    quindi velocissimo)."""
    distanza_massima = float(imp("scostamento_massimo_min"))
    penalita_mancante = penalita_extra = distanza_massima / 2.0
    validi = sorted((c for c in candidati if c["secondi"] is not None), key=lambda c: c["secondi"])
    n, m = len(validi), len(momenti_attesi)
    INF = float("inf")
    costo = [[INF] * (m + 1) for _ in range(n + 1)]
    mossa = [[None] * (m + 1) for _ in range(n + 1)]
    costo[0][0] = 0.0
    for i in range(n + 1):
        for j in range(m + 1):
            c0 = costo[i][j]
            if c0 == INF:
                continue
            if i < n and j < m:
                nome_slot, atteso_sec = momenti_attesi[j]
                distanza = abs(validi[i]["secondi"] - atteso_sec) / 60.0
                if distanza < distanza_massima:
                    evento = str(validi[i]["evento"] or "").upper()
                    atteso_tipo = TIPO_ATTESO_SLOT[nome_slot]
                    penalita_tipo = (
                        PENALITA_TIPO_DIVERSO_MIN
                        if ("ENTRA" in evento or "ESCE" in evento) and atteso_tipo not in evento
                        else 0
                    )
                    nuovo = c0 + distanza + penalita_tipo
                    if nuovo < costo[i + 1][j + 1]:
                        costo[i + 1][j + 1] = nuovo
                        mossa[i + 1][j + 1] = "abbina"
            if i < n and c0 + penalita_extra < costo[i + 1][j]:
                costo[i + 1][j] = c0 + penalita_extra
                mossa[i + 1][j] = "extra"
            if j < m and c0 + penalita_mancante < costo[i][j + 1]:
                costo[i][j + 1] = c0 + penalita_mancante
                mossa[i][j + 1] = "manca"

    assegnazioni = {}
    i, j = n, m
    while i > 0 or j > 0:
        m_ij = mossa[i][j]
        if m_ij == "abbina":
            assegnazioni[momenti_attesi[j - 1][0]] = validi[i - 1]
            i, j = i - 1, j - 1
        elif m_ij == "extra":
            i -= 1
        else:
            j -= 1
    return assegnazioni

ETICHETTE_SLOT = {
    "ingresso": "Ingresso", "inizio_pausa": "Inizio pausa",
    "fine_pausa": "Fine pausa", "uscita": "Uscita",
}
TIPO_ATTESO_SLOT = {
    "ingresso": "ENTRA", "inizio_pausa": "ESCE",
    "fine_pausa": "ENTRA", "uscita": "ESCE",
}


def _rielabora_giorno_da_orario(riga, orari_lavoro, oggi):
    nome = riga["Nome"]
    eventi_puliti = riga.get("EventiPuliti")
    if not isinstance(eventi_puliti, list) or not eventi_puliti or nome not in orari_lavoro:
        return riga
    weekday = riga["Data"].weekday()
    if not giorno_lavorativo_per(orari_lavoro, nome, weekday):
        return riga

    info = orari_lavoro[nome]
    slot_definizioni = [
        ("ingresso", soglia_ritardo_per(orari_lavoro, nome, weekday, None), "Ingresso non timbrato"),
        ("inizio_pausa", info.get("pausa_inizio"), "Inizio pausa non timbrato"),
        ("fine_pausa", info.get("pausa_fine"), "Rientro pausa non timbrato"),
        ("uscita", info.get("uscita"), "Uscita non timbrata"),
    ]
    if all(atteso is None for _, atteso, _ in slot_definizioni):
        return riga

    candidati = []
    for idx, (orario_grezzo, evento) in enumerate(eventi_puliti):
        candidati.append({
            "idx": idx, "orario": orario_grezzo, "evento": evento,
            "secondi": _secondi_di_orario(orario_grezzo), "usato": False,
        })

    momenti_attesi = [
        (nome_slot, atteso.hour * 3600 + atteso.minute * 60 + atteso.second)
        for nome_slot, atteso, _messaggio in slot_definizioni
        if atteso is not None
    ]
    assegnazioni = _abbina_timbrature_a_momenti(candidati, momenti_attesi)
    for c in assegnazioni.values():
        c["usato"] = True

    entrata1 = assegnazioni.get("ingresso", {}).get("orario")
    uscita1 = assegnazioni.get("inizio_pausa", {}).get("orario")
    entrata2 = assegnazioni.get("fine_pausa", {}).get("orario")
    uscita2 = assegnazioni.get("uscita", {}).get("orario")

    ha_ingresso = "ingresso" in assegnazioni
    ha_inizio_pausa = "inizio_pausa" in assegnazioni
    ha_fine_pausa = "fine_pausa" in assegnazioni
    ha_uscita = "uscita" in assegnazioni
    giorno_concluso = ha_uscita or riga["Data"].date() < oggi

    anomalie = []
    if giorno_concluso:
        for nome_slot, atteso, messaggio in slot_definizioni:
            if atteso is None or nome_slot in assegnazioni:
                continue
            if nome_slot == "inizio_pausa":
                if ha_fine_pausa:
                    anomalie.append(messaggio)
            elif nome_slot == "fine_pausa":
                if ha_inizio_pausa:
                    anomalie.append(messaggio)
            elif assegnazioni:
                anomalie.append(messaggio)
        pausa_prevista = info.get("pausa_inizio") is not None and info.get("pausa_fine") is not None
        if pausa_prevista and ha_ingresso and ha_uscita and not ha_inizio_pausa and not ha_fine_pausa:
            anomalie.append("Pausa non timbrata")

    durata1_sec = durata2_sec = 0.0
    # Se manca una (o entrambe) le timbrature della pausa ma ci sono entrata e uscita,
    # si calcola dall'entrata all'uscita togliendo la pausa prevista: così un
    # "dimenticato di timbrare la pausa" non fa perdere mezza giornata di ore
    # (l'anomalia resta comunque segnalata).
    if entrata1 is not None and uscita2 is not None and not (ha_inizio_pausa and ha_fine_pausa):
        t_e, t_u = _secondi_di_orario(entrata1), _secondi_di_orario(uscita2)
        if t_e is not None and t_u is not None and t_u > t_e:
            durata1_sec = t_u - t_e
            p1, p2 = info.get("pausa_inizio"), info.get("pausa_fine")
            if imp("togli_pausa_se_non_timbrata") and p1 is not None and p2 is not None:
                durata1_sec = max(durata1_sec - max(_durata_a_secondi(p2) - _durata_a_secondi(p1), 0), 0)
        else:
            anomalie.append("Orari incoerenti")
    else:
        if entrata1 is not None and uscita1 is not None:
            t_e, t_u = _secondi_di_orario(entrata1), _secondi_di_orario(uscita1)
            if t_e is not None and t_u is not None and t_u > t_e:
                durata1_sec = t_u - t_e
            else:
                anomalie.append("Orari incoerenti (turno 1)")
        if entrata2 is not None and uscita2 is not None:
            t_e, t_u = _secondi_di_orario(entrata2), _secondi_di_orario(uscita2)
            if t_e is not None and t_u is not None and t_u > t_e:
                durata2_sec = t_u - t_e
            else:
                anomalie.append("Orari incoerenti (turno 2)")

    extra = [c for c in candidati if not c["usato"]]
    if extra:
        anomalie.append(f"{len(extra)} timbratura/e extra non riconosciuta/e")

    riga["Entrata1"], riga["Uscita1"] = entrata1, uscita1
    riga["Entrata2"], riga["Uscita2"] = entrata2, uscita2
    riga["Durata1"] = pd.Timedelta(seconds=durata1_sec) if durata1_sec > 0 else None
    riga["Durata2"] = pd.Timedelta(seconds=durata2_sec) if durata2_sec > 0 else None
    durata_tot = durata1_sec + durata2_sec
    riga["DurataTotale"] = pd.Timedelta(seconds=durata_tot) if durata_tot > 0 else None
    riga["Anomalia"] = "; ".join(anomalie) if anomalie else "OK"

    timbrature_grezze = riga.get("TimbratureGrezze")
    if isinstance(timbrature_grezze, list):
        mappa_idx_slot = {c["idx"]: nome_slot for nome_slot, c in assegnazioni.items()}
        nuove_annotazioni = []
        idx_pulito = 0
        for ann in timbrature_grezze:
            nuova = dict(ann)
            if not ann.get("scartata"):
                nome_slot = mappa_idx_slot.get(idx_pulito)
                if nome_slot:
                    nuova["slot"] = ETICHETTE_SLOT[nome_slot]
                    evento_originale = str(ann.get("evento_originale", "")).upper()
                    nuova["corretto"] = bool(evento_originale) and (TIPO_ATTESO_SLOT[nome_slot] not in evento_originale)
                else:
                    nuova["slot"] = None
                    nuova["corretto"] = False
                    nuova["extra"] = True
                idx_pulito += 1
            nuove_annotazioni.append(nuova)
        riga["TimbratureGrezze"] = nuove_annotazioni

    return riga


def assegna_turni_da_orario(df, orari_lavoro, oggi):
    """Rielabora Entrata1/Uscita1/Entrata2/Uscita2 (e le anomalie collegate)
    per i dipendenti che hanno un orario impostato in 'Orari Dipendenti',
    lasciando invariate le righe di chi non ce l'ha (per loro resta il
    vecchio calcolo a turni alternati)."""
    if not orari_lavoro or "EventiPuliti" not in df.columns:
        return df
    df = df.copy()
    df = df.apply(lambda riga: _rielabora_giorno_da_orario(riga, orari_lavoro, oggi), axis=1)
    return df


ETICHETTA_E_DIPENDENTE = "✅ Dipendente (nome giusto)"
COLONNE_EDITOR_ORARI = [
    "Nome", "CorrispondeA", "Ingresso", "InizioPausa", "FinePausa", "Uscita",
    "AltraPausaMin", "GiorniRiposo",
]


def orari_a_editor_df(orari, nomi_trovati, alias_nomi=None):
    """Tabella unica per la pagina 'Orari Dipendenti': una riga per ogni nome
    (dipendenti già impostati, nomi già corretti in passato, e qualunque nome
    comparso nelle timbrature). La seconda colonna dice cosa fare di quel
    nome: è un dipendente vero, corrisponde a un altro dipendente (nome
    scritto male), oppure va ignorato. Vuota = nome nuovo ancora da decidere."""
    alias_nomi = alias_nomi or {}
    orari = orari or {}
    chiavi_note = {n.strip().lower() for n in list(orari.keys()) + list(alias_nomi.keys())}
    nomi = list(orari.keys()) + [n for n in alias_nomi.keys() if n not in orari]
    for n in nomi_trovati or []:
        if n and n.strip().lower() not in chiavi_note:
            nomi.append(n)
            chiavi_note.add(n.strip().lower())

    righe = []
    for nome in nomi:
        info = orari.get(nome, {})
        if nome in orari:
            scelta = ETICHETTA_E_DIPENDENTE
        elif nome in alias_nomi:
            scelta = alias_nomi[nome]
        else:
            scelta = None
        giorni_off = [
            GIORNI_IT[wd] for wd, v in info.get("per_day", {}).items() if v == "OFF"
        ]
        righe.append({
            "Nome": nome,
            "CorrispondeA": scelta,
            "Ingresso": info.get("default"),
            "InizioPausa": info.get("pausa_inizio"),
            "FinePausa": info.get("pausa_fine"),
            "Uscita": info.get("uscita"),
            "AltraPausaMin": round((info.get("pausa_extra_sec") or 0) / 60),
            "GiorniRiposo": [g for g in GIORNI_IT if g in giorni_off],
        })
    # prima i nomi ancora da decidere (così saltano all'occhio), poi in ordine alfabetico
    righe.sort(key=lambda r: (r["CorrispondeA"] is not None, r["Nome"].lower()))
    return pd.DataFrame(righe, columns=COLONNE_EDITOR_ORARI)


def _testo_cella(v):
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    return str(v).strip()


def editor_df_a_orari(df_editor):
    """Legge la tabella modificata e restituisce (orari, correzioni_nomi):
    - righe con 'È un dipendente' (o con la seconda colonna vuota ma con un
      orario compilato, es. un dipendente aggiunto a mano) -> orari;
    - righe che puntano a un altro nome -> correzione del nome scritto male;
    - righe con 'Ignora' -> timbrature da scartare;
    - righe con seconda colonna vuota e nessun orario -> ancora da decidere
      (non salvate: quel nome resta escluso dai conteggi come prima)."""
    abbr_lower = {a.lower(): wd for a, wd in COLONNE_GIORNO_SETTIMANA.items()}
    orari, correzioni = {}, {}
    for _, riga in df_editor.iterrows():
        nome = _testo_cella(riga.get("Nome"))
        if not nome:
            continue
        scelta = _testo_cella(riga.get("CorrispondeA"))
        ha_orario = any(_testo_cella(riga.get(c)) for c in ("Ingresso", "InizioPausa", "FinePausa", "Uscita"))

        if scelta and scelta not in (ETICHETTA_E_DIPENDENTE, nome):
            correzioni[nome] = scelta
            continue
        if not scelta and not ha_orario:
            continue

        per_day = {}
        riposo = riga.get("GiorniRiposo")
        if isinstance(riposo, (list, tuple, set)) or hasattr(riposo, "tolist"):
            pezzi = list(riposo.tolist() if hasattr(riposo, "tolist") else riposo)
        else:
            pezzi = _testo_cella(riposo).split(",")
        for pezzo in pezzi:
            pezzo = str(pezzo).strip().lower()
            if pezzo in abbr_lower:
                per_day[abbr_lower[pezzo]] = "OFF"
        altra_pausa = riga.get("AltraPausaMin")
        try:
            altra_pausa = float(altra_pausa) if _testo_cella(altra_pausa) else 0.0
        except (TypeError, ValueError):
            altra_pausa = 0.0
        orari[nome] = {
            "default": _stringa_a_time(_testo_cella(riga.get("Ingresso"))),
            "uscita": _stringa_a_time(_testo_cella(riga.get("Uscita"))),
            "pausa_inizio": _stringa_a_time(_testo_cella(riga.get("InizioPausa"))),
            "pausa_fine": _stringa_a_time(_testo_cella(riga.get("FinePausa"))),
            "pausa_extra_sec": altra_pausa * 60,
            "per_day": per_day,
        }
    return orari, correzioni


def _elabora_df_completo(df, orari_lavoro, alias_nomi, soglia_default, oggi):
    """Tutti i calcoli sulle timbrature (abbinamento ai turni, ritardi, uscite
    anticipate, anomalie...). Racchiusi in una funzione per poterne mettere
    in memoria (cache) il risultato: si rifanno solo quando cambiano davvero
    i dati, gli orari o le correzioni dei nomi, non ad ogni clic."""
    df["Data"] = pd.to_datetime(df["Data"])

    # Corregge subito eventuali nomi scritti male (es. refusi di un dispositivo
    # di timbratura), PRIMA del filtro sui nomi presenti in "Orari Dipendenti"
    # qui sotto: così una timbratura con un nome scritto male non viene persa,
    # ma corretta e conteggiata per la persona giusta.
    df = applica_alias_nomi(df, alias_nomi)

    if orari_lavoro:
        mappa_nomi = {n.strip().lower(): n for n in orari_lavoro.keys()}
        df["_chiave_nome"] = df["Nome"].str.strip().str.lower()
        df = df[df["_chiave_nome"].isin(mappa_nomi.keys())].copy()
        df["Nome"] = df["_chiave_nome"].map(mappa_nomi)
        df = df.drop(columns="_chiave_nome")
        if df.empty:
            return df, "nessuna_corrispondenza"

    df = assegna_turni_da_orario(df, orari_lavoro, oggi)

    # Con dataset molto piccoli (es. poche timbrature) può capitare che
    # nessun dipendente/giorno abbia mai un secondo turno: in tal caso le
    # colonne Entrata2/Uscita2/Durata2 non vengono create affatto, e il
    # codice sotto che le legge esploderebbe con un KeyError. Le garantiamo
    # sempre presenti (valorizzate a None se assenti) per essere robusti
    # anche con log minimi o parziali.
    for colonna in ("Entrata1", "Uscita1", "Durata1", "Entrata2", "Uscita2", "Durata2"):
        if colonna not in df.columns:
            df[colonna] = None

    df["DurataTotale_sec"] = df["DurataTotale"].map(a_secondi)
    df["Entrata1_time"] = df["Entrata1"].map(a_time)
    df["GiornoSettimana"] = df["Data"].dt.weekday
    df["SogliaRitardo"] = df.apply(
        lambda r: soglia_ritardo_per(orari_lavoro, r["Nome"], r["GiornoSettimana"], soglia_default),
        axis=1,
    )
    df["InRitardo"] = df.apply(
        lambda r: r["SogliaRitardo"] is not None
        and r["Entrata1_time"] is not None
        # confronto al minuto (i secondi non contano: 07:30:45 non è un ritardo sulle 07:30)
        and r["Entrata1_time"].replace(second=0, microsecond=0)
        > _sposta_orario(r["SogliaRitardo"], imp("tolleranza_ritardo_ingresso_min")).replace(second=0, microsecond=0),
        axis=1,
    )

    def _ultima_uscita_time(r):
        u2 = a_time(r["Uscita2"])
        return u2 if u2 is not None else a_time(r["Uscita1"])

    def _uscita_anticipata(r):
        if not orari_lavoro or r["Nome"] not in orari_lavoro:
            return False
        atteso = orari_lavoro[r["Nome"]].get("uscita")
        reale = r["UltimaUscita_time"]
        if atteso is None or reale is None:
            return False
        return reale < _sposta_orario(atteso, -float(imp("tolleranza_uscita_anticipata_min")))

    df["UltimaUscita_time"] = df.apply(_ultima_uscita_time, axis=1)
    df["UscitaAnticipata"] = df.apply(_uscita_anticipata, axis=1)

    def _ritardo_fine_pausa(r):
        if not orari_lavoro or r["Nome"] not in orari_lavoro:
            return False
        atteso = orari_lavoro[r["Nome"]].get("pausa_fine")
        reale = a_time(r["Entrata2"])
        if atteso is None or reale is None:
            return False
        return reale.replace(second=0, microsecond=0) > _sposta_orario(atteso, float(imp("tolleranza_ritardo_fine_pausa_min")))

    df["RitardoFinePausa"] = df.apply(_ritardo_fine_pausa, axis=1)
    df["AnomaliaGrave"] = df["Anomalia"].apply(lambda a: a not in ("OK", None) and pd.notna(a))

    def _tipi_anomalia(r):
        tipi = []
        if r["InRitardo"]:
            tipi.append("Ritardo entrata")
        if r["RitardoFinePausa"]:
            tipi.append("Ritardo fine pausa")
        if r["UscitaAnticipata"]:
            tipi.append("Uscita anticipata")
        if r["AnomaliaGrave"]:
            tipi.append(str(r["Anomalia"]))
        return tipi

    df["TipiAnomalia"] = df.apply(_tipi_anomalia, axis=1)
    df["HaAnomalia"] = df["TipiAnomalia"].map(bool)
    return df, "ok"


def _firma_elenco(voci):
    """Impronta dell'elenco file (nomi + date + dimensioni): se non cambia
    nessun file, l'impronta è identica e si riusa il risultato già calcolato."""
    import hashlib
    return hashlib.md5(repr(voci).encode("utf-8")).hexdigest()


@st.cache_data(show_spinner="Elaborazione timbrature...", max_entries=4)
def prepara_dati(fonte, percorso_cartella, firma_cartella, _voci_cartella, file_bytes, file_name,
                 orari_lavoro, alias_nomi, soglia_default, oggi, impostazioni=None, correzioni=None):
    """Carica ed elabora tutti i dati. Il risultato resta in memoria finché
    non cambia qualcosa (nuovi file di timbratura, orari, correzioni nomi,
    soglia, data di oggi): così un clic sulla dashboard non rifà più tutto
    il lavoro da capo, ma riusa i calcoli già pronti."""
    if fonte == "cartella":
        df_eventi = carica_eventi_grezzi_cartella(percorso_cartella, _voci_cartella)
        if df_eventi.empty:
            return {"df": pd.DataFrame(), "esito": "vuoto", "nomi_grezzi": []}
        nomi_grezzi = sorted(df_eventi["Nome"].unique())
        df_eventi = applica_alias_nomi(df_eventi, alias_nomi)
        df_eventi = applica_correzioni_eventi(df_eventi, correzioni)
        df = elabora_log_grezzi(df_eventi, None)
        for col in COLONNE_STD:
            if col not in df.columns:
                df[col] = None
    else:
        df = load_all_data(file_bytes, file_name)
        nomi_grezzi = sorted(df["Nome"].astype(str).str.strip().unique()) if not df.empty else []

    if df.empty:
        return {"df": df, "esito": "vuoto", "nomi_grezzi": nomi_grezzi}

    df, esito = _elabora_df_completo(df, orari_lavoro, alias_nomi, soglia_default, oggi)
    # Segni della matita: quali orari sono stati inseriti/modificati a mano e
    # quali giorni hanno avuto una correzione (anche solo un'eliminazione).
    eliminati = {(e["data"], e["nome"]) for e in (correzioni or {}).get("eliminate", [])}

    def _orari_manuali(grezze):
        out = []
        for a in grezze if isinstance(grezze, list) else []:
            if "MANUALE" in str(a.get("evento_originale", "")).upper():
                t = a_time(a.get("orario"))
                if t is not None:
                    out.append(t.strftime("%H:%M"))
        return out

    if not df.empty:
        if "TimbratureGrezze" in df.columns:
            df["OrariManuali"] = df["TimbratureGrezze"].map(_orari_manuali)
        else:
            df["OrariManuali"] = [[] for _ in range(len(df))]
        df["CorrettoAMano"] = [
            bool(m) or (d.date().isoformat(), n) in eliminati
            for m, d, n in zip(df["OrariManuali"], df["Data"], df["Nome"])
        ]
    return {"df": df, "esito": esito, "nomi_grezzi": nomi_grezzi}


def _messaggio_festivi(testo):
    st.session_state["_messaggio_festivi"] = testo


def _cb_festivo(data_iso, rendi, nome=""):
    try:
        _aggiorna_festivo(data_iso, rendi, nome)
        d = datetime.date.fromisoformat(data_iso)
        _messaggio_festivi(f"✅ {d:%d/%m/%Y} {'segnato come festivo' if rendi else 'non è più festivo'}.")
    except Exception as e:
        _messaggio_festivi(f"⚠️ Impossibile salvare: {e}")


def _cb_nazionali(valore):
    dati = carica_impostazioni()
    dati["festivi_nazionali"] = bool(valore)
    salva_impostazioni(dati)


def disegna_impostazioni_festivi():
    st.markdown("---")
    st.markdown("#### 🎉 Giorni festivi")
    messaggio = st.session_state.pop("_messaggio_festivi", None)
    if messaggio:
        st.success(messaggio)
    st.caption(
        "Nei giorni festivi nessuno risulta assente. Le **feste nazionali** sono già calcolate in automatico "
        "ogni anno (anche Pasquetta, che cambia data, e dal 2026 il 4 ottobre San Francesco). Qui sotto "
        "aggiungi i giorni in più: **santo patrono, ponti, chiusure aziendali**. Puoi farlo anche dal "
        "🗓️ Calendario della pagina Oggi, quando la ✏️ Modalità correzione è attiva."
    )
    nazionali_attivi = bool(imp("festivi_nazionali"))
    st.checkbox("Considera automaticamente le feste nazionali italiane", value=nazionali_attivi,
                key="chk_festivi_nazionali", on_change=lambda: _cb_nazionali(st.session_state["chk_festivi_nazionali"]))

    anno_corrente = datetime.date.today().year
    col_n, col_e = st.columns(2)
    with col_n:
        st.markdown(f"**Feste nazionali {anno_corrente}**")
        esclusi = set(imp("festivi_esclusi") or [])
        for d, nome in sorted(festivi_nazionali(anno_corrente).items()):
            c1, c2 = st.columns([4, 1.4], vertical_alignment="center")
            attivo = nazionali_attivi and d.isoformat() not in esclusi
            c1.markdown(f"{'🎉' if attivo else '➖'} {GIORNI_IT[d.weekday()]} {d:%d/%m} – {nome}"
                        + ("" if attivo else " *(non considerata)*"))
            if nazionali_attivi:
                c2.button("Togli" if attivo else "Rimetti", key=f"naz_{d.isoformat()}", use_container_width=True,
                          on_click=_cb_festivo, args=(d.isoformat(), not attivo))
    with col_e:
        st.markdown("**Festivi aggiunti da te** (patrono, ponti, chiusure)")
        with st.form("form_nuovo_festivo", clear_on_submit=True):
            f1, f2 = st.columns([1, 1.5])
            data_nuova = f1.date_input("Giorno", value=datetime.date.today(), format="DD/MM/YYYY")
            nome_nuovo = f2.text_input("Descrizione", placeholder="es. Santo patrono")
            if st.form_submit_button("➕ Aggiungi festivo", type="primary"):
                _cb_festivo(data_nuova.isoformat(), True, nome_nuovo.strip() or "Festivo")
                st.rerun()
        extra = imp("festivi_extra") or {}
        if not extra:
            st.caption("Nessun festivo aggiunto.")
        for iso, nome in sorted(extra.items(), reverse=True):
            d = datetime.date.fromisoformat(iso)
            c1, c2 = st.columns([4, 1.4], vertical_alignment="center")
            c1.markdown(f"🎉 {GIORNI_IT[d.weekday()]} {d:%d/%m/%Y} – {nome}")
            c2.button("🗑️ Togli", key=f"extra_{iso}", use_container_width=True, on_click=_cb_festivo, args=(iso, False))


def disegna_impostazioni_regole():
    """Sezione di Impostazioni con tutte le regole modificabili."""
    st.markdown("---")
    st.markdown("#### ⏱️ Regole per riconoscere e valutare le timbrature")
    messaggio = st.session_state.pop("_messaggio_impostazioni", None)
    if messaggio:
        st.success(messaggio)

    with st.expander("ℹ️ Come funziona il riconoscimento delle timbrature (leggimi)"):
        st.markdown(
            """
Ogni dipendente ha **4 momenti previsti** nella giornata (dalla pagina *Orari Dipendenti*):
**ingresso → inizio pausa → fine pausa → uscita**. La dashboard prende le timbrature del
giorno in ordine di orario e le abbina a questi 4 momenti, **rispettando l'ordine**
(la prima timbratura del mattino non può diventare un'uscita, ecc.).

**Scostamento massimo** – quanto può essere lontana una timbratura dall'orario previsto
per essere ancora abbinata a quel momento.
* Prima era fisso a **20 minuti**: serviva ad abbinare una timbratura a un momento solo se
  si era "sicuri" (es. una timbratura alle 10:15 non veniva presa per un inizio pausa
  delle 13:00). Il difetto: chi entrava 37 minuti prima o usciva 42 minuti dopo veniva
  segnato "non riconosciuto" e **la giornata perdeva le ore**.
* Adesso quella sicurezza la dà soprattutto **l'ordine** delle timbrature, quindi si può
  essere molto più larghi (consigliato **240 minuti** = 4 ore). La timbratura alle 10:15
  dell'esempio resta comunque "non riconosciuta" perché fra 07:30 e 13:00 c'è già chi
  corrisponde meglio.
* Se vuoi il comportamento rigido di prima, metti **20**.

**Tolleranze** – decidono solo quando una timbratura già riconosciuta viene *segnalata*
(colore arancione/rosso e conteggio ritardi), non cambiano le ore lavorate:
* *Ritardo ingresso*: minuti di margine oltre l'orario di ingresso prima di segnare "Ritardo".
* *Ritardo fine pausa*: minuti di margine oltre la fine pausa prevista.
* *Uscita anticipata*: minuti di margine prima dell'orario di uscita.

**Mezza giornata (½)** – nella Ricerca Storica un giorno viene segnato ½ invece di P se le
ore lavorate sono meno della percentuale indicata delle ore previste da orario (es. 60%).

**Doppia timbratura** – se la stessa persona timbra due volte entro questi secondi, la
seconda è considerata un doppio clic per errore e non viene contata (resta visibile barrata
nel log).
"""
        )

    with st.form("form_impostazioni_regole"):
        c1, c2 = st.columns(2)
        with c1:
            scostamento = st.number_input(
                "Scostamento massimo per riconoscere una timbratura (minuti)",
                min_value=10, max_value=600, step=5, value=int(imp("scostamento_massimo_min")),
                help="Consigliato 240. Metti 20 per il comportamento rigido di prima.",
            )
            tol_ingresso = st.number_input(
                "Tolleranza ritardo ingresso (minuti)", min_value=0, max_value=120, step=1,
                value=int(imp("tolleranza_ritardo_ingresso_min")),
                help="0 = è ritardo anche un minuto dopo l'orario di ingresso.",
            )
            tol_pausa = st.number_input(
                "Tolleranza ritardo rientro dalla pausa (minuti)", min_value=0, max_value=120, step=1,
                value=int(imp("tolleranza_ritardo_fine_pausa_min")),
            )
            tol_uscita = st.number_input(
                "Tolleranza uscita anticipata (minuti)", min_value=0, max_value=120, step=1,
                value=int(imp("tolleranza_uscita_anticipata_min")),
            )
            soglia_mezza = st.number_input(
                "Mezza giornata (½): se le ore lavorate sono sotto questa % delle ore previste",
                min_value=1, max_value=100, step=5, value=int(imp("soglia_mezza_giornata_pct")),
                help="Esempio: 60 = chi deve fare 8 ore e ne fa meno di 4,8 ha la mezza giornata.",
            )
        with c2:
            doppia = st.number_input(
                "Doppia timbratura: ignora la seconda se entro (secondi)", min_value=0, max_value=600,
                step=10, value=int(imp("finestra_doppia_timbratura_sec")),
            )
            soglia = st.time_input(
                "Soglia ritardo per chi non ha un orario impostato",
                value=a_time(imp("soglia_ritardo_default")) or SOGLIA_RITARDO_DEFAULT,
                step=datetime.timedelta(minutes=5),
            )
            togli_pausa = st.checkbox(
                "Se la pausa non è timbrata (o solo metà), togli comunque dalle ore la pausa prevista",
                value=bool(imp("togli_pausa_se_non_timbrata")),
                help="Esempio: entra 07:30, esce 16:30 senza timbrare la pausa 13:00-13:30 -> 8h30 invece di 9h.",
            )
            chiavi_apertura = list(OPZIONI_GIORNO_APERTURA.keys())
            giorno_apertura = st.radio(
                "Giorno mostrato quando si apre la dashboard",
                chiavi_apertura,
                index=chiavi_apertura.index(imp("giorno_apertura")) if imp("giorno_apertura") in chiavi_apertura else 0,
                format_func=lambda k: OPZIONI_GIORNO_APERTURA[k],
            )
        col_s, col_d = st.columns(2)
        salva = col_s.form_submit_button("💾 Salva impostazioni", type="primary")
        predefinite = col_d.form_submit_button("↺ Ripristina valori consigliati")

    if salva or predefinite:
        attuali = carica_impostazioni()
        if predefinite:
            nuove = dict(IMPOSTAZIONI_PREDEFINITE)
            for k in ("festivi_nazionali", "festivi_extra", "festivi_esclusi"):
                nuove[k] = attuali.get(k, IMPOSTAZIONI_PREDEFINITE[k])
        else:
            nuove = dict(attuali)
            nuove.update({
                "scostamento_massimo_min": int(scostamento),
                "tolleranza_ritardo_ingresso_min": int(tol_ingresso),
                "tolleranza_ritardo_fine_pausa_min": int(tol_pausa),
                "tolleranza_uscita_anticipata_min": int(tol_uscita),
                "finestra_doppia_timbratura_sec": int(doppia),
                "soglia_ritardo_default": soglia.strftime("%H:%M"),
                "giorno_apertura": giorno_apertura,
                "soglia_mezza_giornata_pct": int(soglia_mezza),
                "togli_pausa_se_non_timbrata": bool(togli_pausa),
            })
        try:
            salva_impostazioni(nuove)
            st.session_state["_messaggio_impostazioni"] = "✅ Impostazioni salvate e applicate."
        except Exception as e:
            st.session_state["_messaggio_impostazioni"] = f"⚠️ Impossibile salvare le impostazioni: {e}"
        st.rerun()


def html_stato_giorno(codice, lavorativo=True, data=None):
    """Pallino + testo per i giorni senza timbrature (o con giustificativo)."""
    festa = nome_festivo(data) if data is not None else None
    if festa and codice not in GIUSTIFICATIVI:
        return '<span class="dot dot-azzurro"></span>🎉 Festivo' + ("" if festa == "Festivo" else f" – {festa}")
    if codice == "M":
        return '<span class="dot dot-viola"></span>Malattia'
    if codice == "F":
        return '<span class="dot dot-azzurro"></span>Ferie'
    if codice == "F½":
        return '<span class="dot dot-azzurro"></span>Ferie (mezza giornata)'
    if codice == "A":
        return '<span class="dot dot-rosso"></span>Assente'
    return '<span class="dot dot-azzurro"></span>Riposo' if not lavorativo else '<span class="dot dot-grigio"></span>—'


ETICHETTE_MANCANZE = [
    ("Ingresso non timbrato", "Manca entrata"),
    ("Inizio pausa non timbrato", "Manca uscita per pausa"),
    ("Rientro pausa non timbrato", "Manca rientro dalla pausa"),
    ("Uscita non timbrata", "Manca uscita"),
    ("Uscita mancante", "Manca uscita"),
    ("Pausa non timbrata", "Pausa non timbrata"),
    ("non riconosciuta", "Timbratura non riconosciuta"),
    ("Orari incoerenti", "Orari incoerenti"),
]


def badge_mancanze(r):
    """Etichette rosse per le timbrature mancanti o anomale del giorno."""
    testo = str(r.get("Anomalia") or "")
    if not testo or testo == "OK":
        return ""
    visti, out = set(), ""
    for chiave, etichetta in ETICHETTE_MANCANZE:
        if chiave in testo and etichetta not in visti:
            visti.add(etichetta)
            out += f'<span class="badge-uscita">⚠️ {etichetta}</span>'
    return out


def testo_mancanze(r):
    testo = str(r.get("Anomalia") or "")
    return ", ".join(dict.fromkeys(e for k, e in ETICHETTE_MANCANZE if k in testo)) if testo != "OK" else ""


def orario_con_matita(r, campo):
    """Orario HH:MM di un campo della riga, con ✏️ se è stato inserito/modificato a mano."""
    testo = formatta_orario(r.get(campo))
    manuali = r.get("OrariManuali")
    if testo != "-" and isinstance(manuali, (list, tuple)) and testo in manuali:
        return f"{testo} ✏️"
    return testo


def ore_con_matita(r):
    testo = formatta_orario(r.get("DurataTotale"))
    return f"{testo} ✏️" if (testo != "-" and bool(r.get("CorrettoAMano"))) else testo


def _senza_matita(val):
    return str(val).replace("✏️", "").strip()


def badge_codice(codice):
    if codice in GIUSTIFICATIVI:
        return f'<span class="badge-giust">{GIUSTIFICATIVI[codice]}</span>'
    if codice == "½":
        return '<span class="badge-giust">½ Mezza giornata</span>'
    return ""


def vai_a_correzioni(nome, data):
    st.session_state["corr_nome"] = nome
    st.session_state["corr_data"] = data
    st.session_state["menu_principale"] = ETICHETTE_MENU["correzioni"]


def vai_a_oggi(nome, data):
    st.session_state["giorno_rif"] = data
    st.session_state["dipendente_dettaglio"] = nome
    st.session_state["dettaglio_periodo"] = (data.year, data.month)
    st.session_state["menu_principale"] = ETICHETTE_MENU["oggi"]


def matrice_codici_presenza(df_mese, nomi, anno, mese, colonne_giorni, orari_lavoro, giustificativi, oggi):
    """Tabella nomi x giorni con i simboli P / ½ / A / M / F / -."""
    righe = {(r["Nome"], int(r["Giorno"])): r for _, r in df_mese.iterrows()}
    dati = {
        g: [codice_presenza(n, datetime.date(anno, mese, g), righe.get((n, g)), orari_lavoro, giustificativi, oggi)
            for n in nomi]
        for g in colonne_giorni
    }
    return pd.DataFrame(dati, index=list(nomi))


def _clic_tabella_ricerca(key, anno, mese):
    """Clic su una cella di una tabella della Ricerca Storica:
    - sul nome (o su una colonna che non è un giorno) -> mostra solo quel dipendente;
    - su un giorno -> apre la pagina Oggi su quel dipendente e quel giorno."""
    stato = st.session_state.get(key)
    try:
        selezione = stato["selection"] if isinstance(stato, dict) else stato.selection
        celle = selezione["cells"] if isinstance(selezione, dict) else selezione.cells
    except Exception:
        return
    if not celle:
        return
    cella = celle[0]
    riga, colonna = (cella["row"], cella["column"]) if isinstance(cella, dict) else (cella[0], cella[1])
    nomi = st.session_state.get(f"{key}__nomi") or []
    if not (0 <= int(riga) < len(nomi)):
        return
    nome = nomi[int(riga)]
    try:
        giorno = int(str(colonna))
    except ValueError:
        giorno = None
    if giorno is None:
        st.session_state["widget_ricerca_dipendenti"] = [nome]
        st.session_state["_ricerca_dipendenti_scelti"] = [nome]
    else:
        try:
            if st.session_state.get("modalita_correzione"):
                apri_correzione(nome, datetime.date(anno, mese, giorno))
            else:
                vai_a_oggi(nome, datetime.date(anno, mese, giorno))
        except ValueError:
            return
    # tabella "nuova" al prossimo giro: niente cella rimasta selezionata
    st.session_state["_contatore_clic_ricerca"] = st.session_state.get("_contatore_clic_ricerca", 0) + 1


def _mostra_tutti_ricerca():
    st.session_state["widget_ricerca_dipendenti"] = []
    st.session_state["_ricerca_dipendenti_scelti"] = []


def mostra_matrice_cliccabile(matrice, nome_tab, anno, mese, giorni_weekend, stile_riga=None,
                              stili=None, nomi_righe=None, titolo_prima_colonna="Nome"):
    """Mostra una tabella (righe = dipendenti o voci, colonne = giorni) in cui
    si può cliccare una cella: vedi _clic_tabella_ricerca."""
    contatore = st.session_state.get("_contatore_clic_ricerca", 0)
    key = f"tab_{nome_tab}_{contatore}"
    colonne = list(matrice.columns)
    etichette = [str(i) for i in matrice.index]
    tab = pd.DataFrame({titolo_prima_colonna: etichette})
    for c in colonne:
        tab[str(c)] = [("-" if (v is None or (isinstance(v, float) and pd.isna(v))) else v) for v in matrice[c].tolist()]
    st.session_state[f"{key}__nomi"] = list(nomi_righe) if nomi_righe is not None else etichette

    def _stile(row):
        pos = int(row.name)
        if stili is not None:
            corpo = [stili.iloc[pos, k] for k in range(len(colonne))]
        elif stile_riga is not None:
            valori = pd.Series([row[str(c)] for c in colonne], index=colonne, name=etichette[pos])
            corpo = list(stile_riga(valori))
        else:
            corpo = ["background-color: rgba(128,128,128,0.15)" if c in giorni_weekend else "" for c in colonne]
        return ["font-weight:700; color:#2f6fd1"] + corpo

    styler = tab.style.apply(_stile, axis=1)
    try:
        st.dataframe(
            styler, key=key, on_select=functools.partial(_clic_tabella_ricerca, key, anno, mese),
            selection_mode="single-cell", hide_index=True, use_container_width=True,
            column_config={titolo_prima_colonna: st.column_config.TextColumn(titolo_prima_colonna, pinned=True)},
        )
    except Exception:
        st.dataframe(styler, hide_index=True, use_container_width=True)


def disegna_scheda_dipendente(nome, df_mese, anno, mese, colonne_giorni, giorni_weekend,
                              orari_lavoro, giustificativi, oggi, mese_raw):
    """Vista 'isolata' della Ricerca Storica: un solo dipendente, con tutte le
    informazioni del mese in un'unica tabella (una riga per ogni voce)."""
    col_t, col_c, col_b = st.columns([3, 1.3, 1.3])
    col_t.markdown(f"### 👤 {nome}")
    col_c.button("✏️ Correzioni", key="scheda_vai_correzioni", use_container_width=True,
                 on_click=vai_a_correzioni, args=(nome, oggi - datetime.timedelta(days=1)))
    col_b.button("👥 Mostra tutti", key="scheda_mostra_tutti", use_container_width=True,
                 on_click=_mostra_tutti_ricerca, type="primary")
    for avviso in controlla_orari({nome: (orari_lavoro or {}).get(nome)} if (orari_lavoro or {}).get(nome) else {}):
        st.warning("⚠️ Orario impostato male in 'Orari Dipendenti': " + avviso)

    dati_nome = df_mese[df_mese["Nome"] == nome]
    righe = {int(r["Giorno"]): r for _, r in dati_nome.iterrows()}
    voci = ["Presenza", "Entrata", "Inizio pausa", "Fine pausa", "Uscita", "Ore lavorate", "Anomalie"]
    valori = {v: [] for v in voci}
    stili = {v: [] for v in voci}
    colori_codice = {
        "P": "background-color: rgba(0,170,0,0.20)", "½": "background-color: rgba(240,190,0,0.30)",
        "A": "background-color: rgba(200,0,0,0.15)", "M": "background-color: rgba(142,68,173,0.22)",
        "F": "background-color: rgba(52,152,219,0.25)", "F½": "background-color: rgba(52,152,219,0.25)",
        "-": "color: rgba(120,120,120,0.8)",
    }
    for g in colonne_giorni:
        data_g = datetime.date(anno, mese, g)
        r = righe.get(g)
        codice = codice_presenza(nome, data_g, r, orari_lavoro, giustificativi, oggi)
        base = "background-color: rgba(128,128,128,0.12)" if g in giorni_weekend else ""
        valori["Presenza"].append(codice)
        stili["Presenza"].append(colori_codice.get(codice, ""))
        if r is None:
            for v in voci[1:]:
                valori[v].append("-")
                stili[v].append(base)
            continue
        uscita = r["Uscita2"] if a_time(r["Entrata2"]) is not None else r["UltimaUscita_time"]
        valori["Entrata"].append(orario_con_matita(r, "Entrata1"))
        stili["Entrata"].append("background-color: rgba(230,150,0,0.30)" if r["InRitardo"] else base)
        valori["Inizio pausa"].append(orario_con_matita(r, "Uscita1"))
        stili["Inizio pausa"].append(base)
        valori["Fine pausa"].append(orario_con_matita(r, "Entrata2"))
        stili["Fine pausa"].append("background-color: rgba(230,150,0,0.30)" if r["RitardoFinePausa"] else base)
        campo_uscita = "Uscita2" if a_time(r["Entrata2"]) is not None else "UltimaUscita_time"
        valori["Uscita"].append(orario_con_matita(r, campo_uscita))
        stili["Uscita"].append("background-color: rgba(200,0,0,0.18)" if r["UscitaAnticipata"] else base)
        valori["Ore lavorate"].append(ore_con_matita(r))
        stili["Ore lavorate"].append(base)
        valori["Anomalie"].append(("⚠️ " + testo_mancanze(r)) if r["AnomaliaGrave"] else "")
        stili["Anomalie"].append("background-color: rgba(200,0,0,0.12)" if r["AnomaliaGrave"] else base)

    codici = valori["Presenza"]
    ore_mese = dati_nome["DurataTotale_sec"].sum() / 3600
    m = st.columns(7)
    m[0].metric("⏱️ Ore", f"{ore_mese:.1f} h")
    m[1].metric("✅ Presenze", codici.count("P") + codici.count("½"))
    m[2].metric("½ Mezze", codici.count("½"))
    m[3].metric("❌ Assenze", codici.count("A"))
    m[4].metric("🤒 Malattia", codici.count("M"))
    m[5].metric("🏖️ Ferie", formatta_giorni(codici.count("F") + 0.5 * codici.count("F½")))
    m[6].metric("⏳ Ritardi", int(dati_nome["InRitardo"].sum()))

    matrice = pd.DataFrame(valori, index=colonne_giorni).T
    matrice_stili = pd.DataFrame(stili, index=colonne_giorni).T
    st.caption("👆 Clicca un **giorno** per aprirlo nella pagina Oggi. 🟧 ritardo · 🟥 uscita anticipata.")
    mostra_matrice_cliccabile(
        matrice, "scheda", anno, mese, giorni_weekend, stili=matrice_stili,
        nomi_righe=[nome] * len(voci), titolo_prima_colonna="Voce",
    )
    st.download_button(
        "⬇️ Scarica CSV", matrice.to_csv().encode("utf-8-sig"),
        file_name=f"{nome.replace(' ', '_')}_{mese_raw}.csv", mime="text/csv",
    )


# ----------------------------------------------------------------------------
# Pagina "Correzioni"
# ----------------------------------------------------------------------------

def _salva_con_messaggio(correzioni, messaggio):
    try:
        salva_correzioni(correzioni)
        st.session_state["_messaggio_correzioni"] = ("success", messaggio)
    except Exception as e:
        st.session_state["_messaggio_correzioni"] = ("error", f"⚠️ Impossibile salvare la correzione: {e}")
    # i dati vanno ricalcolati: se siamo dentro la finestra di correzione serve un
    # aggiornamento completo della pagina (vedi dialogo_correzione)
    st.session_state["_serve_ricarica"] = True


def _togli_timbratura_da(corr, nome, data, orario, evento, manuale):
    data_iso = data.isoformat()
    if manuale:
        corr["aggiunte"] = [
            a for a in corr["aggiunte"]
            if not (a["data"] == data_iso and a["nome"] == nome and str(a["orario"]).strip() == str(orario).strip())
        ]
    else:
        corr["eliminate"].append({
            "data": data_iso, "nome": nome, "orario": str(orario).strip(),
            "evento": str(evento).strip(), "quando": datetime.datetime.now().strftime("%d/%m/%Y %H:%M"),
        })


def _aggiungi_timbratura_a(corr, nome, data, orario, tipo, nota=""):
    corr["aggiunte"].append({
        "data": data.isoformat(), "nome": nome, "orario": orario.strftime("%H:%M:00"),
        "evento": tipo, "nota": nota, "quando": datetime.datetime.now().strftime("%d/%m/%Y %H:%M"),
    })


def _elimina_timbratura(nome, data, orario, evento, manuale):
    corr = carica_correzioni()
    _togli_timbratura_da(corr, nome, data, orario, evento, manuale)
    _salva_con_messaggio(corr, f"🗑️ Timbratura delle {str(orario)[:5]} eliminata ({nome}, {data:%d/%m/%Y}).")


def _modifica_timbratura(nome, data, orario, evento, manuale, chiave_nuovo_orario):
    nuovo = st.session_state.get(chiave_nuovo_orario)
    if nuovo is None:
        return
    tipo = "ENTRA" if "ENTRA" in str(evento).upper() else "ESCE"
    corr = carica_correzioni()
    _togli_timbratura_da(corr, nome, data, orario, evento, manuale)
    _aggiungi_timbratura_a(corr, nome, data, nuovo, tipo, nota=f"modificata (era {str(orario)[:5]})")
    _salva_con_messaggio(corr, f"💾 Timbratura delle {str(orario)[:5]} cambiata in {nuovo:%H:%M} ({nome}).")


def _aggiungi_da_widget(nome, data, tipo, chiave_orario, etichetta):
    orario = st.session_state.get(chiave_orario)
    if orario is None:
        return
    corr = carica_correzioni()
    _aggiungi_timbratura_a(corr, nome, data, orario, tipo, nota=etichetta)
    _salva_con_messaggio(corr, f"➕ Aggiunta {etichetta.lower()} alle {orario:%H:%M} ({nome}, {data:%d/%m/%Y}).")


def _aggiungi_giornata_standard(nome, data, momenti):
    corr = carica_correzioni()
    for etichetta, tipo, orario in momenti:
        _aggiungi_timbratura_a(corr, nome, data, orario, tipo, nota=etichetta)
    _salva_con_messaggio(corr, f"➕ Aggiunta la giornata secondo l'orario previsto ({nome}, {data:%d/%m/%Y}).")


def _imposta_giustificativo(nome, data, codice):
    corr = carica_correzioni()
    if codice:
        corr["giustificativi"][chiave_giorno(data, nome)] = codice
    else:
        corr["giustificativi"].pop(chiave_giorno(data, nome), None)
    _salva_con_messaggio(corr, f"💾 {nome} {data:%d/%m/%Y}: {GIUSTIFICATIVI.get(codice, 'giorno normale')}.")


def _ripristina_timbratura(indice):
    corr = carica_correzioni()
    if 0 <= indice < len(corr["eliminate"]):
        e = corr["eliminate"].pop(indice)
        _salva_con_messaggio(corr, f"↩️ Timbratura delle {e['orario'][:5]} ripristinata ({e['nome']}).")


def _annulla_aggiunta(indice):
    corr = carica_correzioni()
    if 0 <= indice < len(corr["aggiunte"]):
        a = corr["aggiunte"].pop(indice)
        _salva_con_messaggio(corr, f"↩️ Timbratura aggiunta a mano delle {a['orario'][:5]} tolta ({a['nome']}).")


def _togli_giustificativo(chiave):
    corr = carica_correzioni()
    corr["giustificativi"].pop(chiave, None)
    _salva_con_messaggio(corr, "↩️ Giustificativo tolto.")


MOMENTI_GIORNATA = [
    # (slot interno, etichetta, tipo evento, campo orario previsto, campo della riga)
    ("Ingresso", "Entrata", "ENTRA", "default", "Entrata1"),
    ("Inizio pausa", "Uscita per la pausa", "ESCE", "pausa_inizio", "Uscita1"),
    ("Fine pausa", "Rientro dalla pausa", "ENTRA", "pausa_fine", "Entrata2"),
    ("Uscita", "Uscita", "ESCE", "uscita", "Uscita2"),
]


def disegna_editor_giorno(df, nome, data, orari_lavoro, prefisso):
    """Correzione di UN giorno di UN dipendente, pensata per essere semplice:
    una riga per ognuno dei 4 momenti della giornata (Entrata, Uscita per la
    pausa, Rientro, Uscita) con orario previsto, orario timbrato (o MANCA) e
    un solo pulsante per l'azione giusta. Sotto: eventuali timbrature in più,
    quelle eliminate, aggiunta libera e malattia/ferie."""
    corr = carica_correzioni()
    chiave = chiave_giorno(data, nome)
    riga_df = df[(df["Nome"] == nome) & (df["Data"].dt.date == data)]
    r = riga_df.iloc[0] if not riga_df.empty else None
    info = (orari_lavoro or {}).get(nome, {})
    lavorativo = lavorativo_il(orari_lavoro, nome, data)
    stato_giust = corr["giustificativi"].get(chiave)
    grezze = r["TimbratureGrezze"] if (r is not None and isinstance(r.get("TimbratureGrezze"), list)) else []

    # ---- riepilogo del giorno ----
    ore = formatta_orario(r["DurataTotale"]) if r is not None else "-"
    riepilogo = f"⏱️ Ore calcolate: **{ore}**"
    if stato_giust:
        riepilogo += f" · Giorno segnato come **{GIUSTIFICATIVI[stato_giust]}**"
    elif nome_festivo(data):
        riepilogo += f" · 🎉 Giorno festivo"
    elif not lavorativo:
        riepilogo += " · Giorno di riposo da orario"
    st.markdown(riepilogo)
    if r is not None and badge_mancanze(r):
        st.markdown("Da sistemare: " + badge_mancanze(r), unsafe_allow_html=True)
    elif r is not None:
        st.markdown('<span style="color:#1e8449;font-weight:700">✅ Giornata completa, niente da sistemare</span>',
                    unsafe_allow_html=True)

    # ---- i 4 momenti della giornata ----
    per_slot = {}
    for ann in grezze:
        if ann.get("slot") and not ann.get("scartata") and not ann.get("extra"):
            per_slot.setdefault(ann["slot"], ann)
    momenti = [m for m in MOMENTI_GIORNATA if info.get(m[3]) is not None or m[0] in per_slot]

    if momenti:
        st.markdown("##### 🕒 La giornata")
        st.caption("Dove c'è **❌ MANCA** l'orario previsto è già scritto: correggilo se serve e premi **➕ Inserisci**. "
                   "Per cambiare un orario sbagliato modificalo e premi **💾 Salva**.")
        intest = st.columns([1.6, 0.8, 1.2, 1.5, 1.1, 0.9])
        for c, t in zip(intest, ["Momento", "Previsto", "Orario", "Stato", "", ""]):
            c.markdown(f"<span style='color:#6b7d8f;font-size:0.85em;font-weight:700'>{t}</span>", unsafe_allow_html=True)
        for slot, etichetta, tipo, campo_prev, _campo_riga in momenti:
            previsto = info.get(campo_prev)
            ann = per_slot.get(slot)
            c1, c2, c3, c4, c5, c6 = st.columns([1.6, 0.8, 1.2, 1.5, 1.1, 0.9], vertical_alignment="center")
            c1.markdown(f"{'➡️' if tipo == 'ENTRA' else '⬅️'} **{etichetta}**")
            c2.markdown(f"<span style='color:#6b7d8f'>{previsto:%H:%M}</span>" if previsto else "-", unsafe_allow_html=True)
            if ann is not None:
                evento = str(ann.get("evento_originale", ""))
                manuale = "MANUALE" in evento.upper()
                k = f"{prefisso}_{chiave}_{slot}_{ann.get('orario')}"
                c3.time_input(etichetta, value=a_time(ann.get("orario")) or previsto or datetime.time(8, 0),
                              step=60, key=k, label_visibility="collapsed")
                c4.markdown("✏️ inserita a mano" if manuale else "✅ timbrata")
                c5.button("💾 Salva", key=f"{k}_salva", use_container_width=True, help="Salva l'orario modificato",
                          on_click=_modifica_timbratura, args=(nome, data, ann.get("orario"), evento, manuale, k))
                c6.button("🗑️", key=f"{k}_elimina", use_container_width=True, help="Elimina questa timbratura",
                          on_click=_elimina_timbratura, args=(nome, data, ann.get("orario"), evento, manuale))
            else:
                k = f"{prefisso}_{chiave}_{slot}_manca_{previsto}"
                c3.time_input(etichetta, value=previsto or datetime.time(8, 0), step=60, key=k,
                              label_visibility="collapsed")
                c4.markdown("<span style='color:#c0392b;font-weight:800'>❌ MANCA</span>", unsafe_allow_html=True)
                c5.button("➕ Inserisci", key=f"{k}_ins", use_container_width=True, type="primary",
                          on_click=_aggiungi_da_widget, args=(nome, data, tipo, k, etichetta))
        mancanti = [(m[1], m[2], info.get(m[3])) for m in momenti if m[0] not in per_slot and info.get(m[3])]
        if len(mancanti) > 1:
            st.button(f"➕ Inserisci tutti i {len(mancanti)} momenti mancanti con l'orario previsto",
                      key=f"{prefisso}_tutti_{chiave}", on_click=_aggiungi_giornata_standard,
                      args=(nome, data, mancanti))

    # ---- timbrature in più (doppie / non riconosciute) ----
    altre = [a for a in grezze if a.get("scartata") or a.get("extra") or not a.get("slot")]
    eliminate_giorno = [(i, e) for i, e in enumerate(corr["eliminate"])
                        if e["data"] == data.isoformat() and e["nome"] == nome]
    if altre or eliminate_giorno:
        st.markdown("##### 🔎 Altre timbrature di questo giorno")
    for i, ann in enumerate(altre):
        evento = str(ann.get("evento_originale", ""))
        manuale = "MANUALE" in evento.upper()
        t = a_time(ann.get("orario"))
        motivo = "❌ doppia (non conteggiata)" if ann.get("scartata") else "⚠️ non riconosciuta"
        c1, c2, c3, c4 = st.columns([1.2, 1.4, 2.2, 1.2], vertical_alignment="center")
        c1.markdown(f"**{t:%H:%M}**" if t else str(ann.get("orario")))
        c2.markdown("➡️ Entrata" if "ENTRA" in evento.upper() else "⬅️ Uscita")
        c3.markdown(motivo)
        c4.button("🗑️ Elimina", key=f"{prefisso}_altra_{chiave}_{i}_{ann.get('orario')}", use_container_width=True,
                  on_click=_elimina_timbratura, args=(nome, data, ann.get("orario"), evento, manuale))
    for i, e in eliminate_giorno:
        c1, c2, c3, c4 = st.columns([1.2, 1.4, 2.2, 1.2], vertical_alignment="center")
        c1.markdown(f"~~{e['orario'][:5]}~~")
        c2.markdown(f"~~{'Entrata' if 'ENTRA' in e['evento'].upper() else 'Uscita'}~~")
        c3.markdown("🗑️ eliminata da te")
        c4.button("↩️ Ripristina", key=f"{prefisso}_ripristina_{chiave}_{i}", use_container_width=True,
                  on_click=_ripristina_timbratura, args=(i,))

    # ---- aggiunta libera ----
    with st.expander("➕ Aggiungi un'altra timbratura con orario a scelta"):
        k_lib, k_tipo = f"{prefisso}_libero_{chiave}", f"{prefisso}_libero_tipo_{chiave}"
        c1, c2, c3 = st.columns([1.1, 1.6, 1.4], vertical_alignment="center")
        c1.time_input("Orario", value=datetime.time(8, 0), step=60, key=k_lib, label_visibility="collapsed")
        tipo_lib = c2.radio("Tipo", ["ENTRA", "ESCE"], horizontal=True, key=k_tipo, label_visibility="collapsed",
                            format_func=lambda v: "➡️ Entrata" if v == "ENTRA" else "⬅️ Uscita")
        c3.button("➕ Aggiungi", key=f"{prefisso}_libero_btn_{chiave}", use_container_width=True,
                  on_click=_aggiungi_da_widget, args=(nome, data, tipo_lib, k_lib, "Timbratura aggiunta"))

    # ---- malattia / ferie ----
    st.markdown("##### 🤒 Malattia / 🏖️ Ferie")
    g1, g2, g3 = st.columns(3)
    g1.button("Giorno normale", key=f"{prefisso}_giust_n_{chiave}", use_container_width=True,
              type="primary" if not stato_giust else "secondary",
              on_click=_imposta_giustificativo, args=(nome, data, ""))
    g2.button("🤒 Malattia", key=f"{prefisso}_giust_m_{chiave}", use_container_width=True,
              type="primary" if stato_giust == "M" else "secondary",
              on_click=_imposta_giustificativo, args=(nome, data, "M"))
    g3.button("🏖️ Ferie", key=f"{prefisso}_giust_f_{chiave}", use_container_width=True,
              type="primary" if stato_giust == "F" else "secondary",
              on_click=_imposta_giustificativo, args=(nome, data, "F"))


def _chiudi_dialogo_correzione():
    st.session_state.pop("_dialogo_correzione", None)


@st.dialog("✏️ Correggi timbrature", width="large", on_dismiss=_chiudi_dialogo_correzione)
def dialogo_correzione(df, orari_lavoro, dipendenti=None):
    if st.session_state.pop("_serve_ricarica", False):
        st.rerun()   # ricalcola tutto con la correzione appena salvata (la finestra resta aperta)
    nome, data = st.session_state["_dialogo_correzione"]
    dipendenti = list(dipendenti or [nome])
    if nome not in dipendenti:
        dipendenti.append(nome)
    i = dipendenti.index(nome)
    precedente, successivo = dipendenti[(i - 1) % len(dipendenti)], dipendenti[(i + 1) % len(dipendenti)]

    # riga dipendente: ◀ precedente | nome | successivo ▶
    c1, c2, c3 = st.columns([1.3, 2, 1.3], vertical_alignment="center")
    c1.button(f"◀ {precedente}", key="dlg_dip_prec", use_container_width=True, on_click=_imposta_stato,
              kwargs={"_dialogo_correzione": (precedente, data)}, help="Dipendente precedente, stesso giorno")
    c2.markdown(f"<div style='text-align:center;font-size:1.25em;font-weight:800'>👤 {nome}</div>",
                unsafe_allow_html=True)
    c3.button(f"{successivo} ▶", key="dlg_dip_succ", use_container_width=True, on_click=_imposta_stato,
              kwargs={"_dialogo_correzione": (successivo, data)}, help="Dipendente successivo, stesso giorno")
    # riga giorno: ◀ giorno prima | data | giorno dopo ▶
    d1, d2, d3 = st.columns([1.3, 2, 1.3], vertical_alignment="center")
    d1.button("◀ Giorno prima", key="dlg_giorno_prec", use_container_width=True, on_click=_imposta_stato,
              kwargs={"_dialogo_correzione": (nome, data - datetime.timedelta(days=1))})
    festa = nome_festivo(data)
    d2.markdown(f"<div style='text-align:center;font-size:1.1em;font-weight:700'>📅 {GIORNI_IT_COMPLETI[data.weekday()]} "
                f"{data:%d/%m/%Y}{' · 🎉 ' + festa if festa else ''}</div>", unsafe_allow_html=True)
    d3.button("Giorno dopo ▶", key="dlg_giorno_succ", use_container_width=True, on_click=_imposta_stato,
              kwargs={"_dialogo_correzione": (nome, data + datetime.timedelta(days=1))})

    messaggio = st.session_state.pop("_messaggio_correzioni", None)
    if messaggio:
        getattr(st, messaggio[0])(messaggio[1])
    st.divider()
    disegna_editor_giorno(df, nome, data, orari_lavoro, prefisso="dlg")
    st.divider()
    if st.button("✅ Chiudi la finestra", use_container_width=True, type="primary", key="dlg_chiudi"):
        st.session_state.pop("_dialogo_correzione", None)
        st.rerun()


def apri_correzione(nome, data):
    """Usata dai clic in modalità correzione: apre la finestra di correzione."""
    st.session_state["_dialogo_correzione"] = (nome, data)


def disegna_pagina_correzioni(df, dipendenti, oggi, orari_lavoro):
    st.markdown("### ✏️ Correzioni timbrature, malattia e ferie")
    if not st.session_state.get("orari_sbloccato"):
        st.caption("🔒 Questa pagina è protetta dalla stessa password di 'Orari Dipendenti'.")
        with st.form("form_password_correzioni", clear_on_submit=True):
            password_inserita = st.text_input("Password", type="password")
            sblocca = st.form_submit_button("🔓 Sblocca")
        if sblocca:
            if password_inserita == carica_password():
                st.session_state["orari_sbloccato"] = True
                st.rerun()
            else:
                st.error("Password errata.")
        return
    if not dipendenti:
        st.info("Nessun dipendente disponibile.")
        return

    st.session_state.pop("_serve_ricarica", None)
    messaggio = st.session_state.pop("_messaggio_correzioni", None)
    if messaggio:
        getattr(st, messaggio[0])(messaggio[1])
    st.caption(
        "Le timbrature originali sul NAS **non vengono mai modificate**: le correzioni vengono salvate a parte "
        "(file `correzioni_timbrature.json` nella cartella della dashboard, uguale per tutti i computer) e "
        "applicate sopra ai dati. Ogni correzione si può annullare. Suggerimento: attiva la "
        "**✏️ Modalità correzione** in alto a destra per correggere direttamente cliccando nelle altre pagine."
    )

    corr = carica_correzioni()
    nome_ricordato = st.session_state.get("corr_nome")
    data_ricordata = st.session_state.get("corr_data") or (oggi - datetime.timedelta(days=1))
    c1, c2, c3 = st.columns([2, 1.2, 1.2])
    nome = c1.selectbox(
        "👤 Dipendente", dipendenti,
        index=dipendenti.index(nome_ricordato) if nome_ricordato in dipendenti else 0,
    )
    data = c2.date_input("📅 Giorno", value=min(data_ricordata, oggi), max_value=oggi, format="DD/MM/YYYY")
    st.session_state["corr_nome"], st.session_state["corr_data"] = nome, data
    with c3:
        st.markdown("<div style='height:1.75rem'></div>", unsafe_allow_html=True)
        st.button("👀 Vedi nella pagina Oggi", use_container_width=True, on_click=vai_a_oggi, args=(nome, data))

    with st.container(border=True):
        st.markdown(f"#### {nome} — {GIORNI_IT_COMPLETI[data.weekday()].lower()} {data:%d/%m/%Y}")
        disegna_editor_giorno(df, nome, data, orari_lavoro, prefisso="pag")

    with st.expander("🏖️ Malattia o ferie per più giorni di seguito"):
        with st.form("form_giustificativo_periodo"):
            f1, f2, f3 = st.columns(3)
            scelta = f1.radio("Tipo", ["M", "F", ""], horizontal=False,
                              format_func=lambda v: {"M": "🤒 Malattia", "F": "🏖️ Ferie", "": "Togli (giorni normali)"}[v])
            dal = f2.date_input("Dal", value=data, format="DD/MM/YYYY")
            al = f3.date_input("Al (compreso)", value=data, format="DD/MM/YYYY")
            salva_p = st.form_submit_button("💾 Salva periodo", type="primary")
        if salva_p:
            if al < dal:
                st.error("La data finale è prima di quella iniziale.")
            else:
                corr = carica_correzioni()
                giorni = [dal + datetime.timedelta(days=k) for k in range((al - dal).days + 1)]
                giorni = [g for g in giorni if lavorativo_il(orari_lavoro, nome, g)]
                for g in giorni:
                    if scelta:
                        corr["giustificativi"][chiave_giorno(g, nome)] = scelta
                    else:
                        corr["giustificativi"].pop(chiave_giorno(g, nome), None)
                _salva_con_messaggio(corr, f"💾 {nome}: {GIUSTIFICATIVI.get(scelta, 'giorni normali')} "
                                           f"per {len(giorni)} giorno/i lavorativi.")
                st.rerun()

    with st.expander("📋 Tutte le correzioni fatte (clicca per vederle e annullarle)"):
        tutte = []
        for i, a in enumerate(corr["aggiunte"]):
            tutte.append((a["data"], "➕ Aggiunta", a["nome"], f"{('Entrata' if a['evento'] == 'ENTRA' else 'Uscita')} "
                          f"{a['orario'][:5]}" + (f" – {a['nota']}" if a.get("nota") else ""), ("agg", i)))
        for i, e in enumerate(corr["eliminate"]):
            tutte.append((e["data"], "🗑️ Eliminata", e["nome"], f"{e['orario'][:5]}", ("eli", i)))
        for k, v in corr["giustificativi"].items():
            d, n = k.split("|", 1)
            tutte.append((d, "🤒 Malattia" if v == "M" else "🏖️ Ferie", n, "", ("giu", k)))
        if not tutte:
            st.caption("Nessuna correzione salvata finora.")
        for d, tipo, n, dettaglio, rif in sorted(tutte, key=lambda x: x[0], reverse=True)[:300]:
            k1, k2, k3, k4, k5 = st.columns([1, 1.2, 2, 2.2, 1], vertical_alignment="center")
            k1.markdown(datetime.date.fromisoformat(d).strftime("%d/%m/%Y"))
            k2.markdown(tipo)
            k3.markdown(n)
            k4.markdown(dettaglio)
            if rif[0] == "agg":
                k5.button("↩️ Annulla", key=f"ann_agg_{rif[1]}", on_click=_annulla_aggiunta, args=(rif[1],))
            elif rif[0] == "eli":
                k5.button("↩️ Annulla", key=f"ann_eli_{rif[1]}", on_click=_ripristina_timbratura, args=(rif[1],))
            else:
                k5.button("↩️ Annulla", key=f"ann_giu_{rif[1]}", on_click=_togli_giustificativo, args=(rif[1],))


# ----------------------------------------------------------------------------
# Scheda "🏖️ Ferie": richieste dei dipendenti, saldi e PIN, riepilogo, email
# ----------------------------------------------------------------------------

def _messaggio_ferie(tipo, testo):
    st.session_state["_messaggio_ferie"] = (tipo, testo)


def _cb_approva_ferie(id_richiesta, orari_lavoro):
    try:
        ok, msg = approva_richiesta_ferie(id_richiesta, "ufficio", orari_lavoro)
        if ok:
            msg += notifica_esito_ferie(id_richiesta)
        _messaggio_ferie("success" if ok else "warning", msg)
    except Exception as e:
        _messaggio_ferie("error", f"⚠️ Impossibile approvare: {e}")


def _cb_rifiuta_ferie(id_richiesta):
    try:
        ok, msg = rifiuta_richiesta_ferie(id_richiesta, st.session_state.get(f"motivo_ferie_{id_richiesta}", ""), "ufficio")
        if ok:
            msg += notifica_esito_ferie(id_richiesta)
        _messaggio_ferie("success" if ok else "warning", msg)
    except Exception as e:
        _messaggio_ferie("error", f"⚠️ Impossibile rifiutare: {e}")


def _aggiorna_voce_ferie(nome, **campi):
    """Cambia solo alcuni campi di un dipendente in ferie_saldi.json (rilegge il file prima)."""
    saldi = carica_saldi_ferie()
    voce = dict(saldi.get(nome) or {})
    for k, v in campi.items():
        if v is None:
            voce.pop(k, None)
        else:
            voce[k] = v
    saldi[nome] = voce
    salva_saldi_ferie(saldi)


def _riga_richiesta_ferie(r):
    dal, al = datetime.date.fromisoformat(r["dal"]), datetime.date.fromisoformat(r["al"])
    testo = f"{GIORNI_IT[dal.weekday()]} {dal:%d/%m/%Y}"
    if al != dal:
        testo += f" → {GIORNI_IT[al.weekday()]} {al:%d/%m/%Y}"
    elif r.get("mezza_giornata"):
        testo += f" (mezza giornata, {r['mezza_giornata']})"
    return testo


def _disegna_ferie_richieste(orari_lavoro, oggi):
    richieste = carica_richieste_ferie()
    in_attesa = sorted([r for r in richieste if r["stato"] == "in_attesa"], key=lambda r: r["dal"])
    saldi = carica_saldi_ferie()
    giustificativi = carica_correzioni()["giustificativi"]
    if not in_attesa:
        st.success("Nessuna richiesta in attesa. 🎉")
    for r in in_attesa:
        giorni, peso = giorni_richiesta_ferie(orari_lavoro, r)
        totale = len(giorni) * peso
        with st.container(border=True):
            st.markdown(f"**{r['nome']}** — {_riga_richiesta_ferie(r)} · **{formatta_giorni(totale)} giorni lavorativi**")
            if r.get("nota"):
                st.caption(f"📝 {r['nota']}")
            saldo = saldo_ferie(r["nome"], datetime.date.fromisoformat(r["dal"]).year, orari_lavoro,
                                giustificativi, richieste, saldi, oggi)
            if saldo["configurato"]:
                resto = saldo["disponibili_se_approvate"]
                st.caption(f"Disponibili {formatta_giorni(saldo['disponibili'])} · dopo questa e le altre in attesa: "
                           f"{formatta_giorni(resto)}")
                if resto < 0:
                    st.warning("⚠️ Supera i giorni disponibili del dipendente.")
            else:
                st.warning("⚠️ I giorni di ferie di questo dipendente per quell'anno non sono impostati (scheda 👥 Saldi e PIN).")
            for g, codice in conflitti_richiesta_ferie(orari_lavoro, r, giustificativi):
                st.warning(f"⚠️ Il {g:%d/%m/%Y} è già segnato come {GIUSTIFICATIVI.get(codice, codice)}"
                           + (" (non verrà sovrascritto)." if codice == "M" else " (verrà riscritto)."))
            if datetime.date.fromisoformat(r["dal"]) < oggi:
                st.caption("ℹ️ Il periodo è già iniziato.")
            c1, c2, c3 = st.columns([1, 2, 1], vertical_alignment="bottom")
            c1.button("✅ Approva", key=f"appr_ferie_{r['id']}", type="primary", use_container_width=True,
                      on_click=_cb_approva_ferie, args=(r["id"], orari_lavoro))
            c2.text_input("Motivo del rifiuto", key=f"motivo_ferie_{r['id']}", placeholder="Es. periodo di picco di lavoro")
            c3.button("❌ Rifiuta", key=f"rif_ferie_{r['id']}", use_container_width=True,
                      on_click=_cb_rifiuta_ferie, args=(r["id"],))

    with st.expander("📜 Storico delle ultime richieste"):
        decise = [r for r in richieste if r["stato"] != "in_attesa"][:40]
        if not decise:
            st.caption("Ancora nessuna richiesta decisa.")
        else:
            st.dataframe(pd.DataFrame([{
                "Dipendente": r["nome"], "Periodo": _riga_richiesta_ferie(r),
                "Giorni": formatta_giorni(r.get("giorni_lavorativi") or 0),
                "Stato": STATI_RICHIESTA.get(r["stato"], r["stato"]),
                "Deciso il": (r.get("deciso_il") or "")[:16].replace("T", " "),
                "Motivo": r.get("motivo_rifiuto") or "",
            } for r in decise]), use_container_width=True, hide_index=True)
        st.caption("Per togliere delle ferie già approvate usa la pagina ✏️ Correzioni (ogni giorno si può annullare).")


def _disegna_ferie_saldi(orari_lavoro, oggi):
    nomi = sorted(orari_lavoro.keys())
    saldi = carica_saldi_ferie()
    cfg = carica_config_ferie()
    nome = st.selectbox("👤 Dipendente", nomi, key="ferie_nome_saldi")
    voce = saldi.get(nome) or {}
    with st.form(f"form_saldo_ferie_{nome}"):
        c1, c2, c3 = st.columns(3)
        anno = c1.number_input("Anno", min_value=2020, max_value=2100, step=1, value=int(voce.get("anno") or oggi.year))
        spettanti = c2.number_input("Giorni spettanti nell'anno", min_value=0.0, max_value=200.0, step=0.5,
                                    value=float(voce.get("giorni_spettanti") or 0))
        residuo = c3.number_input("Residuo anno precedente", min_value=-50.0, max_value=200.0, step=0.5,
                                  value=float(voce.get("residuo_anno_precedente") or 0))
        indirizzo = st.text_input("Email (facoltativa, per avvisarlo dell'esito)", value=voce.get("email") or "")
        salva = st.form_submit_button("💾 Salva giorni ed email", type="primary")
    if salva:
        _aggiorna_voce_ferie(nome, anno=int(anno), giorni_spettanti=float(spettanti),
                             residuo_anno_precedente=float(residuo), email=indirizzo.strip())
        st.success("Salvato.")
        voce = carica_saldi_ferie().get(nome) or {}

    st.markdown("##### 🔑 PIN personale")
    st.caption("Serve al dipendente per entrare nell'app (4-8 cifre). Viene salvato solo come impronta: non si può rileggere, "
               "solo cambiare. Cambiare il PIN fa uscire il dipendente dai telefoni dove era già entrato.")
    messaggio_pin = st.session_state.pop("_messaggio_pin_ferie", None)
    if messaggio_pin:
        st.success(messaggio_pin)
    st.write("PIN impostato: " + ("✅ sì" if voce.get("pin_hash") else "❌ no"))
    p1, p2 = st.columns([2, 1], vertical_alignment="bottom")
    nuovo_pin = p1.text_input("Nuovo PIN", type="password", max_chars=8, key=f"ferie_pin_{nome}")
    if p2.button("🔑 Imposta PIN", use_container_width=True, key=f"ferie_imposta_pin_{nome}"):
        if not pin_valido(nuovo_pin):
            st.error("Il PIN deve avere da 4 a 8 cifre (solo numeri).")
        else:
            _aggiorna_voce_ferie(nome, pin_hash=crea_pin_hash(nuovo_pin))
            st.session_state["_messaggio_pin_ferie"] = f"PIN di {nome} impostato."
            st.rerun()
    if st.button("🎲 Genera un PIN casuale", key=f"ferie_pin_casuale_{nome}"):
        pin = nuovo_pin_casuale()
        _aggiorna_voce_ferie(nome, pin_hash=crea_pin_hash(pin))
        st.session_state["_messaggio_pin_ferie"] = f"Nuovo PIN di {nome}: {pin} — comunicaglielo ora, poi non sarà più visibile."
        st.rerun()

    st.markdown("##### 💳 Carta NFC")
    st.caption("La carta funziona come un link: leggendola con il telefono si apre l'app già collegata al dipendente. "
               "Senza codice basta il nome scritto sulla carta; con un codice personale la carta non si può copiare "
               "conoscendo solo il nome.")
    codice = voce.get("carta_codice")
    base = (cfg.get("url_app") or "https://INDIRIZZO-DEL-NAS/ziwood/ferie/index.html").rstrip("?")
    link = f"{base}?carta={urllib.parse.quote(nome)}" + (f"&c={codice}" if codice else "")
    st.code(link, language=None)
    n1, n2 = st.columns(2)
    if n1.button("🔒 Crea nuovo codice per la carta", use_container_width=True, key=f"ferie_codice_{nome}"):
        _aggiorna_voce_ferie(nome, carta_codice=secrets.token_urlsafe(6))
        st.rerun()
    if codice and n2.button("Togli il codice (basta il nome)", use_container_width=True, key=f"ferie_togli_codice_{nome}"):
        _aggiorna_voce_ferie(nome, carta_codice=None)
        st.rerun()
    if codice:
        st.caption("⚠️ Se crei un nuovo codice, la carta va riscritta con il nuovo link.")


def _disegna_ferie_riepilogo(orari_lavoro, oggi):
    richieste = carica_richieste_ferie()
    saldi = carica_saldi_ferie()
    giustificativi = carica_correzioni()["giustificativi"]
    righe = []
    for nome in sorted(orari_lavoro.keys()):
        s = saldo_ferie(nome, oggi.year, orari_lavoro, giustificativi, richieste, saldi, oggi)
        voce = saldi.get(nome) or {}
        righe.append({
            "Dipendente": nome,
            "Spettanti": formatta_giorni(s["spettanti"]) if s["configurato"] else "—",
            "Residuo anno prec.": formatta_giorni(s["residuo"]) if s["configurato"] else "—",
            "Godute": formatta_giorni(s["godute"]), "Programmate": formatta_giorni(s["programmate"]),
            "In attesa": formatta_giorni(s["in_attesa"]),
            "Disponibili": formatta_giorni(s["disponibili"]) if s["configurato"] else "da impostare",
            "PIN": "✅" if voce.get("pin_hash") else "❌", "Email": "✅" if voce.get("email") else "—",
        })
    st.caption(f"Saldo ferie {oggi.year}. Godute = fino a oggi, programmate = giorni approvati futuri, "
               "disponibili = spettanti + residuo − godute − programmate.")
    st.dataframe(pd.DataFrame(righe), use_container_width=True, hide_index=True)


def _disegna_ferie_email():
    cfg = carica_config_ferie()
    st.caption("Notifiche per posta: all'ufficio quando arriva o viene annullata una richiesta (le invia l'app), "
               "al dipendente quando decidi (le invia la dashboard). Servono i dati di un account email dell'azienda "
               "(server SMTP). Le impostazioni stanno in ferie_config.json, nella cartella della dashboard.")
    with st.form("form_ferie_email"):
        attiva = st.checkbox("Invia le notifiche email", value=bool(cfg["attiva"]))
        c1, c2, c3 = st.columns([3, 1, 1.5])
        host = c1.text_input("Server SMTP", value=cfg["smtp_host"], placeholder="smtp.azienda.it")
        porta = c2.number_input("Porta", min_value=1, max_value=65535, value=int(cfg["smtp_porta"] or 465))
        sic = c3.selectbox("Sicurezza", ["ssl", "tls", "nessuna"], index=["ssl", "tls", "nessuna"].index(cfg["smtp_sicurezza"])
                           if cfg["smtp_sicurezza"] in ("ssl", "tls", "nessuna") else 0,
                           help="ssl = porta 465 · tls = porta 587 (STARTTLS)")
        u1, u2 = st.columns(2)
        utente = u1.text_input("Utente", value=cfg["smtp_utente"])
        password = u2.text_input("Password", type="password", placeholder="(lascia vuoto per non cambiarla)")
        m1, m2 = st.columns(2)
        mittente = m1.text_input("Indirizzo mittente", value=cfg["mittente"], placeholder="ferie@azienda.it")
        nome_mitt = m2.text_input("Nome del mittente", value=cfg["mittente_nome"])
        admin = st.text_input("Chi riceve gli avvisi dell'ufficio (separati da virgola)", value=", ".join(cfg["destinatari_admin"]))
        verifica = st.checkbox("Controlla il certificato del server", value=bool(cfg["verifica_certificato"]),
                               help="Toglila solo se il server usa un certificato interno/autofirmato.")
        url = st.text_input("Indirizzo dell'app ferie (per il link delle carte NFC)", value=cfg["url_app"],
                            placeholder="http://192.168.1.10/ziwood/ferie/index.html")
        salva = st.form_submit_button("💾 Salva impostazioni email", type="primary")
    if salva:
        cfg.update({
            "attiva": attiva, "smtp_host": host.strip(), "smtp_porta": int(porta), "smtp_sicurezza": sic,
            "smtp_utente": utente.strip(), "mittente": mittente.strip(), "mittente_nome": nome_mitt.strip(),
            "destinatari_admin": [a.strip() for a in admin.split(",") if a.strip()],
            "verifica_certificato": verifica, "url_app": url.strip(),
        })
        if password:
            cfg["smtp_password"] = password
        salva_config_ferie(cfg)
        st.success("Salvato.")
    if st.button("✉️ Invia una mail di prova"):
        cfg = carica_config_ferie()
        ok, errore = invia_mail_ferie(cfg["destinatari_admin"], "Prova notifiche ferie",
                                      "Se leggi questo messaggio, le notifiche email delle ferie funzionano.", cfg)
        if ok:
            st.success("Mail di prova inviata: controlla la posta di " + ", ".join(cfg["destinatari_admin"]))
        else:
            st.error(f"Invio non riuscito: {errore}")


def disegna_pagina_ferie(orari_lavoro, oggi):
    st.markdown("### 🏖️ Ferie dei dipendenti")
    if not st.session_state.get("orari_sbloccato"):
        st.caption("🔒 Questa pagina è protetta dalla stessa password di 'Orari Dipendenti'.")
        with st.form("form_password_ferie", clear_on_submit=True):
            password_inserita = st.text_input("Password", type="password")
            sblocca = st.form_submit_button("🔓 Sblocca")
        if sblocca:
            if password_inserita == carica_password():
                st.session_state["orari_sbloccato"] = True
                st.rerun()
            else:
                st.error("Password errata.")
        return
    if not orari_lavoro:
        st.info("Nessun dipendente disponibile: inseriscili prima in 'Orari Dipendenti'.")
        return
    messaggio = st.session_state.pop("_messaggio_ferie", None)
    if messaggio:
        getattr(st, messaggio[0])(messaggio[1])
    n_attesa = sum(1 for r in carica_richieste_ferie() if r["stato"] == "in_attesa")
    sez_richieste, sez_saldi, sez_riepilogo, sez_email = st.tabs(
        [f"📥 Richieste{f' ({n_attesa})' if n_attesa else ''}", "👥 Saldi e PIN", "📊 Riepilogo", "✉️ Notifiche email"])
    with sez_richieste:
        _disegna_ferie_richieste(orari_lavoro, oggi)
    with sez_saldi:
        _disegna_ferie_saldi(orari_lavoro, oggi)
    with sez_riepilogo:
        _disegna_ferie_riepilogo(orari_lavoro, oggi)
    with sez_email:
        _disegna_ferie_email()


def _attiva_correzione():
    st.session_state["modalita_correzione"] = True


def _disattiva_correzione():
    st.session_state["modalita_correzione"] = False
    st.session_state.pop("_dialogo_correzione", None)


def disegna_interruttore_correzione():
    """Pulsante in alto a destra per entrare/uscire dalla modalità correzione
    (stessa password di 'Orari Dipendenti')."""
    if st.session_state.get("modalita_correzione"):
        st.button("✏️ Correzione ATTIVA · esci", type="primary", use_container_width=True,
                  on_click=_disattiva_correzione, key="btn_esci_correzione")
    elif st.session_state.get("orari_sbloccato"):
        st.button("✏️ Modalità correzione", use_container_width=True, on_click=_attiva_correzione,
                  key="btn_entra_correzione")
    else:
        with st.popover("✏️ Modalità correzione", use_container_width=True):
            with st.form("form_password_modalita", clear_on_submit=True):
                password_inserita = st.text_input("Password", type="password")
                entra = st.form_submit_button("🔓 Attiva", type="primary")
            if entra:
                if password_inserita == carica_password():
                    st.session_state["orari_sbloccato"] = True
                    st.session_state["modalita_correzione"] = True
                    st.rerun()
                else:
                    st.error("Password errata.")


@st.cache_resource(show_spinner=False)
def _correggi_lingua_pagina_streamlit():
    """Il browser propone 'Tradurre la pagina?' perché la pagina di Streamlit
    dichiara di essere in inglese (<html lang="en">) mentre il testo è in
    italiano. Correggiamo UNA volta il file della pagina di Streamlit su questo
    computer: lingua italiana e 'non tradurre'. Vale dalla prossima apertura."""
    try:
        import streamlit as _stl
        percorso = os.path.join(os.path.dirname(_stl.__file__), "static", "index.html")
        with open(percorso, "r", encoding="utf-8") as f:
            testo = f.read()
        if 'translate="no"' in testo:
            return True
        import re as _re
        testo = _re.sub(r"<html[^>]*>", '<html lang="it" translate="no" class="notranslate">', testo, count=1)
        testo = _re.sub(r"(<head[^>]*>)", r'\1\n    <meta name="google" content="notranslate" />', testo, count=1)
        with open(percorso, "w", encoding="utf-8") as f:
            f.write(testo)
        return True
    except Exception:
        return False


def blocca_traduzione_browser():
    _correggi_lingua_pagina_streamlit()
    # anche per la pagina già aperta
    components.html(
        "<script>try{var h=window.parent.document.documentElement;"
        "h.setAttribute('lang','it');h.setAttribute('translate','no');h.classList.add('notranslate');}catch(e){}</script>",
        height=0,
    )


def _imposta_stato(**valori):
    """Usata dai pulsanti (on_click): aggiorna lo stato PRIMA che la pagina
    venga ridisegnata, così basta un solo aggiornamento per clic invece di
    due (prima ogni clic ridisegnava la pagina, poi la ridisegnava ancora)."""
    for chiave, valore in valori.items():
        st.session_state[chiave] = valore


ETICHETTE_MENU = {
    "oggi": "📅 Oggi", "ricerca": "🔍 Ricerca Storica", "orari": "🗓️ Orari Dipendenti",
    "correzioni": "✏️ Correzioni", "ferie": "🏖️ Ferie", "impostazioni": "⚙️ Impostazioni",
}


def crea_schede(etichette, key):
    """Schede "pigre": viene calcolata e disegnata SOLO la scheda aperta,
    non tutte insieme come prima (la Ricerca Storica con le sue 8
    sotto-tabelle veniva ricalcolata ad ogni clic anche se non la si stava
    guardando). Con versioni di Streamlit troppo vecchie per supportarlo si
    torna automaticamente al comportamento classico."""
    try:
        return st.tabs(etichette, key=key, on_change="rerun")
    except TypeError:
        return st.tabs(etichette)


def aperta(scheda):
    """True se la scheda è quella aperta (o se la versione di Streamlit non
    permette di saperlo: in quel caso si disegna sempre, come prima)."""
    return getattr(scheda, "open", None) is not False


# ----------------------------------------------------------------------------
# Interfaccia
# ----------------------------------------------------------------------------

inietta_css()
blocca_traduzione_browser()
st.markdown(f"<div class='versione-app'>v{VERSIONE_APP}</div>", unsafe_allow_html=True)
col_titolo_app, col_modalita_app = st.columns([4, 1.3], vertical_alignment="center")
with col_titolo_app:
    st.title("📊 Dashboard Presenze Aziendali")
with col_modalita_app:
    disegna_interruttore_correzione()

if st.session_state.get("modalita_correzione"):
    st.markdown(
        "<div class='banner-correzione'>✏️ <b>MODALITÀ CORREZIONE ATTIVA</b> – clicca la ✏️ accanto a un "
        "dipendente o a un giorno (pagina Oggi), oppure clicca un giorno in una tabella della Ricerca Storica: "
        "si apre la finestra per modificare, eliminare o aggiungere timbrature.</div>",
        unsafe_allow_html=True,
    )
pagina_oggi, pagina_ricerca, pagina_orari, pagina_correzioni, pagina_ferie, pagina_impostazioni = crea_schede(
    list(ETICHETTE_MENU.values()), key="menu_principale"
)

with pagina_impostazioni:
    st.markdown("### Impostazioni")

    st.markdown("#### Origine dei dati di presenza")
    percorso_cartella_salvato = carica_percorso_cartella_log()
    percorso_cartella_input = st.text_input(
        "📁 Cartella di rete con le timbrature grezze (opzionale)",
        value=percorso_cartella_salvato,
        placeholder=r"\\ServerNas\web\ziwood\presenze",
        help="Se indichi qui la cartella dove l'app scrive un CSV per ogni "
             "timbratura (la stessa che Power Query legge dentro presenze.xlsx "
             "per la query 'presenze'), la dashboard leggerà i dati direttamente "
             "da lì: sempre aggiornati in tempo reale, senza dover aprire Excel "
             "e premere 'Aggiorna dati'. Lascia vuoto per continuare a usare il "
             "file Excel come in passato.",
        key="input_cartella_log_grezze",
    )
    if percorso_cartella_input != percorso_cartella_salvato:
        salva_percorso_cartella_log(percorso_cartella_input)
        st.cache_data.clear()
        # Niente st.rerun() qui di proposito: forzare un altro giro completo
        # (che rilegge da capo tutti i dati, compresa la cartella di rete)
        # prima di mostrare l'esito qui sotto rendeva più fragile e più
        # lento tutto il flusso. Il valore appena digitato è già disponibile
        # in questa stessa esecuzione dello script, quindi la diagnostica
        # qui sotto lo usa subito, senza bisogno di un altro rerun.

    try:
        if percorso_cartella_input:
            diag = diagnostica_cartella_log(percorso_cartella_input)
            if diag["errore"]:
                st.error(f"⚠️ Errore leggendo la cartella: {diag['errore']}")
            elif not diag["raggiungibile"]:
                st.warning(
                    "⚠️ Questa cartella non risulta raggiungibile da qui in questo momento "
                    "(dal processo che fa girare la dashboard, anche se magari la vedi "
                    "normalmente in Esplora File): nel frattempo si continua a usare il "
                    "file Excel."
                )
            elif diag["n_file"] == 0:
                st.warning("⚠️ La cartella è raggiungibile ma non contiene nessun file .csv.")
            else:
                data_fmt = (
                    diag["data_piu_recente"].strftime("%d/%m/%Y alle %H:%M")
                    if diag["data_piu_recente"] else "sconosciuta"
                )
                st.success(
                    f"✅ Cartella raggiungibile: trovati **{diag['n_file']} file CSV**. "
                    f"Il più recente è `{diag['file_piu_recente']}` (modificato il {data_fmt})."
                )
                if diag["file_piu_recente_leggibile"] is False:
                    st.error(
                        f"⚠️ Attenzione: il file più recente (`{diag['file_piu_recente']}`) "
                        "non è stato letto correttamente - probabilmente ha un formato "
                        "diverso da quello atteso (intestazione 'Data;Orario;Nome;Evento'). "
                        "Mandami questo file così controllo cosa contiene."
                    )
    except Exception as e:
        # Rete di sicurezza: se per un motivo imprevisto il controllo della
        # cartella fallisce in un modo non già gestito sopra, mostriamo
        # comunque qualcosa di visibile invece di lasciare la sezione vuota
        # in silenzio (come sembra essere successo finora).
        st.error(
            "⚠️ Si è verificato un errore imprevisto controllando la cartella di rete. "
            "Mandami uno screenshot di questo messaggio (compreso il dettaglio tecnico "
            "qui sotto) così lo correggo."
        )
        st.exception(e)

    st.markdown("#### File Excel (usato per gli 'Orari Dipendenti' iniziali, e come fonte dati se la cartella sopra non è impostata)")
    uploaded_file = st.file_uploader("Carica file Excel presenze", type=["xlsx", "xls"])
    soglia_ritardo_input = a_time(imp("soglia_ritardo_default")) or SOGLIA_RITARDO_DEFAULT
    st.caption(
        "Se non carichi un file viene usato automaticamente `presenze.xlsx` "
        "nella cartella dell'app (utile per l'uso su NAS con file aggiornato "
        "automaticamente dall'app Android)."
    )
    if st.button("🔄 Ricarica dati"):
        st.cache_data.clear()
        st.rerun()

    disegna_impostazioni_regole()
    disegna_impostazioni_festivi()

try:
    percorso_trovato = None
    file_bytes, file_name = None, None
    if uploaded_file is not None:
        file_bytes = uploaded_file.getvalue()
        file_name = uploaded_file.name
    else:
        percorsi_possibili = ["presenze.xlsx", "dati/presenze.xlsx"]
        percorso_trovato = next((p for p in percorsi_possibili if os.path.exists(p)), None)
        if percorso_trovato is not None:
            with open(percorso_trovato, "rb") as f:
                file_bytes = f.read()
            file_name = percorso_trovato

    oggi = datetime.date.today()

    cartella_log = carica_percorso_cartella_log()
    elenco_cartella = elenco_file_cartella(cartella_log) if cartella_log else None
    diag_cartella_log = diagnostica_cartella_log(cartella_log, elenco_cartella) if cartella_log else None
    usa_cartella_log = bool(diag_cartella_log) and diag_cartella_log["raggiungibile"] and diag_cartella_log["n_file"] > 0
    if not usa_cartella_log and file_bytes is None:
        raise FileNotFoundError()

    orari_lavoro = ottieni_orari_lavoro(file_bytes, file_name) if file_bytes is not None else carica_orari_da_json()
    alias_nomi = carica_alias_nomi()
    correzioni = carica_correzioni()
    giustificativi = correzioni.get("giustificativi", {})

    if usa_cartella_log:
        voci_cartella = elenco_cartella["voci"]
        risultato_dati = prepara_dati(
            "cartella", cartella_log, _firma_elenco(voci_cartella), voci_cartella, None, None,
            orari_lavoro, alias_nomi, soglia_ritardo_input, oggi, IMPOSTAZIONI, correzioni,
        )
    else:
        risultato_dati = prepara_dati(
            "excel", None, None, None, file_bytes, file_name,
            orari_lavoro, alias_nomi, soglia_ritardo_input, oggi, IMPOSTAZIONI, correzioni,
        )
    df = risultato_dati["df"]
    nomi_grezzi_trovati = risultato_dati["nomi_grezzi"]

    if risultato_dati["esito"] == "vuoto":
        st.error(
            "⚠️ Nessun dato valido trovato" + (
                " nella cartella di rete indicata nelle Impostazioni."
                if usa_cartella_log else
                " nel file Excel. Controlla che il file contenga almeno un "
                "foglio con le colonne 'Data' e 'Nome'."
            )
        )
        st.stop()
    if risultato_dati["esito"] == "nessuna_corrispondenza":
        st.error(
            "⚠️ 'Orari Dipendenti' contiene dei nomi ma nessuno corrisponde a "
            "quelli presenti nel log presenze. Controlla che siano scritti "
            "allo stesso modo."
        )
        st.stop()

    dipendenti_censiti = sorted(orari_lavoro.keys()) if orari_lavoro else sorted(df["Nome"].unique())
    origine_file_locale = percorso_trovato if uploaded_file is None else None

    # ==========================================================================
    # PAGINA 1 — OGGI
    # ==========================================================================
    with pagina_oggi:
        if aperta(pagina_oggi):
            date_disponibili = set(df["Data"].dt.date)
            data_piu_recente_nei_dati = max(date_disponibili) if date_disponibili else None
            if usa_cartella_log:
                st.caption(
                    f"📡 Fonte dati attiva: **cartella di rete** (`{cartella_log}`). "
                    f"Il giorno più recente presente nei dati è "
                    f"**{data_piu_recente_nei_dati.strftime('%d/%m/%Y') if data_piu_recente_nei_dati else 'nessuno'}**."
                )
            else:
                dettaglio_motivo = ""
                if cartella_log and diag_cartella_log:
                    if diag_cartella_log["errore"]:
                        dettaglio_motivo = f" Motivo tecnico: `{diag_cartella_log['errore']}`."
                    elif not diag_cartella_log["raggiungibile"]:
                        dettaglio_motivo = (
                            " La cartella non risulta raggiungibile da qui in questo momento "
                            "(nessun errore specifico riportato da Windows, ma il controllo ha "
                            "dato esito negativo)."
                        )
                    elif diag_cartella_log["n_file"] == 0:
                        dettaglio_motivo = " La cartella è raggiungibile ma risulta vuota (0 file .csv)."
                st.caption(
                    f"📄 Fonte dati attiva: **file Excel** (`{file_name or 'nessuno'}`) - la cartella di "
                    f"rete impostata in Impostazioni non è (ancora) in uso.{dettaglio_motivo} "
                    f"Il giorno più recente presente nei dati è "
                    f"**{data_piu_recente_nei_dati.strftime('%d/%m/%Y') if data_piu_recente_nei_dati else 'nessuno'}**."
                )
            scelta_apertura = imp("giorno_apertura")
            if scelta_apertura == "oggi":
                data_predefinita = oggi
            elif scelta_apertura == "ultimo_lavorativo":
                data_predefinita = oggi - datetime.timedelta(days=1)
                while data_predefinita.weekday() >= 5:
                    data_predefinita -= datetime.timedelta(days=1)
            else:
                data_predefinita = oggi - datetime.timedelta(days=1)
            if "giorno_rif" not in st.session_state:
                st.session_state["giorno_rif"] = data_predefinita
            giorno_rif = st.session_state["giorno_rif"]

            col_fr_ind, col_data_sel, col_cal, col_fr_succ = st.columns([0.08, 0.66, 0.18, 0.08])
            with col_fr_ind:
                st.button(
                    "◀", key="giorno_precedente", use_container_width=True, on_click=_imposta_stato,
                    kwargs={"giorno_rif": giorno_rif - datetime.timedelta(days=1)},
                )
            with col_data_sel:
                nome_giorno = GIORNI_IT_COMPLETI[giorno_rif.weekday()]
                st.markdown(
                    f"<div style='text-align:center;font-size:1.8em;font-weight:800;"
                    f"color:#14324d;line-height:1.3;padding-top:0.1rem'>📅 {nome_giorno} {giorno_rif.strftime('%d/%m/%Y')}</div>",
                    unsafe_allow_html=True,
                )
            with col_cal:
                st.markdown("<div style='height:0.6rem'></div>", unsafe_allow_html=True)
                with crea_popover("🗓️ Calendario", key="popover_calendario_oggi"):
                    disegna_calendario(
                        giorno_rif, oggi, date_disponibili, key_prefix="oggi",
                        chiave_popover="popover_calendario_oggi",
                    )
            with col_fr_succ:
                st.button(
                    "▶", key="giorno_successivo", use_container_width=True, on_click=_imposta_stato,
                    kwargs={"giorno_rif": giorno_rif + datetime.timedelta(days=1)},
                )

            if nome_festivo(giorno_rif):
                _festa = nome_festivo(giorno_rif)
                st.info("🎉 Giorno festivo" + ("" if _festa == "Festivo" else f": **{_festa}**") + " (nessuno risulta assente).")
            elif giorno_rif != oggi:
                st.caption(f"ℹ️ Stai visualizzando una data diversa da oggi ({oggi.strftime('%d/%m/%Y')}).")

            df_oggi = df[df["Data"].dt.date == giorno_rif]
            st.session_state.setdefault("dipendente_dettaglio", None)
            dip_sel = st.session_state.get("dipendente_dettaglio")

            col_lista, col_log = st.columns([1, 1])

            with col_lista:
                if dip_sel and dip_sel in dipendenti_censiti:
                    # ---- Vista dettaglio dipendente: sostituisce l'elenco ----
                    if st.session_state.get("dettaglio_dip_corrente") != dip_sel:
                        st.session_state["dettaglio_dip_corrente"] = dip_sel
                        st.session_state["dettaglio_periodo"] = (giorno_rif.year, giorno_rif.month)
                    anno_d, mese_d = st.session_state.get("dettaglio_periodo", (giorno_rif.year, giorno_rif.month))

                    col_titolo, col_corr, col_chiudi = st.columns([3, 1.4, 0.6])
                    col_titolo.markdown(f"##### 👤 {dip_sel}")
                    col_corr.button(
                        "✏️ Correggi giorno", key="vai_correzioni_dettaglio", use_container_width=True,
                        help="Apre la scheda Correzioni su questo dipendente e sul giorno selezionato",
                        on_click=vai_a_correzioni, args=(dip_sel, giorno_rif),
                    )
                    col_chiudi.button(
                        "✖", key="chiudi_dettaglio", help="Torna all'elenco", on_click=_imposta_stato,
                        kwargs={"dipendente_dettaglio": None},
                    )

                    orario_dip = (orari_lavoro or {}).get(dip_sel, {})
                    pezzi_orario = []
                    if orario_dip.get("default"):
                        pezzi_orario.append(f"Ingresso {formatta_orario(orario_dip['default'])}")
                    if orario_dip.get("pausa_inizio") and orario_dip.get("pausa_fine"):
                        pezzi_orario.append(
                            f"Pausa {formatta_orario(orario_dip['pausa_inizio'])}–{formatta_orario(orario_dip['pausa_fine'])}"
                        )
                    if orario_dip.get("uscita"):
                        pezzi_orario.append(f"Uscita {formatta_orario(orario_dip['uscita'])}")
                    if pezzi_orario:
                        st.caption("Orario previsto: " + " · ".join(pezzi_orario))
                    for avviso in controlla_orari({dip_sel: orario_dip} if orario_dip else {}):
                        st.warning("⚠️ Orario impostato male in 'Orari Dipendenti': " + avviso +
                                   " Per questo le timbrature di uscita non vengono riconosciute.")

                    col_m_prec, col_m_lab, col_m_succ = st.columns([0.18, 0.64, 0.18])
                    with col_m_prec:
                        st.button(
                            "◀", key="mese_prec_dip", use_container_width=True, on_click=_imposta_stato,
                            kwargs={"dettaglio_periodo": _mese_precedente(anno_d, mese_d)},
                        )
                    with col_m_lab:
                        st.markdown(
                            f"<div style='text-align:center;font-weight:600;padding-top:0.35rem'>{MESI_IT[mese_d]} {anno_d}</div>",
                            unsafe_allow_html=True,
                        )
                    with col_m_succ:
                        al_mese_corrente = (anno_d, mese_d) >= (oggi.year, oggi.month)
                        st.button(
                            "▶", key="mese_succ_dip", use_container_width=True, disabled=al_mese_corrente,
                            on_click=_imposta_stato, kwargs={"dettaglio_periodo": _mese_successivo(anno_d, mese_d)},
                        )

                    giorni_nel_mese_d = calendar.monthrange(anno_d, mese_d)[1]
                    ultimo_giorno_d = oggi.day if (anno_d, mese_d) == (oggi.year, oggi.month) else giorni_nel_mese_d

                    df_dip_mese = df[
                        (df["Nome"] == dip_sel) & (df["Data"].dt.year == anno_d) & (df["Data"].dt.month == mese_d)
                    ].copy()
                    df_dip_mese["Giorno"] = df_dip_mese["Data"].dt.day

                    giorni_trascorsi = [
                        g for g in range(1, ultimo_giorno_d + 1)
                        if lavorativo_il(orari_lavoro, dip_sel, datetime.date(anno_d, mese_d, g))
                    ]
                    ore_totali_dip = df_dip_mese["DurataTotale_sec"].sum() / 3600
                    ritardi_dip = int(df_dip_mese["InRitardo"].sum())
                    anomalie_dip = int(df_dip_mese["HaAnomalia"].sum())
                    righe_per_giorno_d = {int(r_["Giorno"]): r_ for _, r_ in df_dip_mese.iterrows()}
                    codici_mese = [
                        codice_presenza(dip_sel, datetime.date(anno_d, mese_d, g), righe_per_giorno_d.get(g),
                                        orari_lavoro, giustificativi, oggi)
                        for g in range(1, ultimo_giorno_d + 1)
                    ]
                    assenti_dip = codici_mese.count("A")
                    malattia_ferie = codici_mese.count("M") + codici_mese.count("F") + codici_mese.count("F½")

                    k1, k2, k3, k4 = st.columns(4)
                    k1.metric("Ore (mese)", f"{ore_totali_dip:.1f} h")
                    k2.metric("Assenze", assenti_dip,
                              help=f"Malattia/Ferie nel mese: {malattia_ferie} (non contano come assenze)")
                    k3.metric("Ritardi", ritardi_dip)
                    k4.metric("Anomalie", anomalie_dip)

                    with st.container(border=True):
                        for g in range(ultimo_giorno_d, 0, -1):
                            data_g = datetime.date(anno_d, mese_d, g)
                            weekday = data_g.weekday()
                            lavorativo = lavorativo_il(orari_lavoro, dip_sel, data_g)
                            riga_g = df_dip_mese[df_dip_mese["Giorno"] == g]
                            etichetta_giorno = f"{GIORNI_IT[weekday]} {g:02d}/{mese_d:02d}"

                            if st.session_state.get("modalita_correzione"):
                                col_gg, col_dd, col_ed = st.columns([1.1, 2.4, 0.35], vertical_alignment="center")
                                col_ed.button("✏️", key=f"edit_dett_{g}", help="Correggi questo giorno",
                                              on_click=apri_correzione, args=(dip_sel, data_g))
                            else:
                                col_gg, col_dd = st.columns([1.1, 2.4])
                            with col_gg:
                                st.markdown(f"**{etichetta_giorno}**")

                            codice_g = codice_presenza(
                                dip_sel, data_g, None if riga_g.empty else riga_g.iloc[0],
                                orari_lavoro, giustificativi, oggi,
                            )
                            if riga_g.empty:
                                with col_dd:
                                    st.markdown(html_stato_giorno(codice_g, lavorativo, data_g), unsafe_allow_html=True)
                            else:
                                r = riga_g.iloc[0]
                                entrata = orario_con_matita(r, "Entrata1")
                                inizio_pausa = orario_con_matita(r, "Uscita1")
                                fine_pausa = orario_con_matita(r, "Entrata2")
                                # se c'è il rientro dalla pausa ma manca l'uscita finale, NON mostriamo
                                # come "uscita" l'inizio pausa (sarebbe fuorviante)
                                uscita_finale = (
                                    orario_con_matita(r, "Uscita2") if a_time(r["Entrata2"]) is not None
                                    else orario_con_matita(r, "UltimaUscita_time")
                                )
                                pausa_txt = f"{inizio_pausa}–{fine_pausa}" if fine_pausa != "-" else "-"
                                badge = ""
                                if r["InRitardo"]:
                                    badge += '<span class="badge-ritardo">Ritardo</span>'
                                if r["RitardoFinePausa"]:
                                    badge += '<span class="badge-ritardo">Ritardo fine pausa</span>'
                                if r["UscitaAnticipata"]:
                                    badge += '<span class="badge-uscita">Uscita anticipata</span>'
                                badge += badge_codice(codice_g) + badge_mancanze(r)
                                if r.get("CorrettoAMano"):
                                    badge += '<span class="badge-giust">✏️ corretto a mano</span>'
                                colore_giorno = _colore_giorno(r)
                                with col_dd:
                                    st.markdown(
                                        f'<span class="dot dot-{colore_giorno}"></span>'
                                        f'Entrata {entrata} · Pausa {pausa_txt} · Uscita {uscita_finale}{badge}',
                                        unsafe_allow_html=True,
                                    )
                else:
                    # ---- Elenco dipendenti (vista normale) ----
                    st.markdown("##### Elenco Dipendenti")
                    st.caption("Clicca su un nome per vederne i dettagli giorno per giorno.")
                    with st.container(border=True):
                        for nome in dipendenti_censiti:
                            riga = df_oggi[df_oggi["Nome"] == nome]
                            weekday = giorno_rif.weekday()
                            lavorativo = lavorativo_il(orari_lavoro, nome, giorno_rif)

                            if st.session_state.get("modalita_correzione"):
                                col_nome, col_edit, col_stato = st.columns([2, 0.45, 3], vertical_alignment="center")
                                col_edit.button("✏️", key=f"edit_oggi_{nome}", help="Correggi le timbrature di questo giorno",
                                                on_click=apri_correzione, args=(nome, giorno_rif))
                            else:
                                col_nome, col_stato = st.columns([2, 3])
                            with col_nome:
                                st.button(
                                    nome, key=f"sel_{nome}", use_container_width=True, on_click=_imposta_stato,
                                    kwargs={"dipendente_dettaglio": nome},
                                )

                            codice_oggi = codice_presenza(
                                nome, giorno_rif, None if riga.empty else riga.iloc[0],
                                orari_lavoro, giustificativi, oggi,
                            )
                            if riga.empty:
                                with col_stato:
                                    st.markdown(html_stato_giorno(codice_oggi, lavorativo, giorno_rif), unsafe_allow_html=True)
                            else:
                                r = riga.iloc[0]
                                eventi_riga = _lista_timbrature(r)
                                badge = badge_codice(codice_oggi) + badge_mancanze(r)
                                if r.get("CorrettoAMano"):
                                    badge += '<span class="badge-giust">✏️ corretto a mano</span>'
                                if r["InRitardo"]:
                                    badge += '<span class="badge-ritardo">Ritardo</span>'
                                if r["RitardoFinePausa"]:
                                    badge += '<span class="badge-ritardo">Ritardo fine pausa</span>'
                                if r["UscitaAnticipata"]:
                                    badge += '<span class="badge-uscita">Uscita anticipata</span>'
                                if eventi_riga:
                                    manuali_riga = r.get("OrariManuali") or []
                                    timbrature_txt = " · ".join(
                                        f"{tipo} {t.strftime('%H:%M')}{' ✏️' if t.strftime('%H:%M') in manuali_riga else ''}"
                                        for t, tipo, _ in eventi_riga
                                    )
                                else:
                                    timbrature_txt = "-"
                                sotto = f"{timbrature_txt}{badge}"
                                with col_stato:
                                    st.markdown(
                                        f'<span class="dot dot-{"arancio" if badge_mancanze(r) else "verde"}"></span>Presente<br>'
                                        f'<span class="sotto-nome">{sotto}</span>',
                                        unsafe_allow_html=True,
                                    )

            with col_log:
                dip_solo_log = dip_sel if (dip_sel and dip_sel in dipendenti_censiti) else None
                if dip_solo_log:
                    st.markdown(f"##### Log Dettagliato — {dip_solo_log}")
                else:
                    st.markdown("##### Log Dettagliato del Giorno")

                eventi_oggi = []
                nomi_da_includere = [dip_solo_log] if dip_solo_log else dipendenti_censiti
                for nome in nomi_da_includere:
                    riga = df_oggi[df_oggi["Nome"] == nome]
                    if riga.empty:
                        continue
                    r = riga.iloc[0]
                    for t, etichetta, stato, segmento in _lista_timbrature_grezze(r):
                        colore = _colore_evento(r, segmento) if segmento else None
                        eventi_oggi.append((t, etichetta, nome, stato, colore))
                eventi_oggi.sort(key=lambda x: x[0], reverse=True)

                if not eventi_oggi:
                    if dip_solo_log:
                        st.info(f"Nessuna timbratura registrata per {dip_solo_log} in questo giorno.")
                    else:
                        st.info("Nessuna timbratura registrata per questo giorno.")
                else:
                    st.caption(
                        "🟢 conforme · 🟡 ritardo fine pausa · 🟠 ritardo entrata/uscita anticipata · 🔴 anomalia grave · "
                        "❌ doppia timbratura (esclusa dal calcolo) · 🔧 = corretta automaticamente"
                    )
                    righe_log = ['<div class="card">']
                    for orario, etichetta, nome, stato, colore in eventi_oggi:
                        descr_nome = "" if dip_solo_log else f" - {nome}"
                        if stato == "doppia":
                            righe_log.append(
                                f'<div class="timeline-riga timeline-riga-scartata">'
                                f'<span class="timeline-orario">{orario.strftime("%H:%M")}</span>'
                                f'<span class="timeline-punto timeline-punto-rosso"></span>'
                                f'<span>❌ Doppia timbratura{descr_nome}</span></div>'
                            )
                        elif stato == "extra":
                            righe_log.append(
                                f'<div class="timeline-riga">'
                                f'<span class="timeline-orario">{orario.strftime("%H:%M")}</span>'
                                f'<span class="timeline-punto timeline-punto-arancio"></span>'
                                f'<span>⚠️ Timbratura non riconosciuta{descr_nome}</span></div>'
                            )
                        else:
                            badge_corretto = ' <span class="badge-ritardo">🔧 Corretto</span>' if stato == "corretto" else ""
                            righe_log.append(
                                f'<div class="timeline-riga">'
                                f'<span class="timeline-orario">{orario.strftime("%H:%M")}</span>'
                                f'<span class="timeline-punto timeline-punto-{colore or "verde"}"></span>'
                                f'<span>{etichetta}{descr_nome}{badge_corretto}</span></div>'
                            )
                    righe_log.append("</div>")
                    st.markdown("".join(righe_log), unsafe_allow_html=True)

    # ==========================================================================
    # PAGINA 2 — RICERCA STORICA
    # ==========================================================================
    with pagina_ricerca:
        if aperta(pagina_ricerca):
            df["AnnoMese_Raw"] = df["Data"].dt.strftime("%Y-%m")
            lista_mesi_raw = sorted(df["AnnoMese_Raw"].unique(), reverse=True)
            opzioni_mesi = {
                f"{MESI_IT.get(int(m.split('-')[1]), m.split('-')[1])} {m.split('-')[0]}": m
                for m in lista_mesi_raw
            }

            st.markdown('<div class="card">', unsafe_allow_html=True)
            col_mese, col_dip = st.columns([1.2, 2])
            with col_mese:
                # La scelta viene ricordata anche passando ad altre schede e tornando qui.
                elenco_mesi = list(opzioni_mesi.keys())
                mese_ricordato = st.session_state.get("_ricerca_mese_scelto")
                mese_scelto_nome = st.selectbox(
                    "📅 Ricerca veloce per mese", elenco_mesi,
                    index=elenco_mesi.index(mese_ricordato) if mese_ricordato in elenco_mesi else 0,
                    key="widget_ricerca_mese",
                )
                st.session_state["_ricerca_mese_scelto"] = mese_scelto_nome
            mese_scelto_raw = opzioni_mesi[mese_scelto_nome]
            anno, mese = map(int, mese_scelto_raw.split("-"))
            giorni_nel_mese = calendar.monthrange(anno, mese)[1]

            with col_dip:
                dipendenti_ricordati = [
                    n for n in st.session_state.get("_ricerca_dipendenti_scelti", []) if n in dipendenti_censiti
                ]
                if "widget_ricerca_dipendenti" not in st.session_state:
                    st.session_state["widget_ricerca_dipendenti"] = dipendenti_ricordati
                else:
                    st.session_state["widget_ricerca_dipendenti"] = [
                        n for n in st.session_state["widget_ricerca_dipendenti"] if n in dipendenti_censiti
                    ]
                dipendenti_selezionati = st.multiselect(
                    "🔎 Dipendenti (uno o più, vuoto = tutti)",
                    options=dipendenti_censiti,
                    placeholder="Scrivi per cercare un nome...",
                    key="widget_ricerca_dipendenti",
                )
                st.session_state["_ricerca_dipendenti_scelti"] = list(dipendenti_selezionati)
            st.markdown("</div>", unsafe_allow_html=True)

            df_mese_completo = df[df["AnnoMese_Raw"] == mese_scelto_raw].copy()
            df_mese_completo["Giorno"] = df_mese_completo["Data"].dt.day
            colonne_giorni = list(range(1, giorni_nel_mese + 1))

            dipendenti_filtrati = dipendenti_selezionati if dipendenti_selezionati else \
                sorted(df_mese_completo["Nome"].unique())

            df_mese = df_mese_completo[
                df_mese_completo["Nome"].isin(dipendenti_filtrati)
            ].copy()

            if not dipendenti_filtrati:
                st.warning("Nessun dipendente disponibile per questo mese.")
                st.stop()

            giorni_weekend = {g for g in colonne_giorni
                              if datetime.date(anno, mese, g).weekday() >= 5 or nome_festivo(datetime.date(anno, mese, g))}

            ore_totali_mese = df_mese["DurataTotale_sec"].sum() / 3600
            n_anomalie = int(df_mese["HaAnomalia"].sum())
            n_ritardi = int(df_mese["InRitardo"].sum())

            col_k1, col_k2, col_k3, col_k4 = st.columns(4)
            col_k1.metric("👥 Dipendenti", len(dipendenti_filtrati))
            col_k2.metric("⏱️ Ore totali", f"{ore_totali_mese:,.1f} h".replace(",", "."))
            col_k3.metric("⏳ Ritardi", n_ritardi)
            col_k4.metric("⚠️ Anomalie", int(n_anomalie))

            st.markdown("---")

            if len(dipendenti_selezionati) == 1:
                disegna_scheda_dipendente(
                    dipendenti_selezionati[0], df_mese, anno, mese, colonne_giorni, giorni_weekend,
                    orari_lavoro, giustificativi, oggi, mese_scelto_raw,
                )
            else:
                st.caption("👆 Nelle tabelle: clicca un **nome** per vedere solo quel dipendente con tutte le sue "
                           "informazioni; clicca un **giorno** per aprirlo nella pagina Oggi.")
                (
                    sub_ore, sub_presenza, sub_entrate, sub_uscite, sub_pausa_in, sub_pausa_fine,
                    sub_riepilogo, sub_riepilogo_generale,
                ) = crea_schede([
                    "🕒 Ore Effettuate", "✅ Presente / ❌ Assente", "⏳ Orario Entrata", "🚪 Orario Uscita",
                    "☕ Inizio Pausa", "☕ Fine Pausa", "📋 Riepilogo per Dipendente", "🗂️ Riepilogo Generale",
                ], key="sottomenu_ricerca")

                def applica_stile_celle(styler, func):
                    if hasattr(styler, "map"):
                        return styler.map(func)
                    return styler.applymap(func)

                def stile_weekend(styler_data):
                    def _stile(col):
                        if col.name in giorni_weekend:
                            return ["background-color: rgba(128,128,128,0.15)"] * len(col)
                        return [""] * len(col)
                    return styler_data.apply(_stile, axis=0)

                def stile_riga_presenza(row):
                    stili = []
                    for giorno in row.index:
                        val = row[giorno]
                        if val == "P":
                            stile = "background-color: rgba(0,170,0,0.20)"
                        elif val == "½":
                            stile = "background-color: rgba(240,190,0,0.30)"
                        elif val == "M":
                            stile = "background-color: rgba(142,68,173,0.22)"
                        elif val in ("F", "F½"):
                            stile = "background-color: rgba(52,152,219,0.25)"
                        elif val == "-":
                            stile = "background-color: rgba(150,150,150,0.10); color: rgba(120,120,120,0.8)"
                        else:
                            stile = "background-color: rgba(200,0,0,0.15)"
                        if giorno in giorni_weekend:
                            stile += ";border-bottom: 3px solid rgba(90,90,90,0.5)"
                        stili.append(stile)
                    return stili

                def stile_riga_entrate(row):
                    nome = row.name
                    stili = []
                    for giorno in row.index:
                        val = row[giorno]
                        t = a_time(_senza_matita(val)) if val != "-" else None
                        weekday = datetime.date(anno, mese, giorno).weekday()
                        soglia = soglia_ritardo_per(orari_lavoro, nome, weekday, soglia_ritardo_input)
                        stile = ""
                        if soglia is not None and t is not None and t > _sposta_orario(soglia, imp("tolleranza_ritardo_ingresso_min")):
                            stile = "background-color: rgba(230,150,0,0.25)"
                        if giorno in giorni_weekend:
                            stile += (";" if stile else "") + "border-bottom: 3px solid rgba(90,90,90,0.5)"
                        stili.append(stile)
                    return stili

                def stile_riga_uscita(row):
                    nome = row.name
                    stili = []
                    for giorno in row.index:
                        val = row[giorno]
                        t = a_time(_senza_matita(val)) if val != "-" else None
                        stile = ""
                        atteso = orari_lavoro.get(nome, {}).get("uscita") if orari_lavoro else None
                        if atteso is not None and t is not None:
                            soglia = _sposta_orario(atteso, -float(imp("tolleranza_uscita_anticipata_min")))
                            if t < soglia:
                                stile = "background-color: rgba(200,0,0,0.15)"
                        if giorno in giorni_weekend:
                            stile += (";" if stile else "") + "border-bottom: 3px solid rgba(90,90,90,0.5)"
                        stili.append(stile)
                    return stili

                def stile_riga_pausa(row, campo_orario):
                    nome = row.name
                    stili = []
                    for giorno in row.index:
                        val = row[giorno]
                        t = a_time(_senza_matita(val)) if val != "-" else None
                        stile = ""
                        atteso = orari_lavoro.get(nome, {}).get(campo_orario) if orari_lavoro else None
                        if atteso is not None and t is not None:
                            soglia = _sposta_orario(atteso, float(imp("tolleranza_ritardo_fine_pausa_min")))
                            if t > soglia:
                                stile = "background-color: rgba(230,150,0,0.25)"
                        if giorno in giorni_weekend:
                            stile += (";" if stile else "") + "border-bottom: 3px solid rgba(90,90,90,0.5)"
                        stili.append(stile)
                    return stili

                with sub_ore:
                    if aperta(sub_ore):
                        df_mese["Durata_Fmt"] = df_mese.apply(ore_con_matita, axis=1) if not df_mese.empty else []
                        matrice_ore = df_mese.pivot_table(index="Nome", columns="Giorno", values="Durata_Fmt", aggfunc="first")
                        matrice_ore = matrice_ore.reindex(index=dipendenti_filtrati, columns=colonne_giorni).fillna("-")
                        mostra_matrice_cliccabile(matrice_ore, "ore", anno, mese, giorni_weekend)
                        st.download_button(
                            "⬇️ Scarica CSV", matrice_ore.to_csv().encode("utf-8-sig"),
                            file_name=f"ore_{mese_scelto_raw}.csv", mime="text/csv",
                        )

                with sub_presenza:
                    if aperta(sub_presenza):
                        matrice_presenze = matrice_codici_presenza(
                            df_mese, dipendenti_filtrati, anno, mese, colonne_giorni, orari_lavoro, giustificativi, oggi
                        )
                        mostra_matrice_cliccabile(matrice_presenze, "presenze", anno, mese, giorni_weekend, stile_riga_presenza)
                        st.caption(
                            "**P** presente · **½** mezza giornata (ore sotto la soglia impostata) · **A** assente · "
                            "**M** malattia · **F** ferie (Malattia e Ferie si inseriscono nella scheda ✏️ Correzioni) · "
                            "**-** giorno futuro o di riposo (non conta come assenza)."
                        )

                with sub_entrate:
                    if aperta(sub_entrate):
                        df_mese["Entrata1_Fmt"] = df_mese.apply(lambda r: orario_con_matita(r, "Entrata1"), axis=1) if not df_mese.empty else []
                        matrice_entrate = df_mese.pivot_table(index="Nome", columns="Giorno", values="Entrata1_Fmt", aggfunc="first")
                        matrice_entrate = matrice_entrate.reindex(index=dipendenti_filtrati, columns=colonne_giorni).fillna("-")
                        mostra_matrice_cliccabile(matrice_entrate, "entrate", anno, mese, giorni_weekend, stile_riga_entrate)
                        st.caption("🟧 Evidenziati gli ingressi oltre l'orario personalizzato di ciascun dipendente.")

                with sub_uscite:
                    if aperta(sub_uscite):
                        df_mese["UscitaFinale_Fmt"] = df_mese.apply(lambda r: orario_con_matita(r, "UltimaUscita_time"), axis=1) if not df_mese.empty else []
                        matrice_uscite = df_mese.pivot_table(index="Nome", columns="Giorno", values="UscitaFinale_Fmt", aggfunc="first")
                        matrice_uscite = matrice_uscite.reindex(index=dipendenti_filtrati, columns=colonne_giorni).fillna("-")
                        mostra_matrice_cliccabile(matrice_uscite, "uscite", anno, mese, giorni_weekend, stile_riga_uscita)
                        st.caption("🟥 Evidenziate le uscite anticipate rispetto all'orario personalizzato di ciascun dipendente.")

                with sub_pausa_in:
                    if aperta(sub_pausa_in):
                        df_mese["Pausa1_Fmt"] = df_mese.apply(lambda r: orario_con_matita(r, "Uscita1"), axis=1) if not df_mese.empty else []
                        matrice_pausa_in = df_mese.pivot_table(index="Nome", columns="Giorno", values="Pausa1_Fmt", aggfunc="first")
                        matrice_pausa_in = matrice_pausa_in.reindex(index=dipendenti_filtrati, columns=colonne_giorni).fillna("-")
                        mostra_matrice_cliccabile(matrice_pausa_in, "pausa_in", anno, mese, giorni_weekend)
                        st.caption("Orario di inizio pausa registrato (non segnalato come ritardo).")

                with sub_pausa_fine:
                    if aperta(sub_pausa_fine):
                        df_mese["Pausa2_Fmt"] = df_mese.apply(lambda r: orario_con_matita(r, "Entrata2"), axis=1) if not df_mese.empty else []
                        matrice_pausa_fine = df_mese.pivot_table(index="Nome", columns="Giorno", values="Pausa2_Fmt", aggfunc="first")
                        matrice_pausa_fine = matrice_pausa_fine.reindex(index=dipendenti_filtrati, columns=colonne_giorni).fillna("-")
                        mostra_matrice_cliccabile(
                            matrice_pausa_fine, "pausa_fine", anno, mese, giorni_weekend,
                            lambda row: stile_riga_pausa(row, "pausa_fine"),
                        )
                        st.caption("🟧 Evidenziato il rientro dalla pausa oltre l'orario previsto per ciascun dipendente.")

                with sub_riepilogo:
                    if aperta(sub_riepilogo):
                        riepilogo = (
                            df_mese.groupby("Nome")
                            .agg(
                                Giorni_Presenti=("DurataTotale_sec", lambda x: int((x > 0).sum())),
                                Ore_Totali=("DurataTotale_sec", lambda x: round(x.sum() / 3600, 2)),
                                Ritardi=("InRitardo", "sum"),
                                Anomalie=("HaAnomalia", "sum"),
                            )
                            .reindex(dipendenti_filtrati)
                            .fillna(0)
                        )
                        riepilogo["Ritardi"] = riepilogo["Ritardi"].astype(int)
                        riepilogo["Anomalie"] = riepilogo["Anomalie"].astype(int)

                        def _giorni_lavorativi_attesi(nome):
                            return sum(
                                1 for g in colonne_giorni
                                if datetime.date(anno, mese, g) <= oggi
                                and lavorativo_il(orari_lavoro, nome, datetime.date(anno, mese, g))
                            )

                        codici_tutti = matrice_codici_presenza(
                            df_mese, dipendenti_filtrati, anno, mese, colonne_giorni, orari_lavoro, giustificativi, oggi
                        )
                        riepilogo["Giorni_Assenti"] = [int((codici_tutti.loc[n] == "A").sum()) for n in riepilogo.index]
                        riepilogo["Mezze_Giornate"] = [int((codici_tutti.loc[n] == "½").sum()) for n in riepilogo.index]
                        riepilogo["Malattia"] = [int((codici_tutti.loc[n] == "M").sum()) for n in riepilogo.index]
                        riepilogo["Ferie"] = [float((codici_tutti.loc[n] == "F").sum() + 0.5 * (codici_tutti.loc[n] == "F½").sum()) for n in riepilogo.index]

                        ore_previste_giorno = {n: ore_previste_secondi(orari_lavoro, n) for n in riepilogo.index}
                        if any(v is not None for v in ore_previste_giorno.values()):
                            riepilogo["Ore_Previste"] = [
                                round(ore_previste_giorno[n] * _giorni_lavorativi_attesi(n) / 3600, 2)
                                if ore_previste_giorno[n] is not None else None
                                for n in riepilogo.index
                            ]
                            riepilogo["Scarto_Ore"] = riepilogo["Ore_Totali"] - riepilogo["Ore_Previste"]

                        mostra_matrice_cliccabile(riepilogo, "riepilogo", anno, mese, set())
                        st.download_button(
                            "⬇️ Scarica CSV", riepilogo.to_csv().encode("utf-8-sig"),
                            file_name=f"riepilogo_{mese_scelto_raw}.csv", mime="text/csv",
                        )

                        righe_anomale = df_mese[df_mese["HaAnomalia"]].copy()
                        if not righe_anomale.empty:
                            with st.expander(f"⚠️ Dettaglio {len(righe_anomale)} anomalie del periodo (entrata e uscita)"):
                                righe_anomale["Tipo"] = righe_anomale["TipiAnomalia"].map(lambda t: ", ".join(t))
                                righe_anomale["Entrata1"] = righe_anomale["Entrata1"].map(formatta_orario)
                                righe_anomale["Uscita1"] = righe_anomale["Uscita1"].map(formatta_orario)
                                st.dataframe(
                                    righe_anomale[["Data", "Nome", "Tipo", "Entrata1", "Uscita1"]].sort_values("Data"),
                                    use_container_width=True,
                                )

                with sub_riepilogo_generale:
                    if aperta(sub_riepilogo_generale):
                        st.caption(
                            "Per ogni giorno del mese: entrata, uscita e ore lavorate nella stessa cella, "
                            "su tre righe separate. Sempre visibili tutti i giorni del mese."
                        )
                        righe_html = []
                        for nome in dipendenti_filtrati:
                            dati_nome = df_mese[df_mese["Nome"] == nome]
                            celle = []
                            for g in colonne_giorni:
                                sotto = dati_nome[dati_nome["Giorno"] == g]
                                classe_we = " cella-weekend" if g in giorni_weekend else ""
                                if sotto.empty:
                                    celle.append(f'<td class="cella-riepilogo{classe_we}">-</td>')
                                    continue
                                r = sotto.iloc[0]
                                e = orario_con_matita(r, "Entrata1")
                                u = orario_con_matita(r, "UltimaUscita_time")
                                ore = ore_con_matita(r)
                                celle.append(
                                    f'<td class="cella-riepilogo{classe_we}">'
                                    f'<div class="riga-e"><span class="icona-mov">→</span>{e}</div>'
                                    f'<div class="riga-u"><span class="icona-mov">←</span>{u}</div>'
                                    f'<div class="riga-o"><span class="icona-mov">⏱</span>{ore}</div>'
                                    f'</td>'
                                )
                            nome_html = str(nome).replace("<", "&lt;").replace(">", "&gt;")
                            righe_html.append(
                                f'<tr><td class="cella-nome">{nome_html}</td>' + "".join(celle) + "</tr>"
                            )

                        intestazione = "".join(
                            f'<th class="{"cella-weekend" if g in giorni_weekend else ""}">{g}</th>'
                            for g in colonne_giorni
                        )
                        tabella_html = (
                            '<div class="tabella-riepilogo-wrap"><table class="tabella-riepilogo">'
                            f'<thead><tr><th class="cella-nome-header">Dipendente</th>{intestazione}</tr></thead>'
                            f'<tbody>{"".join(righe_html)}</tbody></table></div>'
                        )
                        st.markdown(tabella_html, unsafe_allow_html=True)
                        st.caption("→ Entrata · ← Uscita · ⏱ Ore lavorate")

    # ==========================================================================
    # PAGINA 3 — ORARI DIPENDENTI
    # ==========================================================================
    with pagina_orari:
        if aperta(pagina_orari):
            st.markdown("### Orario di lavoro previsto per dipendente")

            st.session_state.setdefault("orari_sbloccato", False)

            if not st.session_state["orari_sbloccato"]:
                st.caption(
                    "🔒 La modifica degli orari è protetta da password, per evitare che chiunque "
                    "possa cambiare gli orari registrati."
                )
                with st.form("form_password_orari", clear_on_submit=True):
                    password_inserita = st.text_input("Password", type="password")
                    sblocca = st.form_submit_button("🔓 Sblocca")
                if sblocca:
                    if password_inserita == carica_password():
                        st.session_state["orari_sbloccato"] = True
                        st.rerun()
                    else:
                        st.error("Password errata.")
                st.stop()

            col_sblocco_info, col_blocca = st.columns([4, 1])
            with col_sblocco_info:
                st.caption("🔓 Modifica sbloccata per questa sessione.")
            with col_blocca:
                if st.button("🔒 Blocca", use_container_width=True):
                    st.session_state["orari_sbloccato"] = False
                    st.rerun()

            with st.expander("🔑 Cambia password"):
                with st.form("form_cambia_password", clear_on_submit=True):
                    vecchia = st.text_input("Password attuale", type="password")
                    nuova = st.text_input("Nuova password", type="password")
                    nuova_conferma = st.text_input("Ripeti nuova password", type="password")
                    cambia = st.form_submit_button("Cambia password")
                if cambia:
                    if vecchia != carica_password():
                        st.error("La password attuale non è corretta.")
                    elif not nuova:
                        st.error("La nuova password non può essere vuota.")
                    elif nuova != nuova_conferma:
                        st.error("Le due password inserite non coincidono.")
                    else:
                        salva_password(nuova)
                        st.success("Password cambiata correttamente.")

            st.caption(
                "Qui puoi impostare l'orario di ciascun dipendente (ingresso, pausa, uscita) "
                "e i suoi giorni di riposo. Le modifiche vengono salvate in un file separato "
                "(`orari_lavoro.json`) e da quel momento hanno sempre la priorità, così un "
                "aggiornamento automatico di `presenze.xlsx` da parte dell'app Android non "
                "può mai sovrascriverle."
            )

            messaggio_salvataggio = st.session_state.pop("_messaggio_orari", None)
            if messaggio_salvataggio:
                getattr(st, messaggio_salvataggio[0])(messaggio_salvataggio[1])

            st.caption(
                "La **seconda colonna** dice cosa fare di ogni nome che compare nelle "
                "timbrature: **✅ Dipendente** (nome giusto, usa l'orario della riga), "
                "oppure scegli **il nome giusto** a cui corrisponde (per i nomi scritti male, "
                "es. 'Paolo Angona' → 'Paolo Ancona': le sue timbrature vengono sommate a "
                "quelle della persona giusta), oppure **🚫 Ignora** per scartarle. "
                "Le righe in cima con la seconda colonna **vuota** sono nomi nuovi trovati "
                "nelle timbrature e ancora da decidere: finché resta vuota non vengono conteggiati."
            )

            df_editor = orari_a_editor_df(orari_lavoro or {}, nomi_grezzi_trovati, alias_nomi)
            # Giorni di riposo: si scelgono da un elenco (Lun ... Dom), niente da scrivere a mano.
            if hasattr(st.column_config, "MultiselectColumn"):
                colonna_giorni_riposo = st.column_config.MultiselectColumn(
                    "Giorni di riposo", options=GIORNI_IT,
                    help="Clicca la cella e scegli i giorni dall'elenco (puoi sceglierne più di uno).",
                )
            else:
                df_editor["GiorniRiposo"] = df_editor["GiorniRiposo"].map(lambda l: ", ".join(l))
                colonna_giorni_riposo = st.column_config.TextColumn(
                    "Giorni di riposo", help="Abbreviazioni separate da virgola, es: Sab, Dom"
                )
            avvisi_orari = controlla_orari(orari_lavoro)
            if avvisi_orari:
                st.error("⚠️ Alcuni orari sono impostati in modo incoerente e vanno corretti nella tabella qui sotto:\n\n- "
                         + "\n- ".join(avvisi_orari))
            nomi_da_decidere = df_editor.loc[df_editor["CorrispondeA"].isna(), "Nome"].tolist()
            if nomi_da_decidere:
                st.warning(
                    f"⚠️ {len(nomi_da_decidere)} nome/i trovati nelle timbrature non sono ancora "
                    "assegnati (seconda colonna vuota): " + ", ".join(nomi_da_decidere)
                )

            nomi_ufficiali = sorted((orari_lavoro or {}).keys())
            opzioni_corrisponde = [ETICHETTA_E_DIPENDENTE] + nomi_ufficiali + [ETICHETTA_ALIAS_IGNORA]
            opzioni_corrisponde += sorted({
                v for v in (alias_nomi or {}).values() if v not in opzioni_corrisponde
            })

            editor_result = st.data_editor(
                df_editor,
                num_rows="dynamic",
                use_container_width=True,
                hide_index=True,
                column_order=COLONNE_EDITOR_ORARI,
                column_config={
                    "Nome": st.column_config.TextColumn("Nome (come nelle timbrature)", required=True),
                    "CorrispondeA": st.column_config.SelectboxColumn(
                        "È un dipendente / Corrisponde a",
                        options=opzioni_corrisponde,
                        help="✅ = nome giusto di un dipendente; un altro nome = questo è un "
                             "modo sbagliato di scrivere quel dipendente; 🚫 = ignora.",
                        width="large",
                    ),
                    "Ingresso": st.column_config.TimeColumn("Ingresso", format="HH:mm", step=60),
                    "InizioPausa": st.column_config.TimeColumn("Inizio pausa", format="HH:mm", step=60),
                    "FinePausa": st.column_config.TimeColumn("Fine pausa", format="HH:mm", step=60),
                    "Uscita": st.column_config.TimeColumn("Uscita", format="HH:mm", step=60),
                    "AltraPausaMin": st.column_config.NumberColumn("Altra pausa (min)", min_value=0, step=5),
                    "GiorniRiposo": colonna_giorni_riposo,
                },
                key="editor_orari",
            )

            col_salva, col_ripristina = st.columns([1, 1])
            with col_salva:
                if st.button("💾 Salva", type="primary"):
                    nuovo_orari, nuove_correzioni = editor_df_a_orari(editor_result)
                    destinazioni_sbagliate = sorted({
                        v for v in nuove_correzioni.values()
                        if v != ETICHETTA_ALIAS_IGNORA and v not in nuovo_orari
                    })
                    salva_orari_lavoro(nuovo_orari)
                    salva_alias_nomi(nuove_correzioni)
                    tipo, testo = "success", "✅ Salvato: orari e corrispondenze dei nomi aggiornati."
                    if origine_file_locale:
                        try:
                            scrivi_orari_in_excel(origine_file_locale, nuovo_orari)
                            testo += " Aggiornato anche il foglio 'OrariLavoro' del file Excel."
                        except Exception as e_excel:
                            tipo = "warning"
                            testo += (
                                " Non è stato possibile aggiornare anche il file Excel "
                                f"(magari è aperto in un altro programma): {e_excel}"
                            )
                    avvisi_nuovi = controlla_orari(nuovo_orari)
                    if avvisi_nuovi:
                        tipo = "warning"
                        testo += " ATTENZIONE, controlla questi orari: " + " ".join(avvisi_nuovi)
                    if destinazioni_sbagliate:
                        tipo = "warning"
                        testo += (
                            " Attenzione: questi nomi scelti nella seconda colonna non sono "
                            "(più) segnati come dipendenti, quindi le timbrature a loro "
                            "collegate non verranno conteggiate: " + ", ".join(destinazioni_sbagliate)
                        )
                    st.session_state["_messaggio_orari"] = (tipo, testo)
                    st.cache_data.clear()
                    st.rerun()
            with col_ripristina:
                if st.button("↺ Ripristina dal foglio Excel 'OrariLavoro'"):
                    if os.path.exists(PERCORSO_ORARI_SALVATI):
                        os.remove(PERCORSO_ORARI_SALVATI)
                    st.cache_data.clear()
                    st.rerun()

    with pagina_correzioni:
        if aperta(pagina_correzioni):
            disegna_pagina_correzioni(df, dipendenti_censiti, oggi, orari_lavoro)

    with pagina_ferie:
        if aperta(pagina_ferie):
            disegna_pagina_ferie(orari_lavoro, oggi)

    if st.session_state.get("modalita_correzione") and st.session_state.get("_dialogo_correzione"):
        dialogo_correzione(df, orari_lavoro, dipendenti_censiti)

except FileNotFoundError:
    st.error(
        "⚠️ Nessun file trovato. Carica un file Excel dalla barra laterale oppure "
        "posiziona un file chiamato `presenze.xlsx` nella stessa cartella dell'app."
    )
except Exception as e:
    st.error(f"Errore durante l'elaborazione del file: {e}")
    st.exception(e)
